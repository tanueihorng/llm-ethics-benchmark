"""IEEE citation-order gate for the REWRITTEN document line.

Why this exists: `verify_report_claims.py` sets ``BUILDER =
scripts/build_fyp_report_v5.js`` and its ``ieee_citations`` lock parses that
file's ``const refs = [`` block. It therefore checks the frozen v5 JavaScript
line and nothing else — the rewritten Python builders (v1-v4) and the DOCX they
produce have never been covered. On 2026-08-23 the v4 report shipped with its
first-use sequence running 1, 2, 3, 5, 6, 7, 8, 9, 10, 18, 24, 20, ... while
`make verify-claims` reported 93/93 green, because the gate could not see the
document. This is the same gate-gap shape as D54.

Two extraction traps this test is built to avoid, both of which under-report
citations and produce false "uncited reference" findings:

  * ``[1]-[3]`` range citations, which are valid IEEE and name every number in
    the span, not just the endpoints;
  * citations inside TABLE CELLS, which ``Document.paragraphs`` does not yield.

A FAIL means the document's reference numbering does not follow first use, or a
listed reference is never cited. Fix the numbering — never relax this test.
"""

from __future__ import annotations

import re
from pathlib import Path

import pytest

docx = pytest.importorskip("docx")
from docx import Document  # noqa: E402
from docx.table import Table  # noqa: E402
from docx.text.paragraph import Paragraph  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parents[1]

# The rewritten line. DOCX are untracked (D60), so a missing file is a skip, not
# a failure: CI legitimately checks fewer documents than a local run.
DOCUMENTS = [
    "docs/rewritten_fyp_report_v1.docx",
    "docs/rewritten_fyp_report_v2.docx",
    "docs/rewritten_fyp_report_v3.docx",
    "docs/rewritten_fyp_report_v4.docx",
    "docs/rewritten_fyp_report_v5.docx",
    "docs/rewritten_fyp_interim_v1.docx",
    "docs/rewritten_fyp_interim_v3.docx",
    "docs/rewritten_fyp_interim_v4.docx",
    "docs/rewritten_fyp_interim_v5.docx",
    "docs/rewritten_fyp_interim_compact.docx",
    "docs/rewritten_fyp_interim_progress.docx",
]

CITATION = re.compile(r"\[(\d+)\]")
CITATION_RANGE = re.compile(r"\[(\d+)\]\s*[–—-]\s*\[(\d+)\]")
REFERENCE_ENTRY = re.compile(r"^\[(\d+)\]\s+\S")


def expand_ranges(text: str) -> str:
    """``[1]-[3]`` names 1, 2 and 3 — not just the endpoints."""

    def expand(match: re.Match) -> str:
        low, high = int(match.group(1)), int(match.group(2))
        if high < low:
            return match.group(0)
        return "".join(f"[{n}]" for n in range(low, high + 1))

    return CITATION_RANGE.sub(expand, text)


def iter_block_paragraphs(document):
    """Every paragraph in body order, INCLUDING those inside tables."""
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, document).rows:
                for cell in row.cells:
                    yield from cell.paragraphs


def split_document(document):
    """One traversal, body order: (reference numbers, body paragraph texts).

    A single pass classifies each paragraph as it is reached, because matching
    paragraphs ACROSS two traversals is unsafe: python-docx builds a fresh
    wrapper each time, and ``id(paragraph._element)`` can be recycled once a
    temporary proxy is collected — which silently skips the wrong paragraphs.
    """
    numbers, body, seen_heading = [], [], False
    for paragraph in iter_block_paragraphs(document):
        text = paragraph.text.strip()
        if text == "References":
            seen_heading = True
            body.append(text)
            continue
        match = REFERENCE_ENTRY.match(text)
        if seen_heading and match:
            numbers.append(int(match.group(1)))
        else:
            body.append(text)
    return numbers, body


def first_use_order(body, listed):
    order = []
    for text in body:
        for match in CITATION.finditer(expand_ranges(text)):
            number = int(match.group(1))
            if number in listed and number not in order:
                order.append(number)
    return order


@pytest.mark.parametrize("relative", DOCUMENTS)
def test_references_numbered_by_first_use(relative: str) -> None:
    path = REPO_ROOT / relative
    if not path.exists():
        pytest.skip(f"{relative} not built (DOCX are untracked, D60)")

    document = Document(str(path))
    listed, body = split_document(document)
    assert listed, f"{relative}: no bibliography entries found under 'References'"
    assert listed == list(range(1, len(listed) + 1)), (
        f"{relative}: bibliography is not numbered 1..N; starts {listed[:12]}"
    )

    order = first_use_order(body, set(listed))
    uncited = sorted(set(listed) - set(order))
    assert not uncited, f"{relative}: references listed but never cited: {uncited}"
    assert order == listed, (
        f"{relative}: citations are not in strict first-use order.\n"
        f"  first-use sequence: {order[:14]}\n"
        f"  expected          : {listed[:14]}"
    )


def test_range_citations_are_expanded() -> None:
    """Self-test: the endpoint-only reading is what mis-diagnosed v4 as broken."""
    assert expand_ranges("cited [1]-[3] here") == "cited [1][2][3] here"
    assert expand_ranges("cited [8]–[10]") == "cited [8][9][10]"
    assert expand_ranges("[5] alone") == "[5] alone"


def _synthetic(body_lines, table_lines=()):
    """A minimal document: body paragraphs, an optional table, then References."""
    document = Document()
    for line in body_lines:
        document.add_paragraph(line)
    if table_lines:
        table = document.add_table(rows=len(table_lines), cols=1)
        for row, line in zip(table.rows, table_lines):
            row.cells[0].text = line
    document.add_paragraph("References")
    document.add_paragraph("[1] A. Author, \u201cFirst work,\u201d 2020.")
    document.add_paragraph("[2] B. Author, \u201cSecond work,\u201d 2021.")
    document.add_paragraph("[3] C. Author, \u201cThird work,\u201d 2022.")
    return document


def test_gate_fires_on_out_of_order_citations() -> None:
    """Self-test: the exact v4 defect must fail, or this file is decoration."""
    document = _synthetic(["Sentence one [2].", "Sentence two [1].", "Sentence three [3]."])
    listed, body = split_document(document)
    assert listed == [1, 2, 3]
    assert first_use_order(body, set(listed)) == [2, 1, 3] != listed


def test_gate_sees_citations_inside_tables() -> None:
    """Self-test: a table-blind scan reports [3] uncited, which is how v4 was mis-read."""
    document = _synthetic(["Sentence one [1].", "Sentence two [2]."], table_lines=["Row citing [3]."])
    listed, body = split_document(document)
    order = first_use_order(body, set(listed))
    assert order == [1, 2, 3], "citation inside a table cell was not seen"
    assert not set(listed) - set(order)


def test_gate_excludes_the_reference_list_itself() -> None:
    """Self-test: the entries' own [N] prefixes must not count as citations."""
    document = _synthetic(["Only [3] is cited in the body."])
    listed, body = split_document(document)
    order = first_use_order(body, set(listed))
    assert order == [3], f"reference-list prefixes leaked into the body scan: {order}"
    assert sorted(set(listed) - set(order)) == [1, 2]


@pytest.mark.parametrize(
    "relative",
    ["docs/rewritten_fyp_report_v4.docx", "docs/rewritten_fyp_interim_v4.docx"],
)
def test_v4_citations_preserve_following_prose_spacing(relative: str) -> None:
    """Citation renumbering must not glue a marker to the next word."""
    path = REPO_ROOT / relative
    if not path.exists():
        pytest.skip(f"{relative} not built (DOCX are untracked, D60)")
    document = Document(str(path))
    _, body = split_document(document)
    glued = [text for text in body if re.search(r"\]\w", text)]
    assert not glued, f"{relative}: citation marker glued to following text: {glued[:3]}"


def test_v4_opening_names_and_cites_all_model_families() -> None:
    """The opening must cite the four families IN THE ORDER IT NAMES THEM.

    A range citation defeats the first-use-order lock: ``[1]-[4]`` expands to
    1, 2, 3, 4 and passes however the names beside it are arranged. v4 shipped
    naming Mistral third and Phi fourth while [3] was Phi and [4] was Mistral.
    This checks the pairing the range hides, by resolving each reference number
    against the bibliography entry it actually points to.
    """
    path = REPO_ROOT / "docs/rewritten_fyp_report_v4.docx"
    if not path.exists():
        pytest.skip("v4 report not built (DOCX are untracked, D60)")
    document = Document(str(path))
    _, body = split_document(document)
    opening = body[body.index("1.1 Background and Motivation") + 1]
    assert "[1]–[4]" in opening

    # split_document returns reference NUMBERS; the entry text is needed here,
    # so collect it in the same single-pass, body-order way.
    entries: dict[int, str] = {}
    seen_heading = False
    for paragraph in iter_block_paragraphs(document):
        text = paragraph.text.strip()
        if text == "References":
            seen_heading = True
            continue
        match = REFERENCE_ENTRY.match(text)
        if seen_heading and match:
            entries[int(match.group(1))] = text

    families = ["Qwen3", "the Llama 3 family", "Phi-4-mini", "Mistral-7B"]
    named = ", ".join(families[:-1]) + ", and " + families[-1]
    assert named in opening, f"opening does not name the families in citation order: {opening[:220]}"

    # Each [n] is resolved against the bibliography entry it actually points
    # to, rather than against a hard-coded number.
    expected_source = {
        "Qwen3": "Qwen3",
        "the Llama 3 family": "Llama 3",
        "Phi-4-mini": "Phi-4-Mini",
        "Mistral-7B": "Mistral 7B",
    }
    for position, family in enumerate(families, start=1):
        entry = entries.get(position, "")
        assert expected_source[family] in entry, (
            f"the opening names {family} in position {position}, but reference "
            f"[{position}] is {entry[:90]!r}"
        )


def test_v4_problem_statement_uses_academic_capability_wording() -> None:
    """The report should use reader-facing terminology rather than repository dialect."""
    path = REPO_ROOT / "docs/rewritten_fyp_report_v4.docx"
    if not path.exists():
        pytest.skip("v4 report not built (DOCX are untracked, D60)")
    document = Document(str(path))
    _, body = split_document(document)
    text = "\n".join(body).lower()
    assert "capability benchmarks" in text
    assert "capability anchors" not in text
