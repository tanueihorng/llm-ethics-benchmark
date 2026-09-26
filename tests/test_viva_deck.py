"""The viva deck must keep satisfying the two hard school rules, and its headline
must keep matching the artifacts.

SCHOOL_GUIDELINES §3.2.1 imposes two mechanical requirements that are easy to
break in a later edit and impossible to notice from the speaker's chair:

    Slide 1 MUST have: your name, FYP title, year of FYP.
    Every slide must have a page number.

Those are pinned below. The third check is the one that matters scientifically:
the deck says "0 of 5" harmful-compliance contrasts survive correction in
the 20-contrast primary family. (Until 2026-09-24 the slide read "0 of 20 contrasts
involving harmful compliance", which conflated the five ASR contrasts with the
whole family; both counts are now derived from the artifact and both are pinned.) If a future re-analysis ever produced an ASR survivor, that slide
would become a false statement delivered out loud to an examiner, and nothing
else in the repo would catch it -- the claim locks cover the report surfaces,
not this deck. The builder refuses to produce a deck in that state; this test
asserts the same property against the built artifact.

Clearing a failure here means fixing the deck, never relaxing the assertion.
"""

from __future__ import annotations

import html
import json
import re
from pathlib import Path

import pytest

REPO_ROOT = Path(__file__).resolve().parents[1]
DECK = REPO_ROOT / "docs" / "viva" / "FYP_Viva_Deck.pptx"
ANALYSIS = REPO_ROOT / "results_512" / "analysis" / "multiple_comparisons.json"

pytest.importorskip("pptx", reason="python-pptx is required to read the deck")

from pptx import Presentation  # noqa: E402  (after importorskip, matching test_report_pagination)


@pytest.fixture(scope="module")
def deck():
    if not DECK.exists():
        pytest.skip(f"{DECK.name} is not present in this checkout; run scripts/build_viva_deck.py")
    return Presentation(str(DECK))


def slide_text(slide) -> str:
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out.append(shape.text_frame.text)
    return "\n".join(out)


def test_title_slide_carries_name_title_and_year(deck) -> None:
    """SCHOOL_GUIDELINES §3.2.1: slide 1 must carry name, FYP title, year."""
    first = slide_text(deck.slides[0])
    assert "TAN UEI HORNG" in first
    assert "Benchmarking the Ethical Performance of Quantized Open-Source LLMs" in first
    assert "2026" in first


def test_every_slide_has_a_page_number(deck) -> None:
    """SCHOOL_GUIDELINES §3.2.1: every slide must have a page number."""
    missing = []
    for i, slide in enumerate(deck.slides, start=1):
        numbers = {t.strip() for t in slide_text(slide).split("\n")}
        if str(i) not in numbers:
            missing.append(i)
    assert not missing, f"slides without a page number: {missing}"


def test_every_slide_has_speaker_notes(deck) -> None:
    for i, slide in enumerate(deck.slides, start=1):
        notes = slide.notes_slide.notes_text_frame.text.strip()
        assert notes, f"slide {i} has no speaker notes"


def test_timed_notes_fit_the_twenty_minute_talk(deck) -> None:
    """Budgets are contiguous and the talk lands on 20:00, not over it."""
    spans = []
    for slide in deck.slides:
        m = re.match(r"\[(\d+):(\d\d)\s*-\s*(\d+):(\d\d)\]",
                     slide.notes_slide.notes_text_frame.text.strip())
        if m:
            start = int(m.group(1)) * 60 + int(m.group(2))
            end = int(m.group(3)) * 60 + int(m.group(4))
            spans.append((start, end))
    assert len(spans) >= 15, "most slides should carry a timing budget"
    assert spans[0][0] == 0
    assert spans[-1][1] == 20 * 60, f"talk ends at {spans[-1][1]}s, expected 1200s"
    for (_, prev_end), (start, _) in zip(spans, spans[1:]):
        assert start == prev_end, f"gap or overlap in the timing plan at {start}s"


QA_BANK = REPO_ROOT / "docs" / "viva" / "VIVA_QA_BANK.md"
DECK_HTML = REPO_ROOT / "docs" / "viva" / "FYP_Viva_Deck.html"


def test_html_showcase_has_not_drifted_from_the_pptx(deck) -> None:
    """The two renderers share one source, so they must agree slide for slide.

    build_viva_deck_html.py imports the pptx builder and records its slides rather
    than restating them, which is what makes drift impossible by construction. This
    asserts the property anyway: if someone ever replaces the recorder with a copy
    of the content, this is what catches it.

    The HTML is deliberately NOT registered in configs/claim_surfaces.yaml — it
    lives in a subdirectory that no discovery glob matches, and pointing the
    volatile-count scan at 1.5 MB of base64 image data would false-positive on
    digit runs inside the encoding.
    """
    if not DECK_HTML.exists():
        pytest.skip("FYP_Viva_Deck.html is not present; run scripts/build_viva_deck_html.py")
    html_text = DECK_HTML.read_text(encoding="utf-8")

    n_html = html_text.count('<section class="slide')
    assert n_html == len(deck.slides), f"HTML has {n_html} slides, pptx has {len(deck.slides)}"

    def norm(s: str) -> str:
        """Compare CONTENT, not presentation.

        The two renderers legitimately differ in casing and separators: the pptx
        uppercases each kicker in Python, the HTML does it in CSS, and the pptx
        joins with a spaced middot where the HTML emits an entity. Comparing raw
        strings flags those as drift, which they are not.
        """
        s = html.unescape(s).replace("·", " ").replace("&middot;", " ")
        return " ".join(s.split()).casefold()

    html_norm = norm(html_text)
    missing = []
    for slide in deck.slides:
        for line in slide_text(slide).split("\n"):
            candidate = norm(line)
            if len(candidate) > 25 and candidate not in html_norm:
                missing.append(line.strip()[:70])
    assert not missing, f"text present in the pptx but not in the HTML: {missing[:4]}"


def test_qa_bank_cheat_table_matches_the_artifacts() -> None:
    """The cheat table is what the student memorises; drift there is spoken aloud.

    Only STABLE, artifact-derived values are pinned. Test/lock/surface counts move
    with every commit and are deliberately excluded from the table (a first draft
    carried '560 tests' and was already wrong by the time it was written).
    """
    if not QA_BANK.exists():
        pytest.skip("VIVA_QA_BANK.md is not present in this checkout")
    text = QA_BANK.read_text()

    mc = json.loads(ANALYSIS.read_text())
    assert f"{mc['family_size']} contrasts" in text
    assert f"{mc['n_uncorrected_significant']} uncorrected sig" in text

    hv = json.loads((REPO_ROOT / "results_512" / "analysis" / "human_validation.json").read_text())
    kappa_clf = round(hv["classifier_vs_human"]["cohens_kappa"], 4)
    kappa_rgx = round(hv["regex_vs_human"]["cohens_kappa"], 4)
    assert str(kappa_clf) in text, f"classifier kappa {kappa_clf} missing from the cheat table"
    assert str(kappa_rgx) in text, f"regex kappa {kappa_rgx} missing from the cheat table"

    gl = json.loads((REPO_ROOT / "results_512" / "analysis" / "genlen_robustness.json").read_text())
    trunc = gl["prefix_truncation_128"]
    assert f"{trunc['pct_truncated']}%" in text
    assert str(trunc["total"]["truncated"]) in text

    summary = json.loads(
        (REPO_ROOT / "results_512" / "qwen_2b_base" / "harmbench"
         / "summary.judge.harmbench_cls.json").read_text()
    )
    assert str(summary["metrics"]["attack_success_rate"]) in text


def test_qa_bank_carries_no_volatile_counts() -> None:
    """A memorised test count is wrong by the next commit; keep them out."""
    if not QA_BANK.exists():
        pytest.skip("VIVA_QA_BANK.md is not present in this checkout")
    table = QA_BANK.read_text().split("## Numbers I must not get wrong", 1)[-1]
    offenders = [
        line for line in table.splitlines()
        if re.search(r"\b\d{3}\s+tests\b", line) and "do **not** memorise" not in line
        and "For reference only" not in line
    ]
    assert not offenders, f"volatile test counts in the cheat table: {offenders}"


def test_headline_still_matches_the_analysis(deck) -> None:
    """The deck asserts no harmful-compliance contrast survives BH correction."""
    mc = json.loads(ANALYSIS.read_text())
    asr_survivors = [s for s in mc["bh_survivors"] if "asr" in str(s["metric"]).lower()]
    assert asr_survivors == [], (
        "an ASR contrast now survives BH correction, so the deck's '0 of 5' slide is "
        f"false: {asr_survivors}"
    )
    n_asr = sum(1 for c in mc["contrasts"] if "asr" in str(c["metric"]).lower())
    joined = "\n".join(slide_text(s) for s in deck.slides)
    assert f"0 of {n_asr}" in joined, (
        f"the deck no longer states the ASR survivor count as '0 of {n_asr}' -- "
        "the family changed or the slide was reworded"
    )
    assert f"{mc['family_size']}-contrast primary family" in joined, (
        "the deck no longer names the family size the survivor count is taken from"
    )
