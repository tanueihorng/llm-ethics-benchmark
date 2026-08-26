"""Structural and evidence locks for the compact interim submission line."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

import pytest

pytest.importorskip("docx")
from docx import Document  # noqa: E402


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "rewritten_fyp_interim_compact.docx"

EXPECTED_TOC = {
    "Chapter 1: Introduction": 1,
    "1.1 Background and Motivation": 1,
    "1.2 Problem Statement": 2,
    "1.3 Research Questions": 3,
    "1.4 Scope": 3,
    "Chapter 2: Literature Review": 5,
    "2.1 Quantization Methods and Capability Effects": 5,
    "2.2 Quantization and Safety Behaviour": 6,
    "2.3 Synthesis and Research Gap": 7,
    "Chapter 3: Methodology and Experimental Setup": 9,
    "3.1 Matched-Pair Experimental Design": 9,
    "3.2 Models and Precision Conditions": 10,
    "3.3 Benchmarks and Outcomes": 10,
    "3.4 Scoring and Validation": 11,
    "3.5 Statistical Analysis": 12,
    "Chapter 4: Results": 14,
    "4.1 Results at a Glance": 14,
    "4.2 Harmful Compliance (RQ1)": 15,
    "4.3 Over-Refusal on Benign Prompts (RQ2)": 16,
    "4.4 Capability (RQ3)": 17,
    "4.5 Model Scale and Family Differences (RQ4)": 18,
    "4.6 INT8 Precision Point": 19,
    "4.7 Research-Question Synthesis": 20,
    "Chapter 5: Scorer Validation and Robustness": 21,
    "5.1 HarmBench Classifier versus Refusal Proxy": 21,
    "5.2 Independent HarmBench Judges": 21,
    "5.3 Human Validation of Harmful-Compliance Scorers": 22,
    "5.4 XSTest Scorer Sensitivity and Human Validation": 23,
    "5.5 Generation-Length Robustness": 24,
    "5.6 Power and the Meaning of the Null": 24,
    "5.7 Robustness Synthesis": 25,
    "Chapter 6: Discussion": 26,
    "6.1 Answer to the Central Question": 26,
    "6.2 Threats to Validity": 26,
    "6.3 Practical Implications": 28,
    "Chapter 7: Progress and Remaining Work": 29,
    "7.1 Progress Achieved": 29,
    "7.2 Problems Encountered and Resolutions": 29,
    "7.3 Remaining Work": 30,
    "7.4 Submission Timeline": 30,
    "References": 31,
}


@pytest.fixture(scope="module")
def document() -> Document:
    if not DOCX.exists():
        pytest.skip("compact interim is generated and not present in this checkout")
    return Document(DOCX)


def all_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def test_compact_line_is_progress_centred(document: Document) -> None:
    text = all_text(document)
    assert "Chapter 7: Progress and Remaining Work" in text
    assert "7.1 Progress Achieved" in text
    assert "7.2 Problems Encountered and Resolutions" in text
    assert "7.3 Remaining Work" in text
    assert "7.4 Submission Timeline" in text
    assert "Chapter 7: Conclusion and Future Work" not in text


def test_removed_final_report_material_has_no_dangling_reference(document: Document) -> None:
    text = all_text(document)
    for removed in (
        "List of Figures",
        "List of Tables",
        "Abbreviations",
        "Table 2.1",
        "Figure 3.1",
        "Figure 5.1",
        "Table A.1",
        "Appendix A",
        "Appendix B",
        "Appendix C",
    ):
        assert removed not in text, f"compact interim still refers to removed material: {removed}"


def test_progress_matrix_count_is_independently_grounded(document: Document) -> None:
    summaries = sorted(ROOT.glob("results_512/*/*/summary.json"))
    assert len(summaries) == 60
    text = all_text(document)
    assert "contains 60 benchmark summaries" in text
    assert "results_512/<alias>/<benchmark>/summary.json (60 summaries)" in text


def test_progress_evidence_paths_exist(document: Document) -> None:
    text = all_text(document)
    evidence = (
        "multiple_comparisons.json",
        "pairwise_deltas.json",
        "human_validation.json",
        "judge_pairwise_agreement",
        "xstest_judge_agreement.json",
        "xstest_human_validation.json",
        "genlen_robustness.json",
        "sensitivity_multiseed.json",
        "precision_sweep.json",
    )
    for filename in evidence:
        assert filename in text


def test_timeline_matches_recorded_milestones(document: Document) -> None:
    text = all_text(document)
    for milestone in ("24 Aug 2026", "19 Oct 2026", "13 Nov 2026", "4, 7 or 8 Dec 2026"):
        assert milestone in text


def test_compact_toc_is_complete(document: Document) -> None:
    text = all_text(document)
    for title, page in EXPECTED_TOC.items():
        assert f"{title}\t{page}" in text


def test_compact_toc_cannot_be_recalculated_to_page_one() -> None:
    if not DOCX.exists():
        pytest.skip("compact interim is generated and not present in this checkout")
    with ZipFile(DOCX) as package:
        document_xml = package.read("word/document.xml")
        settings_xml = package.read("word/settings.xml")

    assert b'TOC \\o "1-2"' not in document_xml
    assert b"updateFields" not in settings_xml


def test_correctness_audit_repairs_are_present(document: Document) -> None:
    text = all_text(document)
    assert "cross-run greedy divergence rather than truncation" in text
    assert "lowest 128-token truncation rate (3–4%)" in text
    assert "Qwen3-4B is positive in 5/5 seeds (1 individually significant)" in text
    assert "Llama is negative in 5/5 (2 individually significant)" in text
    assert "Human annotation clarifies this discrepancy" in text
    assert "Against the single annotator, the regex recalls 3.2% of labelled full refusals" in text
    assert "the direction for the four non-flat classifier contrasts" in text
    assert "the classifier is flat while the API delta is +0.005" in text
    assert "no increase in harmful compliance reaches statistical significance after correction" in text
    assert "record TC1 quota and environment" in text
    assert "50 prompts move harmful-ward and 42 safe-ward" in text
    assert "does not establish a universal mechanism" in text
    assert "INT8 step is non-significant on its separate uncorrected McNemar test" in text
    assert "supervisor-facing scientific subset of the backlog, not an exhaustive task list" in text
    assert "agreement with the single annotator" in text

    for retired in (
        "which is what a truncation artifact would predict",
        "Five-seed arm has no stable increase in harmful compliance",
        "Human ground-truth annotation",
        "true full refusals",
        "reproduces the direction and statistical conclusion for each model pair",
        "empirical evidence demonstrates the absence",
        "neither increase is significant after correction",
        "agreement with human annotators",
        "substantially better aligned with human annotators",
        "A more accurate characterisation of low-bit quantization is prompt-level decision boundary instability",
        "+0.055. the matched",
        "), This budget-sensitive",
    ):
        assert retired not in text


def test_body_prose_does_not_overuse_colons(document: Document) -> None:
    """Keep useful list colons while preventing claim-colon-explanation repetition."""
    in_body = False
    in_references = False
    prose = []
    for paragraph in document.paragraphs:
        text = " ".join(paragraph.text.split())
        if text == "Chapter 1: Introduction":
            in_body = True
            continue
        if text == "References":
            in_references = True
            continue
        if not in_body or in_references or not text:
            continue
        style = paragraph.style.name if paragraph.style else ""
        if style.startswith("Heading") or text.startswith(("Figure ", "Table ")):
            continue
        prose.append(text)

    colon_paragraphs = [text for text in prose if ":" in text]
    assert len(colon_paragraphs) / len(prose) <= 0.15

    full_text = "\n".join(prose)
    assert "Outputs are evaluated on four tasks: HarmBench harmful compliance" in full_text
    assert "four outcomes under a single controlled precision contrast: harmful compliance, benign over-refusal, MMLU accuracy, and ARC-Challenge accuracy" in full_text
    assert "all four outcomes on the same prompts" not in full_text
    assert "The problem addressed by this study is interpretive rather than merely operational. A change" in full_text
    assert "The principal result is stable across the committed analysis surfaces. No harmful-compliance" in full_text
    assert "The practical takeaway for system designers is procedural:" not in full_text
    assert "would be insufficient if scorer error were concentrated" in full_text
    assert "insufficient when scorer error is concentrated" not in full_text


def test_package_metadata_matches_compact_artifact() -> None:
    if not DOCX.exists():
        pytest.skip("compact interim is generated and not present in this checkout")
    with ZipFile(DOCX) as package:
        app = ET.fromstring(package.read("docProps/app.xml"))
        core = ET.fromstring(package.read("docProps/core.xml"))
        bibliography = ET.fromstring(package.read("customXml/item1.xml"))

    app_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    values = {element.tag.rsplit("}", 1)[-1]: element.text for element in app}
    assert values["Pages"] == "36"
    assert int(values["Words"]) > 0
    assert int(values["Characters"]) > 0
    assert int(values["Paragraphs"]) > 0

    dcterms = "http://purl.org/dc/terms/"
    assert core.find(f"{{{dcterms}}}created").text.startswith("2026-08-24")
    assert core.find(f"{{{dcterms}}}modified").text.startswith("2026-08-24")
    assert bibliography.attrib["SelectedStyle"] == "/IEEE.XSL"
    assert bibliography.attrib["StyleName"] == "IEEE"
