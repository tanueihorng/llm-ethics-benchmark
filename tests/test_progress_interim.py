"""Scope, navigation, and evidence locks for the user-directed interim."""

from __future__ import annotations

from pathlib import Path
from zipfile import ZipFile

import pytest

pytest.importorskip("docx")
from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "rewritten_fyp_interim_progress.docx"

EXPECTED_TOC = {
    "Chapter 1: Introduction": 1,
    "1.1 Background and Motivation": 1,
    "1.2 Problem Statement": 2,
    "1.3 Research Questions": 3,
    "1.4 Scope": 3,
    "1.5 Contributions": 4,
    "Chapter 2: Literature Review": 5,
    "2.1 Trustworthiness and Safety Evaluation": 5,
    "2.2 Quantization Methods and Capability Effects": 6,
    "2.3 Quantization and Safety Behaviour": 7,
    "2.4 Refusal, Harmfulness, and Threat Models": 8,
    "2.5 Automatic Evaluators and Human Validation": 10,
    "2.6 Synthesis and Research Gap": 11,
    "Chapter 3: Methodology and Experimental Setup": 13,
    "3.1 Matched-Pair Experimental Design": 13,
    "3.2 Models and Precision Conditions": 14,
    "3.3 Benchmarks and Outcomes": 15,
    "3.4 Scoring and Validation": 16,
    "3.5 Statistical Analysis": 16,
    "3.6 Implementation and Execution": 18,
    "Chapter 4: Results": 19,
    "4.1 Results at a Glance": 19,
    "4.2 Harmful Compliance (RQ1)": 20,
    "4.3 Over-Refusal on Benign Prompts (RQ2)": 21,
    "4.4 Capability (RQ3)": 22,
    "4.5 Model Scale and Family Differences (RQ4)": 23,
    "4.6 INT8 Precision Point": 24,
    "4.7 Research-Question Synthesis": 25,
    "Chapter 5: Progress and Remaining Work": 26,
    "5.1 Progress Achieved": 26,
    "5.2 Remaining Work": 26,
    "References": 28,
}


@pytest.fixture(scope="module")
def document() -> Document:
    if not DOCX.exists():
        pytest.skip("progress interim is generated and not present in this checkout")
    return Document(DOCX)


def all_text(document: Document) -> str:
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.extend(paragraph.text for paragraph in cell.paragraphs)
    return "\n".join(parts)


def test_requested_chapter_contract(document: Document) -> None:
    headings = {
        paragraph.text.strip()
        for paragraph in document.paragraphs
        if paragraph.style and paragraph.style.name.startswith("Heading")
    }
    assert set(EXPECTED_TOC) <= headings
    for removed in (
        "1.6 Report Structure",
        "Chapter 5: Scorer Validation and Robustness",
        "Chapter 6: Discussion",
        "Chapter 7: Conclusion and Future Work",
        "2.7 Implications for Experimental Design",
        "3.7 Reproducibility Controls",
        "3.8 Analysis Workflow and Quality Control",
        "Appendix A: Reproducibility Summary",
    ):
        assert removed not in headings


def test_user_introduction_and_four_outcomes_are_present(document: Document) -> None:
    text = all_text(document)
    assert "The problem addressed by this study is interpretive rather than merely operational." in text
    assert (
        "four outcomes under a single controlled precision contrast: harmful compliance, benign over-refusal, "
        "MMLU accuracy, and ARC-Challenge accuracy"
    ) in text
    assert "complicance" not in text
    assert "together thus" not in text


def test_progress_is_repository_grounded_and_remaining_scope_is_narrow(document: Document) -> None:
    text = all_text(document)
    assert len(list(ROOT.glob("results_512/*/*/summary.json"))) == 60
    assert "no harmful-compliance contrast survives Benjamini–Hochberg" in text
    assert "The corrected family contains three survivors" in text
    assert "classifier reaches Cohen's κ = 0.59, compared with κ = 0.11 for the refusal regex" in text
    assert "Scorer Validation and Robustness" in text
    assert "Discussion" in text
    for excluded in (
        "supervisor feedback",
        "submission preparation",
        "formatting and submission",
        "second annotator",
        "MyTCinfo",
    ):
        assert excluded not in text


def test_surviving_citations_keep_the_correct_source_identities(document: Document) -> None:
    text = all_text(document)
    assert "GPTQ / AWQ / SmoothQuant [2], [1], [22]" in text
    assert "GPTQ / AWQ / SmoothQuant [5], [6], [22]" not in text
    assert (
        "HarmBench and XSTest scoring are cross-checked using the HarmBench classifier, an independent API "
        "judge, Llama Guard, and benchmark-specific human-labelled samples."
    ) in text
    assert "and two human-labelled samples" not in text


def test_toc_pages_and_internal_links(document: Document) -> None:
    text = all_text(document)
    for title, page in EXPECTED_TOC.items():
        assert f"{title}\t{page}" in text

    bookmarks = {
        element.get(qn("w:name"))
        for element in document.element.body.iter(qn("w:bookmarkStart"))
    }
    anchors = [
        element.get(qn("w:anchor"))
        for element in document.element.body.iter(qn("w:hyperlink"))
        if element.get(qn("w:anchor"))
    ]
    assert len(anchors) == len(EXPECTED_TOC)
    assert len(set(anchors)) == len(anchors)
    assert set(anchors) <= bookmarks


def test_static_toc_and_reference_count(document: Document) -> None:
    references = [p.text for p in document.paragraphs if p.text.startswith("[")]
    assert len(references) == 57
    for removed in (
        "A survey on LLM-as-a-judge",
        "No free labels: Limitations of LLM-as-a-judge without human grounding",
        "Individual comparisons by ranking methods",
        "Statistical Power Analysis for the Behavioral Sciences",
    ):
        assert removed not in "\n".join(references)
    with ZipFile(DOCX) as package:
        document_xml = package.read("word/document.xml")
        settings_xml = package.read("word/settings.xml")
    assert b'TOC \\o "1-2"' not in document_xml
    assert b"updateFields" not in settings_xml
