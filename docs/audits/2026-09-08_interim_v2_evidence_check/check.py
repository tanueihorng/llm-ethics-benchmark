from pathlib import Path
import json,re,math,zipfile,hashlib
import numpy as np
from docx import Document
from lxml import etree as E
R=Path(__file__).resolve().parents[3];D=Document(R/'docs/rewritten_fyp_interim_progress_v2.docx')
pairs=['qwen_2b','qwen_4b','llama_3_2_3b','mistral_7b','phi4_mini'];out=[]
for pair in pairs:
 for bench,field,src in [('harmbench','judge_harmful','scores.judge.harmbench_cls.jsonl'),('xstest','is_refusal','scores.v2.jsonl'),('mmlu','is_correct','raw.jsonl'),('arc','is_correct','raw.jsonl')]:
  maps=[]
  for suffix in ['base','4bit']:
   rows=[json.loads(l) for l in (R/f'results_512/{pair}_{suffix}/{bench}/{src}').read_text().splitlines() if l.strip()]
   vals={str(r['prompt_id']):int(r[field] if bench=='harmbench' else r['score_fields'][field]) for r in rows}
   assert len(vals)==len(rows);maps.append(vals)
  assert maps[0].keys()==maps[1].keys();ids=sorted(maps[0]);a=np.array([maps[0][i] for i in ids]);z=np.array([maps[1][i] for i in ids]);n=len(a);b=int(((a==0)&(z==1)).sum());c=int(((a==1)&(z==0)).sum());k=b+c
  p=min(1,2*sum(math.comb(k,i) for i in range(min(b,c)+1))/2**k) if k else 1
  draws=(z-a)[np.random.default_rng(42).integers(0,n,size=(10000 if bench=='harmbench' else 2000,n))].mean(axis=1);lo,hi=np.quantile(draws,[.025,.975]);out.append(dict(pair=pair,benchmark=bench,n=n,base=float(a.mean()),nf4=float(z.mean()),delta=(b-c)/n,b=b,c=c,p=p,lo=float(lo),hi=float(hi)))
order=sorted(range(20),key=lambda i:out[i]['p']);q=1
for rank,i in reversed(list(enumerate(order,1))):q=min(q,out[i]['p']*20/rank);out[i]['q']=q
errors=[]
for i,pair in enumerate(pairs):
 for bench,ti,off,dp in [('harmbench',4,1,3),('xstest',4,4,3),('mmlu',5,1,3),('arc',5,4,4)]:
  v=next(v for v in out if v['pair']==pair and v['benchmark']==bench);row=D.tables[ti].rows[i+1]
  got=[float(s.replace('−','-')) for j in range(off,off+3) for s in re.findall(r'[+−-]?\d+\.\d+',row.cells[j].text)]
  want=[round(v['base'],3),round(v['nf4'],3),round(v['delta'],dp),round(v['lo'],dp),round(v['hi'],dp),round(v['p'],4),round(v['q'],4)]
  if got!=want:errors.append(dict(pair=pair,benchmark=bench,reported=got,recomputed=want))
refs=[p.text for p in D.paragraphs if p.style.name=='Reference'];assert len(refs)==61
assert [int(re.match(r'\[(\d+)\]',s)[1]) for s in refs]==list(range(1,62))
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile(R/'docs/rewritten_fyp_interim_progress_v2.docx') as z:
 root=E.fromstring(z.read('word/document.xml'));texts=[''.join(p.xpath('.//w:t/text()',namespaces=ns)) for p in root.xpath('//w:p',namespaces=ns)]
 body=[s for s in texts if s not in refs];seq=[]
 for s in body:
  expanded=re.sub(r'\[(\d+)\]\s*[–-]\s*\[(\d+)\]',lambda m:' '.join(f'[{n}]' for n in range(int(m[1]),int(m[2])+1)),s)
  seq.extend(map(int,re.findall(r'\[(\d+)\]',expanded)))
 unique=list(dict.fromkeys(seq));assert unique==list(range(1,62)),unique
 extra=[]
 for name in z.namelist():
  if name.endswith('.xml') and any(s in name for s in ['footnote','endnote','header','footer']):
   t=''.join(E.fromstring(z.read(name)).xpath('//w:t/text()',namespaces=ns))
   if re.search(r'\[\d+\]',t):extra.append(name)
 result=dict(document_sha256=hashlib.sha256((R/'docs/rewritten_fyp_interim_progress_v2.docx').read_bytes()).hexdigest(),references=len(refs),citation_occurrences_expanded=len(seq),all_references_cited=True,first_use_in_order=True,extra_citation_parts=extra,primary_comparisons=out,table_discrepancies=errors)
Path(__file__).with_name('numeric-citation-check.json').write_text(json.dumps(result,indent=2));print(json.dumps({k:v for k,v in result.items() if k!='primary_comparisons'},indent=2));print('Survivors',[(v['pair'],v['benchmark'],v['q']) for v in out if v['q']<.05])

assert not errors, errors
