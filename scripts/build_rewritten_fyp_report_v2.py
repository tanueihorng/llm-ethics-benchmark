#!/usr/bin/env python3
"""Build the plain-academic-English v2 rewrite without modifying v1.

The v1 builder remains the evidence and layout source. This builder renders it
to a temporary DOCX, rewrites the prose layer, verifies that numerical tokens
and citation markers are unchanged, and writes a separate v2 artifact.
"""

from __future__ import annotations

import importlib.util
import re
import tempfile
from pathlib import Path

from docx import Document


ROOT = Path(__file__).resolve().parents[1]
BASE_BUILDER = ROOT / "scripts" / "build_rewritten_fyp_report_v1.py"
OUT = ROOT / "docs" / "rewritten_fyp_report_v2.docx"


HEADING_REWRITES = {
    "2.5 Automatic Evaluators and Human Grounding":
        "2.5 Automatic Evaluators and Human Validation",
    "5.3 Human Grounding of Harmful-Compliance Scorers":
        "5.3 Human Validation of Harmful-Compliance Scorers",
    "6.3 Capability as the Deployment Cost":
        "6.3 Capability Effects and Deployment",
    "6.4 Scorer Validity as a Research Result":
        "6.4 Validity of the Scoring Methods",
    "Figure 1.1. Capability-anchored safety space using MMLU and HarmBench deltas":
        "Figure 1.1. Changes in MMLU accuracy and HarmBench harmful compliance",
    "Table 6.1. Competing explanations for the HarmBench null":
        "Table 6.1. Possible explanations for the HarmBench result",
    "Table 6.2. Published quantization effects placed beside this study's":
        "Table 6.2. Comparison with published quantization results",
}

# Cached front-matter pagination from the final 59-page LibreOffice render.
# Word refreshes the TOC field when opened, but the visible cached TOC and the
# lists of figures/tables must still be correct before that refresh occurs.
PAGE_NUMBER_REWRITES = {
    "1.3 Research Questions": 2,
    "1.5 Contributions": 3,
    "1.6 Report Structure": 3,
    "Chapter 3: Methodology and Experimental Setup": 14,
    "3.1 Matched-Pair Experimental Design": 14,
    "3.2 Models and Precision Conditions": 15,
    "3.3 Benchmarks and Outcomes": 16,
    "3.4 Scoring and Validation": 17,
    "3.5 Statistical Analysis": 17,
    "3.6 Implementation and Execution": 19,
    "3.7 Reproducibility Controls": 20,
    "3.8 Analysis Workflow and Quality Control": 20,
    "Chapter 4: Results": 22,
    "4.1 Results at a Glance": 22,
    "4.2 Harmful Compliance (RQ1)": 23,
    "4.3 Over-Refusal on Benign Prompts (RQ2)": 24,
    "4.4 Capability (RQ3)": 25,
    "4.5 Model Scale and Family Differences (RQ4)": 26,
    "4.6 INT8 Precision Point": 26,
    "4.7 Research-Question Synthesis": 28,
    "Chapter 5: Scorer Validation and Robustness": 29,
    "5.1 HarmBench Classifier versus Refusal Proxy": 29,
    "5.2 Independent HarmBench Judges": 30,
    "5.3 Human Validation of Harmful-Compliance Scorers": 31,
    "5.4 XSTest Scorer Sensitivity and Human Validation": 31,
    "5.5 Generation-Length Robustness": 32,
    "5.6 Multi-Seed Decoding Sensitivity": 33,
    "5.7 Power and the Meaning of the Null": 34,
    "5.8 Robustness Synthesis": 35,
    "Chapter 6: Discussion": 36,
    "6.1 Answer to the Central Question": 36,
    "6.2 Relationship to Prior Work": 36,
    "6.3 Capability Effects and Deployment": 39,
    "6.4 Validity of the Scoring Methods": 40,
    "6.5 Threats to Validity": 41,
    "6.6 Practical Implications": 42,
    "Chapter 7: Conclusion and Future Work": 43,
    "7.1 Conclusions": 43,
    "7.2 Prioritised Future Work": 43,
    "References": 45,
    "Appendix A: Reproducibility Summary": 48,
    "Appendix B: Minimum Detectable Capability Effects": 49,
    "Appendix C: Human-Validation Boundaries": 50,
    "Figure 1.1. Changes in MMLU accuracy and HarmBench harmful compliance": 2,
    "Figure 3.1. Matched pairs, precision arms, benchmark battery, and analysis flow": 15,
    "Figure 4.1. Per-pair HarmBench classifier ΔASR with paired confidence intervals": 24,
    "Figure 4.2. Precision sweep across fp16, INT8, and NF4": 28,
    "Figure 5.1. HarmBench classifier ASR versus the refusal-regex proxy": 30,
    "Figure 5.2. Qwen3-1.7B greedy result and five-seed HarmBench sensitivity": 34,
    "Table 3.1. Model pairs and precision conditions": 15,
    "Table 3.2. Benchmark roles, sample sizes, and primary outcomes": 16,
    "Table 4.1. NF4 safety and over-refusal results": 22,
    "Table 4.2. NF4 capability results": 22,
    "Table 5.1. Per-pair HarmBench deltas under three judges": 31,
    "Table 5.2. HarmBench classifier deltas at 128 and 512 tokens": 33,
    "Table 5.3. Five-seed HarmBench sensitivity": 33,
    "Table 6.1. Possible explanations for the HarmBench result": 36,
    "Table 6.2. Comparison with published quantization results": 37,
    "Table A.1. Primary run controls": 48,
    "Table B.1. Capability effects and minimum detectable effects": 49,
}


# Full-paragraph rewrites are used where the original argument was compressed
# into project-internal terminology. Prefixes are intentionally long and must
# each match exactly one paragraph in the freshly built v1 document.
PARAGRAPH_REWRITES = [
    (
        "Post-training quantization makes compact instruction-tuned language models practical",
        "Post-training quantization reduces the resources needed to run instruction-tuned language "
        "models, but it may also change their behaviour. This report examines whether four-bit "
        "quantization affects harmful compliance or refusal of benign prompts, and whether apparent "
        "safety changes can instead be explained by lower general capability. Five matched model "
        "pairs from Qwen, Llama, Mistral, and Phi were evaluated at fp16, INT8, and NF4 precision. "
        "HarmBench measured harmful compliance, XSTest measured refusal on benign prompts, and MMLU "
        "and ARC-Challenge measured general capability. The analysis used paired bootstrap intervals, "
        "exact McNemar tests, and Benjamini-Hochberg correction. The HarmBench classifier was checked "
        "against an API judge, Llama Guard, and a human-labelled sample. At the HarmBench reference "
        "limit of 512 new tokens, no NF4 change in harmful compliance remained significant after "
        "multiple-comparison correction. Three results remained significant: lower MMLU accuracy for "
        "Qwen3-1.7B, lower ARC accuracy for Llama-3.2-3B, and lower regex-scored over-refusal for "
        "Phi-4-mini. Human validation found that the semantic refusal judge agreed much more closely "
        "with the annotator than the regex did, so the Phi result depends on the scoring method. Within "
        "the direct-request setting evaluated here, NF4 caused detectable capability losses but did not "
        "produce a statistically significant increase in harmful compliance after correction.",
    ),
    (
        "Compact instruction-tuned language models have moved capable generation into a deployment range",
        "Smaller instruction-tuned language models can now run on consumer GPUs and some edge devices. "
        "Qwen3 [1], the Llama 3 family [2], Phi-4-mini [3], and Gemma 2 [4] are examples of models with "
        "far fewer parameters than frontier systems. Local inference can keep user inputs on the device "
        "when no remote service is used. This report does not measure latency, energy use, privacy, or "
        "monetary cost, and therefore makes no claim about benefits on those dimensions. It focuses on "
        "model behaviour. Memory remains a practical constraint because full-precision inference may "
        "exceed the capacity of the target device [5].",
    ),
    (
        "Post-training quantization addresses that constraint by representing weights",
        "Post-training quantization reduces memory use by representing weights, and sometimes "
        "activations, with fewer bits. AWQ reports four-bit compression and faster inference on desktop "
        "and mobile GPUs [5]. GPTQ demonstrates one-shot weight quantization for models with as many as "
        "175 billion parameters [6]. A large scaling-law study also finds that four-bit inference often "
        "offers a favourable balance under a fixed memory budget [7]. These studies show that four-bit "
        "quantization is a realistic deployment configuration rather than only a laboratory experiment.",
    ),
    (
        "The quantization method papers reviewed here report perplexity, accuracy, latency, or memory",
        "Most quantization studies report perplexity, task accuracy, latency, or memory use [5]-[7]. "
        "Large comparisons show that many tasks remain usable at four bits [8], [9], but a small average "
        "accuracy change does not show that safety behaviour is unchanged. Safety-focused studies report "
        "different effects across models, bit widths, and quantization methods [10]. Quantization may also "
        "affect model confidence when accuracy remains stable [11]. Other studies have identified "
        "quantization-related vulnerabilities [12] and language-dependent performance losses that are "
        "understated by automatic metrics [13].",
    ),
    (
        "Refusal behaviour gives a concrete reason for concern",
        "Refusal behaviour is one reason to test quantized models directly. Training an assistant to be "
        "helpful and harmless teaches it when to refuse a request [14]. Mechanistic studies associate "
        "refusal with a concentrated direction in activation space [15], while other work suggests that "
        "safety alignment can be weak during generation [16]. Fine-tuning on a small adversarial dataset "
        "can also weaken safeguards [17]. These findings do not show that ordinary NF4 quantization harms "
        "safety. They do show why behaviour should be measured again after the model weights are changed.",
    ),
    (
        "Safety evaluation is difficult because refusal is not synonymous with safety",
        "Refusal alone is not a sufficient measure of safety. A model that refuses every request may "
        "appear harmless but provide little practical value. XSTest was designed to detect excessive "
        "refusal of benign prompts [18], and OR-Bench later studied the same problem at a larger scale "
        "[19]. Conversely, a response without an explicit refusal phrase is not necessarily harmful. A "
        "safety score may change because the model's behaviour changed, because its general capability "
        "changed, or because the scorer interpreted the response differently.",
    ),
    (
        "This project asks whether quantization changes the balance among harmful compliance",
        "This project examines how quantization affects harmful compliance, refusal of benign prompts, "
        "and general capability. The central distinction is between a change in alignment and a loss of "
        "capability. A lower attack-success rate does not show better alignment if the model has also "
        "become less able to answer ordinary questions. The study therefore evaluates all three outcomes "
        "and validates the scorers used to measure them.",
    ),
    (
        "Figure 1.1 previews the answer",
        "Figure 1.1 compares each model pair's change in MMLU accuracy with its change in "
        "classifier-scored harmful compliance. Paired bootstrap intervals are shown for both measures. "
        "None of the five intervals provides evidence of an increase in harmful compliance. The Llama "
        "interval excludes zero only for a decrease in harmful compliance, while all five MMLU point "
        "estimates are negative. Chapter 4 reports which changes remain significant after correction, "
        "and Chapter 5 examines whether the conclusions change when different scorers are used.",
    ),
    (
        "The capability instruments are a six-subject MMLU subset and ARC-Challenge",
        "General capability is measured with a six-subject MMLU subset and ARC-Challenge. These two "
        "multiple-choice benchmarks do not cover every form of model capability. The primary HarmBench "
        "generation limit is 512 new tokens and the main decoding condition is greedy. A separate "
        "five-seed analysis tests sensitivity to stochastic decoding and is not included in the primary "
        "set of statistical comparisons.",
    ),
    (
        "A reproducible benchmark implementation with redacted derived sidecars",
        "A reproducible benchmark implementation that preserves the original run outputs and stores "
        "revised scorer labels in separate redacted files.",
    ),
    (
        "The literature also sets the report's claim boundary",
        "The literature also defines the scope of this study. Quantization-conditioned attacks show that "
        "worst-case failures are possible [12], [39], while jailbreak benchmarks show that adversarial "
        "robustness requires an explicit attack evaluation [43]. Research on automatic judges shows why "
        "their labels should be checked against human annotations [45], [46]. The conclusion of this "
        "report is therefore limited to the tested models, quantization methods, prompts, generation "
        "limit, and scorers: capability losses were detected, but no increase in harmful compliance "
        "remained significant after correction.",
    ),
    (
        "The primary multiplicity family contains 20 NF4 contrasts",
        "The primary analysis contains 20 NF4 comparisons. Three results remain significant after "
        "Benjamini-Hochberg correction, and none concerns HarmBench harmful compliance. They are the "
        "Qwen3-1.7B MMLU decrease of −0.090, the Llama-3.2-3B ARC decrease of −0.0316, and the "
        "Phi-4-mini over-refusal decrease of −0.048.",
    ),
    (
        "The correction therefore moves the report away from a pair-by-pair search",
        "After correction, the clearest effects concern capability. The only significant over-refusal "
        "result is a decrease, and Chapter 5 shows that it is not reproduced by the semantic scorer.",
    ),
    (
        "Tables 4.1 and 4.2 separate safety from capability",
        "Tables 4.1 and 4.2 report the safety and capability results separately. The baseline and NF4 "
        "columns show the observed rates. The delta and confidence-interval columns show the estimated "
        "change and its uncertainty, while the p- and q-values distinguish individual significance from "
        "significance after correcting the full set of comparisons.",
    ),
    (
        "The sign convention is consistent: positive ASR and over-refusal deltas are adverse",
        "Positive changes in ASR or over-refusal are undesirable, while negative capability changes are "
        "undesirable. A decrease in ASR does not offset a simultaneous loss of capability and should not "
        "by itself be interpreted as improved alignment.",
    ),
    (
        "Capability contains the clearest costs",
        "The largest corrected effects concern capability. Qwen3-1.7B MMLU accuracy decreases from "
        "0.643 to 0.553, a decrease of 0.090 that remains significant after BH-FDR correction with "
        "q = 0.008. In practical terms, the NF4 model answers about nine fewer questions correctly per "
        "hundred than its fp16 counterpart.",
    ),
    (
        "The independent three-class refusal judge does not reproduce the Phi-4-mini regex decrease",
        "The independent three-class refusal judge does not reproduce the Phi-4-mini decrease found by "
        "the regex. The strict definition gives a small positive change, while the broad definition is "
        "close to zero. Neither result is statistically significant. Across all five model pairs, the "
        "semantic judge finds no statistically significant increase in over-refusal.",
    ),
    (
        "Human validation explains the discrepancy",
        "The human-labelled sample helps explain this disagreement. Under the strict definition, kappa "
        "is 0.485 for the semantic judge and −0.006 for the regex. The regex identifies only 3.2% of full "
        "refusals, compared with 96.8% for the judge. Under the broad definition, kappa is 0.662 for the "
        "judge and 0.054 for the regex. These results suggest that the significant Phi-4-mini decrease "
        "found by the regex is more likely to reflect scoring error than a stable behavioural change.",
    ),
    (
        "The XSTest human audit contains a disclosed protocol deviation",
        "The XSTest human audit departed from the preregistered procedure in one respect. The blindness "
        "exclusion was applied to every response associated with an affected prompt, rather than only to "
        "the specified model-response items. This decision was made before sampling and annotation and "
        "reduced possible priming, but it did not follow the written rule. The analysis therefore should "
        "not be described as a fully preregistered confirmation. Repeating the comparison on the common "
        "items shows that no more than 14 of the 200 sampled items were affected and that the large kappa "
        "difference between the judge and regex remains.",
    ),
    (
        "The broad mapping increases the refusal class by counting partial refusals",
        "The broad definition also counts partial refusals. The judge's kappa increases under this "
        "definition, while the regex remains close to zero because it misses most refusals that are not "
        "expressed through its fixed phrases. The semantic judge therefore agrees more closely with the "
        "annotator under both definitions.",
    ),
    (
        "The validation-informed multiplicity family replaces only the over-refusal instrument",
        "A separate sensitivity analysis replaces the regex over-refusal score with the strict semantic "
        "judge score while leaving the five model pairs and all other outcomes unchanged. Only the Qwen "
        "MMLU and Llama ARC capability losses remain significant. No harmful-compliance or over-refusal "
        "comparison remains significant after correction.",
    ),
    (
        "Failure to reject is not proof that the true effect is zero",
        "A non-significant result does not prove that the true effect is exactly zero. With 200 paired "
        "HarmBench items per model, the minimum detectable absolute change in ASR is approximately "
        "0.04–0.09, depending on the number of discordant pairs. The experiment can therefore detect a "
        "large and consistent increase across models, but smaller effects may remain unresolved.",
    ),
    (
        "The null is nevertheless more informative than a single non-significant p-value",
        "Several checks support the same qualified conclusion. No ASR comparison remains significant "
        "after correction. Three semantic judges find no significant increase in harmful compliance for "
        "any model pair. The Qwen3-1.7B increase observed with a 128-token limit is absent at 512 tokens, "
        "and the multi-seed analysis finds no stable increase. These checks address multiple comparisons, "
        "scorer validity, response truncation, and variation across decoding seeds.",
    ),
    (
        "Capability nulls have different resolution because MMLU and ARC sample sizes differ",
        "The capability tests have different sensitivity because MMLU and ARC use different sample "
        "sizes. Appendix B shows that some observed changes are smaller than their estimated minimum "
        "detectable effects. The report describes these cases as having no detectable change and uses "
        "equivalence language only for the INT8 comparisons that were tested against an explicit margin.",
    ),
    (
        "The defensible claim is the absence of a corrected increase large enough",
        "The evidence shows no statistically significant increase in harmful compliance after correction "
        "at the resolution of this experiment. It does not prove exact preservation or rule out smaller "
        "effects.",
    ),
    (
        "RQ5 asks whether the conclusion survives measurement and inference changes",
        "RQ5 tests whether the conclusion changes under different scorers, generation limits, and "
        "decoding seeds. Neither the API judge nor Llama Guard finds a significant NF4 increase in "
        "harmful compliance for any model pair. Human annotations favour the HarmBench classifier over "
        "the refusal proxy and the semantic XSTest judge over the regex. The Qwen3-1.7B increase observed "
        "at 128 tokens is absent at 512 tokens, and the multi-seed analysis finds no stable increase. "
        "These checks support the same qualified conclusion through different forms of evidence.",
    ),
    (
        "The two anchors prevent overgeneralisation from one benchmark",
        "The two capability benchmarks prevent overgeneralisation from a single task. Qwen3-1.7B has a "
        "large MMLU loss but only a small ARC decrease, whereas Llama's reduction is clearer on ARC than "
        "on MMLU. The largest corrected effects therefore concern capability, but their size depends on "
        "the benchmark.",
    ),
    (
        "The sample is stratified toward scorer disagreements",
        "The sample contains a larger share of scorer disagreements than the full HarmBench dataset, so "
        "its class proportions should not be interpreted as population ASR. It also has only one "
        "annotator. The appropriate comparison is therefore between scorers within this audit sample: "
        "the classifier agrees more closely with the human labels than the refusal proxy does. A second "
        "annotator would allow inter-annotator agreement to be measured.",
    ),
    (
        "The comparison also changes absolute baseline rates because longer responses",
        "The longer generation limit also changes the absolute baseline rates because models have more "
        "opportunity to comply, refuse, or qualify an answer. The 512-token results are therefore used "
        "for the primary analysis, while the 128-token results are retained only as a sensitivity "
        "comparison. When truncation is common, the generation limit can change what the benchmark "
        "effectively measures.",
    ),
    (
        "Mistral and Phi remain close to zero on average",
        "The average Mistral and Phi changes remain close to zero. Their greedy estimates lie just "
        "outside narrow seed ranges that include or approach zero, which shows that greedy and stochastic "
        "decoding are different experimental conditions. Across all five model pairs, the multi-seed "
        "analysis finds no stable statistically significant increase in harmful compliance.",
    ),
    (
        "Figure 5.2 shows the Qwen3-1.7B greedy headline",
        "Figure 5.2 compares the Qwen3-1.7B greedy estimate with the five stochastic-decoding seeds. The "
        "greedy estimate matches the two lowest seed estimates and falls within one standard deviation "
        "of the seed mean. It is therefore consistent with the variation observed across decoding seeds.",
    ),
    (
        "A seed is not an independent model replicate",
        "Changing the decoding seed does not create an independent model replicate. The model weights "
        "and prompts remain fixed, so this analysis measures variation caused by stochastic generation "
        "rather than uncertainty across separately trained models.",
    ),
    (
        "This pattern is more valuable than agreement for its own sake",
        "Disagreement between scorers is examined rather than averaged away. When two scorers disagree, "
        "the report identifies the difference in what they measure and gives more weight to the scorer "
        "that agrees more closely with human annotations. When they agree, the conclusion remains limited "
        "to the scope and statistical resolution of the experiment.",
    ),
    (
        "The validation results can be read as a sequence of eliminated explanations",
        "The validation analyses address several alternative explanations. Human labels show that the "
        "refusal proxy over-flags harmful compliance. The API judge and Llama Guard test whether the "
        "classifier result depends on one judge. The 512-token rerun addresses response truncation, and "
        "the five stochastic seeds test sensitivity to decoding randomness.",
    ),
    (
        "Scorer validation is not an auxiliary engineering check",
        "Scorer validation changes the interpretation of the experiment. The HarmBench refusal proxy "
        "overestimates harmful compliance, while the XSTest regex misses most full refusals in the human "
        "sample. For both benchmarks, the semantic scorer agrees more closely with the human annotations.",
    ),
    (
        "For the XSTest judge sensitivity study, its dated pre-registration",
        "The original preregistered XSTest analysis and the later validation analysis are both reported. "
        "Retaining the regex result prevents the scorer from being replaced simply because of the "
        "outcome. Reporting the semantic-judge result incorporates the stronger validation evidence. "
        "Both analyses reach the same conclusion relevant to RQ2: neither scorer finds a statistically "
        "significant increase in over-refusal.",
    ),
    (
        "The study began with a distinction: an apparent safety change may reflect alignment",
        "The results show why harmful compliance, general capability, and scorer behaviour must be "
        "considered together. With the primary HarmBench classifier and a 512-token generation limit, no "
        "NF4 increase in harmful compliance remains significant after correction. Two of the three "
        "significant results are capability losses. The third is a decrease in regex-scored over-refusal "
        "that the semantic scorer does not reproduce.",
    ),
    (
        "This conclusion is deliberately bounded",
        "This conclusion applies only to the evaluated models, quantization settings, direct-request "
        "prompts, generation limits, and detectable effect sizes. It does not show that NF4 preserves all "
        "forms of safety or that every smaller model will behave in the same way. Within the experiment, "
        "the clearest costs of NF4 were losses in general capability rather than increases in harmful "
        "compliance.",
    ),
    (
        "The evidence also rejects two tempting simplifications",
        "The results do not support a general claim that quantization improves safety. Changes differ "
        "across model pairs, and the only individually significant decrease in harmful compliance occurs "
        "alongside a capability loss. The results also do not support a general claim that NF4 increases "
        "harmful compliance, because none of those increases remains significant after correction or is "
        "consistently reproduced across judges.",
    ),
    (
        "A more useful description is model-specific boundary instability under a capability cost",
        "Quantization changes some prompt-level classifications in both directions even when the average "
        "change is close to zero. One possible explanation is that small numerical changes move borderline "
        "responses across the scorer's decision boundary. The behavioural results cannot identify the "
        "internal mechanism, so this remains a hypothesis rather than a causal conclusion.",
    ),
    (
        "The practical message is procedural rather than promotional",
        "The results do not justify describing NF4 as either unsafe or risk-free. A quantized model should "
        "be evaluated in its final deployment configuration, using task-relevant capability tests and a "
        "scorer that measures the intended safety outcome.",
    ),
    (
        "Read together, the table supports the narrower reading this report has taken throughout",
        "Taken together, the prior studies support a conclusion limited to the evaluated models and "
        "settings. Their reported increases in harmful compliance arise under attacks, calibrated offline "
        "quantization methods that rewrite the checkpoint, or models deliberately prepared to fail after "
        "quantization. They do not show that ordinary NF4 loading increases harmful compliance in an "
        "unmodified instruction-tuned checkpoint. This experiment tested that narrower question and found "
        "no detectable increase at its statistical resolution.",
    ),
    (
        "The 512-token INT8 rerun also resolves a historical warning",
        "The 512-token INT8 rerun also clarifies an earlier result. At 128 tokens, Llama INT8 appeared to "
        "increase harmful compliance under both judges. At 512 tokens, the classifier change is +0.005 "
        "with p = 1.000 and the API-judge change is +0.010 with p = 0.688. The apparent increase is "
        "therefore not present at the reference generation limit.",
    ),
    (
        "This project evaluated five matched compact-language-model pairs",
        "This project evaluated five matched pairs of smaller language models on HarmBench, XSTest, MMLU, "
        "and ARC-Challenge. NF4 was the primary quantization method, and INT8 provided an additional "
        "precision comparison. The experiment evaluated each quantized model directly instead of assuming "
        "that it retained the behaviour of its full-precision counterpart.",
    ),
    (
        "RQ1 is answered by a bounded null",
        "For RQ1, no NF4 change in harmful compliance remains significant after correcting the 20 primary "
        "comparisons at the 512-token reference limit. For RQ2, no scorer finds a statistically significant "
        "increase in refusal of benign prompts. For RQ3, the largest corrected losses concern capability, "
        "particularly Qwen3-1.7B on MMLU and Llama-3.2-3B on ARC. RQ4 shows that the size of the change "
        "depends on the model, family, task, and precision condition; capability generally remains closer "
        "to fp16 under INT8 than under NF4. For RQ5, the main harmful-compliance conclusion is unchanged "
        "when alternative judges, human annotations, a longer generation limit, and multiple decoding "
        "seeds are considered.",
    ),
    (
        "The central lesson is methodological",
        "A refusal count alone is not enough to interpret the safety effect of quantization. Harmful "
        "compliance, refusal of benign prompts, and general capability should be measured together, and "
        "the scorer should be checked against the behaviour it is intended to identify.",
    ),
    (
        "This report also changes how the contribution is presented",
        "The main contribution is the evidence produced by the controlled experiment and the validation "
        "of the scoring methods. The benchmarking framework, cluster scripts, and verification tools made "
        "the study possible, but they are supporting implementation rather than the principal research "
        "result.",
    ),
    (
        "Within its boundaries, the study provides a clear deployment conclusion",
        "Under the tested conditions, INT8 preserves capability more consistently than NF4, while NF4 can "
        "cause substantial losses on particular tasks. Neither precision condition produces a "
        "statistically significant increase in harmful compliance after correction in the 512-token "
        "DirectRequest setting. Deployment decisions should therefore include task-specific capability "
        "tests as well as safety evaluation of the final quantized model.",
    ),
    (
        "These extensions should preserve the discipline used here",
        "Future experiments should preserve the original run outputs, store revised scorer labels in "
        "separate redacted files, use matched checkpoint revisions, and distinguish the primary analysis "
        "from sensitivity analyses.",
    ),
    (
        "The first two extensions are the most urgent",
        "The first two extensions are the most urgent. An adversarial evaluation would extend the study "
        "beyond direct requests, while a second annotator would measure the reliability of both "
        "human-labelled datasets. Measurements of efficiency and additional quantization methods would "
        "then allow safety and capability changes to be compared with resource savings.",
    ),
]


# These replacements cover the same register throughout the remaining chapters.
# Specific grammatical forms appear before general terms.
PHRASE_REWRITES = [
    ("Capability anchors", "Benchmarks of general capability"),
    ("Capability collapse masquerading as safety", "Capability loss with lower harmful compliance"),
    ("Capability anchor", "Capability benchmark"),
    ("clean pre-registered badge", "fully preregistered status"),
    ("capability collapse masquerading as safety", "capability loss accompanied by lower harmful compliance"),
    ("capability cliff", "large capability decrease"),
    ("current-facing surfaces", "current report documents"),
    ("public result surface", "public repository"),
    ("result surface", "reported results"),
    ("claim boundary", "scope of the conclusions"),
    ("precision-loading path", "precision configuration"),
    ("deployment form", "deployment configuration"),
    ("neighbouring explanations", "alternative explanations"),
    ("procedural loyalty", "strict adherence to the original procedure"),
    ("human-grounded", "validated against human annotations"),
    ("Human grounding", "Comparison with human annotations"),
    ("human grounding", "comparison with human annotations"),
    ("better-grounded instrument", "scorer that agreed more closely with the human annotations"),
    ("better-grounded semantic scorer", "semantic scorer that agreed more closely with the human annotations"),
    ("scorer ranking", "comparison between scorers"),
    ("capability anchors", "benchmarks of general capability"),
    ("capability anchor", "capability benchmark"),
    ("capability anchoring", "using capability results to interpret safety results"),
    ("primary multiplicity family", "set of 20 primary NF4 comparisons"),
    ("validation-informed multiplicity family", "validation-informed sensitivity analysis"),
    ("multiplicity family", "set of corrected comparisons"),
    ("corrected family", "results after multiple-comparison correction"),
    (
        "measurement error from the three surviving effects, so scorer validation remains necessary after statistical correction",
        "measurement error. Exactly three primary contrasts survived Benjamini-Hochberg false-discovery-rate correction, so scorer validation remains necessary after statistical correction",
    ),
    (
        "the validation-informed family replaces it with judge-strict over-refusal and leaves only two survivors, both capability losses",
        "the validation-informed sensitivity analysis replaces that score with judge-strict over-refusal. In that validation-informed sensitivity analysis, exactly two contrasts survived correction, both capability losses",
    ),
    ("registered survivors", "results that remained significant after correction"),
    ("three survivors", "three results that remained significant after correction"),
    ("two survivors", "two results that remained significant after correction"),
    ("single over-refusal survivor", "single over-refusal result that remained significant after correction"),
    ("Phi survivor", "Phi result"),
    ("survivors", "results that remained significant after correction"),
    ("survivor", "result that remained significant after correction"),
    ("harmful-ward sign convention", "sign convention for harmful compliance"),
    ("harmful-ward significance conclusion", "conclusion about the significance of the harmful-compliance change"),
    ("harmful-ward point estimate", "point estimate for increased harmful compliance"),
    ("significant harmful-ward seed", "seed with a significant increase in harmful compliance"),
    ("harmful-ward increase", "increase in harmful compliance"),
    ("harmful-ward shift", "shift towards greater harmful compliance"),
    ("harmful-ward effect", "increase in harmful compliance"),
    ("harmful-ward pair", "model pair with increased harmful compliance"),
    ("harmful-ward estimates", "harmful-compliance estimates"),
    ("harmful-ward", "towards greater harmful compliance"),
    ("safe-ward ASR movement", "decrease in ASR"),
    ("safe-ward decrease", "decrease in harmful compliance"),
    ("safe-ward", "towards lower harmful compliance"),
    ("bounded null", "qualified null result"),
    ("load-bearing assumptions", "central assumptions"),
    ("load-bearing distinction", "central distinction"),
    ("load-bearing", "central"),
    ("the headline", "the main conclusion"),
    ("headline", "main conclusion"),
    ("retired signal", "earlier result"),
    ("retired Qwen signal", "earlier Qwen result"),
    ("signal dissolves", "result is no longer present"),
    ("dissolves", "is no longer present"),
    ("retired", "earlier"),
    ("primary tree", "primary set of results"),
    ("512-token tree", "512-token results"),
    ("results tree", "set of results"),
    ("result trees", "stored results"),
    ("derived sidecars", "derived scoring files"),
    ("sidecars", "derived scoring files"),
    ("sidecar", "derived scoring file"),
    ("registry supplies recurring values", "analysis files provide recurring values"),
    ("claim registry", "verified analysis record"),
    ("registry", "analysis record"),
    ("compare instruments", "compare scorers"),
    ("semantic instrument", "semantic scorer"),
    ("automatic labels require grounding", "automatic labels require validation"),
    ("inconsistent surfaces", "inconsistent documents"),
    ("generated surfaces", "generated documents"),
    ("surface patterns", "text patterns"),
    ("clean badge", "claim of full preregistration"),
    ("comparisons are therefore grounded", "comparisons are therefore supported by human annotations"),
    ("grounded against human labels", "compared with human labels"),
    ("better grounded than", "more consistent with the human labels than"),
    ("committed summaries", "stored summaries"),
    ("Committed summaries", "Stored summaries"),
    ("committed primary configuration", "primary configuration"),
    ("committed analysis", "stored analysis"),
    ("committed redacted scores", "stored redacted scores"),
    ("community artefacts", "community-produced model files"),
    ("deployed artefact", "deployed model"),
    ("attack artefacts", "attack resources"),
    ("full-precision artefacts", "full-precision models"),
    ("deployed artefacts", "deployed models"),
    ("run artefacts", "run files"),
    ("intended artefacts", "intended files"),
    ("earliest artefacts", "earliest run files"),
    ("historical artefacts", "historical run files"),
    ("generation-length artefact", "effect of the shorter generation limit"),
    ("Quantized artefact", "Quantized model"),
    ("INT8 arm", "INT8 comparison"),
    ("multi-seed arm", "multi-seed analysis"),
    ("stochastic arm", "stochastic-decoding analysis"),
    ("adversarial arm", "adversarial evaluation"),
    ("precision point", "precision condition"),
    ("model-specific numerics", "model-specific responses to quantization"),
    ("a condition that an attacker can exploit", "vulnerabilities that attackers may exploit"),
    ("disturb confidence", "affect model confidence"),
    ("derives the research gap", "identifies the research gap"),
    ("A separate literature studies", "A separate body of work examines"),
    ("purchased by inability", "caused by reduced capability"),
    ("arrives with lower capability", "occurs alongside lower capability"),
    ("protect different numerical structures", "preserve different parts of the model's numerical representation"),
    ("the scale of the calibration data", "the amount of calibration data"),
    ("average resilience can conceal subgroup damage", "average performance can hide degradation in particular model groups"),
    ("The parallel for safety is direct", "The same concern applies to safety"),
    ("aggressiveness from method", "the degree of compression from the quantization method"),
    ("model-specific numerics", "differences in how individual models respond numerically"),
    ("not about all low-bit arithmetic", "and cannot be generalised to every low-bit quantization method"),
    ("capability contains the clearest costs", "the clearest effects of NF4 concern capability"),
    ("capability produces the largest corrected effects", "the largest corrected effects concern capability"),
    ("capability axis carries a negative displacement", "MMLU point estimate decreases"),
    ("does not reproduce at the same magnitude", "is much smaller"),
    ("capability is more orderly", "capability changes more consistently with precision"),
    ("historical warning", "earlier result"),
    ("gold-standard claim", "strength of the human reference standard"),
    ("human reliability ceiling", "agreement between independent human annotators"),
    ("counterfactual overlap analysis bounds the displacement", "comparison on the items shared by both sampling rules limits the possible effect"),
    ("generation cap can become part of the measured construct", "generation limit can change what the benchmark effectively measures"),
    ("different exposure across families", "different truncation rates across model families"),
    ("is an accident of one decode", "is specific to one decoding run"),
    ("accident of one decode", "result specific to one decoding run"),
    ("answered in the negative within its detection bounds", "found no detectable increase at the resolution of this experiment"),
    ("English null should not be exported to multilingual safety", "English-language findings should not be generalised to multilingual safety"),
    ("distinct deployed artefact", "distinct deployed model"),
    ("committed validation artefacts", "stored validation files"),
    ("committed validation artifacts", "stored validation files"),
    ("it estimates a precision effect under the DirectRequest condition", "it estimates how precision affects responses to direct requests"),
    ("Breadth increases coverage, but it can make the treatment and scorer less visible", "Broader evaluations cover more behaviours, but they can make the effects of quantization and scoring harder to identify"),
    ("the route from a prompt-level decision to the reported metric can be inspected", "each prompt-level scoring decision can be traced to the reported metric"),
    ("treats scorer validation as a research result rather than a preprocessing detail", "evaluates scorer validity explicitly"),
    ("MMLU and ARC also differ in useful ways", "MMLU and ARC measure different forms of multiple-choice capability"),
    ("Agreement between them would strengthen a broad capability interpretation", "Similar changes on both benchmarks would support a more general conclusion about capability"),
    ("Accuracy can remain stable while other numerical properties move", "Accuracy can remain stable while other numerical properties change"),
    ("borderline behavioural decisions may move", "responses on borderline cases may change"),
    ("The scale of the calibration data also differs", "The amount of calibration data also differs"),
    ("Capability-focused evaluations provide a useful prior expectation", "Capability studies provide context for the expected effects"),
    ("Mechanistic studies of safety layers offer another perspective", "Mechanistic studies examine how refusal is represented inside the model"),
    ("compression could move safety decisions unevenly", "quantization may affect safety decisions differently across prompts and models"),
    ("The safety-specific studies differ in their comparison units", "Safety studies compare different types of model pairs"),
    ("They also differ in outcome construction", "They also use different scoring methods"),
    ("The adversarial studies establish possibility rather than prevalence", "Adversarial studies show that failures can occur, but not how common they are"),
    ("A controlled benign-checkpoint study is needed", "A matched study of normally trained checkpoints is therefore needed"),
    ("Decoding Compressed Trust is the most compatible prior result", "Decoding Compressed Trust is the prior study most similar to this experiment"),
    ("Direct requests and jailbreaks also allocate causal responsibility differently", "Direct-request and jailbreak evaluations fail for different reasons"),
    ("naming its absence of attack augmentation wherever the safety conclusion is stated", "stating explicitly that no attack augmentation was used whenever a safety conclusion is presented"),
    ("XSTest and OR-Bench offer another complementary pair", "XSTest and OR-Bench complement each other"),
    ("it cannot claim the breadth of the larger corpus", "it does not cover the full range of the larger dataset"),
    ("That binary refusal is itself a simplification is shown instead by finer-grained annotation schemes", "Finer-grained annotation schemes show that binary refusal labels are a simplification"),
    ("Partial refusals, caveated answers, and mixed responses occupy the boundary", "Some responses contain partial refusals, caveats, or a mixture of refusal and compliance"),
    ("A pooled coefficient can hide a failure concentrated in the very model that determines the conclusion", "A pooled agreement score can hide poor performance on the model pair that matters most to the conclusion"),
    ("one annotator as an infallible oracle", "one annotator as automatically correct"),
    ("Table 2.1 is a scoped positioning aid", "Table 2.1 compares the studies relevant to this experiment"),
    ("Five gaps follow", "The literature review identifies five gaps"),
    ("a declared family of paired tests", "a predefined set of paired statistical tests"),
    ("It does not attempt to establish a universal law of quantization", "The conclusions are limited to the evaluated models and quantization settings"),
    ("The missing link is inferential", "The main gap concerns interpretation"),
    ("A null result also needs positive design", "A non-significant result is informative only when the experiment can detect effects of practical importance"),
    ("challenge conditions that could hide or manufacture an effect", "conditions that could obscure an effect or create a misleading one"),
    ("make the null more informative", "make the non-significant result easier to interpret"),
    ("The baseline is loaded at the resolved full precision used by the TC1 CUDA environment", "The baseline model uses the full precision selected by the TC1 CUDA environment"),
    ("isolates the load-time precision intervention within a single execution environment", "keeps the execution environment fixed while changing only the loading precision"),
    ("a compute dtype that follows the resolved model dtype", "a compute data type matched to the model data type"),
    ("chat-template path", "chat-template procedure"),
    ("not a post hoc statistical convenience", "specified before the final statistical analysis"),
    ("increase architectural diversity while confounding family, training data, alignment recipe, tokenizer, and size", "broaden the range of architectures, although family, training data, alignment method, tokenizer, and size cannot be separated"),
    ("Each pair shares its model revision exactly", "Both members of each pair use the same model revision"),
    ("Its pre-specified deterministic scorer remains the primary scorer of record", "The preregistered deterministic regex remains the primary scorer"),
    ("a validation-informed sensitivity analysis", "a sensitivity analysis based on the validated scorer"),
    ("thinking disabled where the template exposes that option", "the model's optional thinking mode disabled"),
    ("raw-prompt fallback is not used for the successful primary runs", "all successful primary runs use the model's chat template rather than a raw prompt"),
    ("a stratified 200-item human-labelled sample grounds both", "a stratified 200-item human-labelled sample evaluates both"),
    ("A stratified 200-item human-labelled sample grounds both", "A stratified 200-item human-labelled sample evaluates both"),
    ("a human-labelled sample grounds both", "a human-labelled sample is used to evaluate both"),
    ("This arrangement keeps the original analysis visible", "This preserves the original analysis"),
    ("This difference is stated before agreement is interpreted", "The report accounts for this difference when interpreting agreement"),
    ("documented blindness rule and its disclosed deviation", "documented blindness procedure, including the disclosed deviation"),
    ("rather than treating the verbal bands of Landis and Koch [57] as universal truth", "instead of assuming that the conventional verbal bands of Landis and Koch [57] apply equally to every setting"),
    ("Agreement methodology follows the caution that", "The agreement analysis recognises that"),
    ("Llama Guard supplies the open-weight safety-moderation cross-check", "Llama Guard provides an open-weight comparison for safety classification"),
    ("both remain visible", "both results are reported"),
    (
        "construct error from the three surviving effects",
        "measurement error. Exactly three primary contrasts survived Benjamini-Hochberg false-discovery-rate correction",
    ),
    ("broad degradation denotes adverse capability movement without a clean safety interpretation", "broad degradation describes lower capability without clear evidence of a safety change"),
    ("Matrix orchestration loads one model once", "The matrix runner loads each model once"),
    ("offline Hugging Face modes", "Hugging Face offline settings"),
    ("new user-generated configurations", "configurations created through the dashboard"),
    ("changing the evidence underlying the report", "altering the results used in the report"),
    ("The precision figure reads the fp16–INT8–NF4 analysis", "The precision figure shows the fp16, INT8, and NF4 comparison"),
    ("the precision figure reads the fp16–INT8–NF4 analysis", "the precision figure shows the fp16, INT8, and NF4 comparison"),
    ("keeps the underlying path out of the reader-facing caption", "uses captions that describe the figure without repository paths"),
    ("capability moves in the same adverse direction", "capability also decreases"),
    ("The cross-family pattern is heterogeneous in direction but homogeneous in inferential status", "The direction of the change differs across model families, but none of the harmful-compliance increases remains significant after correction"),
    ("the direction under the original scorer is in any case benign", "the original scorer reports a decrease rather than an increase"),
    ("The direction under the original scorer is in any case benign", "The original scorer reports a decrease rather than an increase"),
    ("it is a null rather than a borderline finding", "there is insufficient evidence of a change"),
    ("The validation-informed family replaces it", "The validation-informed sensitivity analysis replaces that score"),
    ("the validation-informed family replaces it", "the validation-informed sensitivity analysis replaces that score"),
    ("the validation-informed family", "the alternative judge-based analysis"),
    ("250-item condition", "set of 250 prompts"),
    ("This baseline dependence is another reason not to treat over-refusal as a cross-model leaderboard", "These baseline differences make over-refusal rates unsuitable for ranking models directly"),
    ("Qwen3-1.7B is broad degradation", "Qwen3-1.7B is classified as broad degradation"),
    ("lean toward alignment degradation", "show point estimates in the direction of higher harmful compliance"),
    ("would confound treatment effect with measurement fit", "would mix the effect of quantization with differences between the benchmarks"),
    ("Capability is more orderly", "Capability changes more consistently across precision levels"),
    ("engineering benefit side of that choice", "engineering benefits of the chosen precision"),
    ("Safety does not follow the same shape", "The safety results do not follow the same pattern"),
    ("a single precision-driven alignment mechanism", "one alignment effect caused by lower precision"),
    ("retained only as a transparent secondary view", "retained as a secondary analysis"),
    ("an auditable before-and-after record", "a record of both scoring versions"),
    ("significance status of the substantive per-pair conclusion", "statistical conclusion for each model pair"),
    ("a looser construct cross-check", "a comparison using a broader definition of unsafe content"),
    ("This disagreement is informative rather than proof that one scorer is mechanically wrong", "This disagreement reflects the different behaviours measured by the two scorers and does not by itself show that either scorer is incorrect"),
    ("Every other alias falls in a single lower band", "All other model conditions have lower agreement"),
    ("The 200-item human sample grounds the scorer choice", "The 200-item human sample supports the choice of scorer"),
    ("The validation results can be read as a sequence of eliminated explanations", "The validation analyses rule out several alternative explanations"),
    ("The classifier is not trusted alone either", "The classifier is also checked against other scorers"),
    ("Remaining uncertainty is concentrated in", "The main remaining uncertainties concern"),
    ("Table 6.2 places the published figures beside this study's", "Table 6.2 compares published results with those of this study"),
    ("Reading the columns without that conversion inverts the apparent size", "Comparing the columns without that conversion would give a misleading impression of their relative magnitudes"),
    ("That row agrees with the present null in direction", "That published result is consistent with the absence of an increase in this study"),
    ("Read together, the table supports the narrower reading this report has taken throughout", "Together, these studies support a conclusion limited to the evaluated models and settings"),
    ("The HarmBench regex counted willingness", "The HarmBench regex inferred harmful compliance from refusal phrases"),
    ("In both cases, the evaluator is part of the empirical question", "In both benchmarks, the scoring method affects the reported result"),
    ("The direct-request boundary explains", "The direct-request scope explains"),
    ("The missing cost axis limits the deployment recommendation", "The deployment recommendation is limited because latency, memory use, and energy consumption were not measured"),
    ("A responsible process begins with the application", "Deployment evaluation should begin with the requirements of the intended application"),
    ("Safety should be evaluated on the final serving stack", "Safety should be evaluated using the final deployment configuration"),
    ("Evidence is therefore triangulated", "The conclusion is checked"),
    ("the historical environment-provenance gap", "the incomplete record of the historical software environment"),
    ("It still removes the clean confirmatory label", "The analysis still cannot be described as a fully preregistered confirmation"),
    ("the numerical results should remain attached to the evaluated models and conditions", "the numerical findings should not be generalised beyond the evaluated models and conditions"),
    ("a deployment-relevant GGUF path", "a GGUF deployment configuration"),
    ("strengthen the human reference standard on both safety axes", "add a second annotator to both human-labelled safety datasets"),
    ("judge–regex strict-kappa gap", "difference in strict kappa between the judge and regex"),
    ("a main conclusion average", "an overall average"),
    ("That move is conceptually close to using capability results to interpret safety results", "This serves a similar purpose to checking independent capability results"),
    ("a towards lower harmful compliance shift", "a decrease in harmful compliance"),
    ("Per-alias summaries", "Summaries for each model condition"),
    ("Per-alias classifier-versus-proxy kappa", "Classifier-versus-proxy kappa for each model condition"),
    ("per-alias", "per model condition"),
    ("Per-alias", "For each model condition"),
    ("per alias", "per model condition"),
    ("vary by alias", "vary by model condition"),
    ("precision arm", "precision comparison"),
    ("selected scorer derived scoring files", "derived scoring files for the selected scorer"),
    ("shift towards greater harmful compliance", "increase in harmful compliance"),
    ("directional signal worth replication", "result worth repeating in a larger study"),
    ("co-movement", "simultaneous changes"),
    ("after correction after correction", "after correction"),
    ("common towards greater harmful compliance NF4 effect", "consistent NF4 increase in harmful compliance"),
    ("joint movement", "the combined changes"),
    ("The remaining labels are directional rather than confirmed", "The remaining classifications are based only on the direction of the point estimates"),
    ("Mistral leans toward alignment improvement", "Mistral has a point estimate in the direction of lower harmful compliance"),
    ("corrected movement", "result that remains significant after correction"),
    ("a increase in harmful compliance", "an increase in harmful compliance"),
    ("no towards greater harmful compliance NF4 increase", "no NF4 increase in harmful compliance"),
    ("a effect of the shorter generation limit", "an effect of the shorter generation limit"),
    ("adversarial HarmBench arm", "adversarial HarmBench evaluation"),
    ("clean pre-registered confirmation", "fully preregistered confirmation"),
    ("none of those directional safety labels survives the corrected inferential test", "none of those safety classifications is supported after correction"),
    ("INT8 is the lower-risk precision condition when the memory budget permits it", "INT8 is preferable when preserving capability is more important than achieving greater compression"),
    ("misleading impression of their relative magnitudes of the effects", "misleading impression of the relative effect sizes"),
    ("The present null does not contradict those findings", "The absence of a corrected increase in this study does not contradict those findings"),
    ("sensitivity arms", "sensitivity analyses"),
    ("MMLU and ARC are useful anchors", "MMLU and ARC are useful capability benchmarks"),
    ("The triangulation is useful precisely because these dependencies differ", "Using scorers with different dependencies reduces reliance on any one method"),
    ("Extend using capability results to interpret safety results to", "Extend capability evaluation to"),
]


FORBIDDEN_TERMS = [
    "harmful-ward",
    "safe-ward",
    "bounded null",
    "clean pre-registered badge",
    "claim boundary",
    "public result surface",
    "procedural loyalty",
    "capability cliff",
    "human grounding",
]


def _load_base_builder():
    spec = importlib.util.spec_from_file_location("build_rewritten_fyp_report_v1_for_v2", BASE_BUILDER)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import {BASE_BUILDER}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _all_paragraphs(doc: Document):
    yield from doc.paragraphs
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                yield from cell.paragraphs


def _token_inventory(doc: Document) -> tuple[list[str], list[str]]:
    text = "\n".join(p.text for p in _all_paragraphs(doc))
    citations = re.findall(r"\[(?:\d+)(?:\]-\[\d+)?\]", text)
    text_without_citations = re.sub(r"\[(?:\d+)(?:\]-\[\d+)?\]", "", text)
    numbers = re.findall(
        r"(?<![A-Za-z])[-+−]?(?:\d{2,}(?:\.\d+)?|\d+\.\d+)%?",
        text_without_citations,
    )
    numbers = [token.replace("-", "−") for token in numbers]
    return numbers, citations


def _replace_paragraph_text(paragraph, new_text: str) -> None:
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def _apply_heading_rewrites(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for old, new in HEADING_REWRITES.items():
            if paragraph.text == old:
                _replace_paragraph_text(paragraph, new)
            elif paragraph.text.startswith(old + "\t") and paragraph.runs:
                paragraph.runs[0].text = new
            elif paragraph.text.startswith(old + ".") and paragraph.runs:
                paragraph.runs[0].text = new + paragraph.runs[0].text[len(old):]


def _apply_cached_page_numbers(doc: Document) -> None:
    seen: set[str] = set()
    for paragraph in doc.paragraphs:
        text = paragraph.text
        if "\t" not in text:
            continue
        title, page = text.rsplit("\t", 1)
        if title not in PAGE_NUMBER_REWRITES or not page.isdigit():
            continue
        _replace_paragraph_text(paragraph, f"{title}\t{PAGE_NUMBER_REWRITES[title]}")
        seen.add(title)
    missing = sorted(set(PAGE_NUMBER_REWRITES) - seen)
    if missing:
        raise AssertionError(f"Cached pagination rewrite did not match: {missing}")


def _rewrite_prose(doc: Document) -> dict[str, int]:
    counts = {prefix: 0 for prefix, _ in PARAGRAPH_REWRITES}
    in_report = False
    for paragraph in _all_paragraphs(doc):
        style = paragraph.style.name if paragraph.style is not None else ""
        old_text = paragraph.text

        # The Abstract sits before Chapter 1. Apply full-paragraph rewrites
        # before the chapter-scope guard so it is covered without exposing the
        # table of contents to the general phrase rules.
        matched_full_rewrite = False
        for prefix, replacement in PARAGRAPH_REWRITES:
            if old_text.startswith(prefix):
                counts[prefix] += 1
                if len(paragraph.runs) > 1:
                    raise AssertionError(f"Unexpected multi-run prose paragraph: {old_text[:100]!r}")
                _replace_paragraph_text(paragraph, replacement)
                matched_full_rewrite = True
                break
        if matched_full_rewrite:
            continue

        if paragraph.text == "Chapter 1: Introduction":
            in_report = True
        if not in_report or style == "Reference" or not paragraph.text.strip():
            continue
        if style.startswith("Heading") or style.endswith("Caption"):
            continue

        new_text = old_text
        for old, replacement in PHRASE_REWRITES:
            new_text = re.sub(
                rf"(?<!\w){re.escape(old)}(?!\w)",
                replacement,
                new_text,
            )

        if new_text != old_text:
            if len(paragraph.runs) > 1:
                raise AssertionError(f"Unexpected multi-run prose paragraph: {old_text[:100]!r}")
            _replace_paragraph_text(paragraph, new_text)
    return counts


def build_report() -> Path:
    base = _load_base_builder()
    with tempfile.TemporaryDirectory(prefix="fyp_report_v2_") as tmp:
        tmp_docx = Path(tmp) / "v1_evidence_source.docx"
        original_out = base.OUT
        try:
            base.OUT = tmp_docx
            base.build_report()
        finally:
            base.OUT = original_out

        source_doc = Document(tmp_docx)
        source_numbers, source_citations = _token_inventory(source_doc)
        source_paragraph_count = len(list(_all_paragraphs(source_doc)))

        doc = Document(tmp_docx)
        _apply_heading_rewrites(doc)
        counts = _rewrite_prose(doc)

        missing = [prefix for prefix, count in counts.items() if count != 1]
        if missing:
            details = {prefix: counts[prefix] for prefix in missing}
            raise AssertionError(f"Paragraph rewrite must-fire failure: {details}")

        new_numbers, new_citations = _token_inventory(doc)
        source_number_counts = {token: source_numbers.count(token) for token in set(source_numbers)}
        new_number_counts = {token: new_numbers.count(token) for token in set(new_numbers)}
        missing_numbers = {
            token: (count, new_number_counts.get(token, 0))
            for token, count in source_number_counts.items()
            if new_number_counts.get(token, 0) < count
        }
        invented_values = sorted(set(new_numbers) - set(source_numbers))
        if missing_numbers or invented_values:
            raise AssertionError(
                "Numerical inventory changed during prose rewrite: "
                f"missing={missing_numbers!r}, invented={invented_values!r}"
            )
        if source_citations != new_citations:
            raise AssertionError("Citation marker sequence changed during prose rewrite")
        if len(list(_all_paragraphs(doc))) != source_paragraph_count:
            raise AssertionError("Paragraph count changed during prose rewrite")

        _apply_cached_page_numbers(doc)

        body_text = "\n".join(p.text.lower() for p in _all_paragraphs(doc) if p.style.name != "Reference")
        remaining = [term for term in FORBIDDEN_TERMS if term in body_text]
        if remaining:
            raise AssertionError(f"Legacy report language remains: {remaining}")

        doc.core_properties.subject = "Plain-academic-English FYP report rewrite v2"
        doc.core_properties.comments = (
            "Non-destructive prose rewrite derived from the v1 evidence and layout source. "
            "Numerical tokens, citation markers, structure, tables, and figures are preserved."
        )
        OUT.parent.mkdir(parents=True, exist_ok=True)
        doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(build_report())
