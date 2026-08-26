#!/usr/bin/env python3
"""Build the corpus-informed rewrite of the FYP report.

This builder is intentionally separate from every existing v5/v6 report
builder.  It reads the generated claim registry and redacted analysis sidecars,
then writes ``docs/rewritten_fyp_report_v1.docx`` without changing the
canonical report.

Design preset
-------------
Base preset: ``narrative_proposal``.
Named academic override: ``ntu_empirical_research_report``.

The override follows the observed conventions in the fifteen-paper NTU corpus:
A4 portrait; Times New Roman 11 pt body; 1.5 line spacing; restrained black
chapter hierarchy; sparse title page; Roman-numbered front matter; Arabic
numbering from Chapter 1; separate lists of figures and tables; table captions
above tables; figure captions below figures; IEEE numeric citations; and a
research-report structure centred on method, evidence, and discussion.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Iterable, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = DOCS / "rewritten_fyp_report_v1.docx"
CLAIMS_PATH = ROOT / "generated" / "claim_registry.json"
ANALYSIS = ROOT / "results_512" / "analysis"
FIGURES = DOCS / "figures_512"
FIGURES_FALLBACK = DOCS / "figures"
V6_BUILDER = ROOT / "scripts" / "build_fyp_report_v6.js"

FONT = "Times New Roman"
MUTED = "595959"
HEADER_FILL = "D9E2F3"
LIGHT_FILL = "F2F2F2"
WHITE = "FFFFFF"
BLACK = "000000"
CONTENT_WIDTH_CM = 15.6
NUMERIC_MINUS = re.compile(r"(?<![A-Za-z0-9])-(?=\d)")


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def loadClaimRegistry(path: Path):
    """Load the generated registry through the name recognised by surface checks."""
    return load_json(path)


CLAIM_REGISTRY = loadClaimRegistry(CLAIMS_PATH)
CLAIMS = CLAIM_REGISTRY["claims"]
PAIRS = CLAIMS["pairs"]
PAIR_ORDER = CLAIMS["study"]["pair_order"]
PAIR_NAMES = {pair_id: PAIRS[pair_id]["name"] for pair_id in PAIR_ORDER}

HUMAN_HARMBENCH = load_json(ANALYSIS / "human_validation.json")
HUMAN_XSTEST = load_json(ANALYSIS / "xstest_human_validation.json")
JUDGE_API = load_json(ANALYSIS / "judge_pairwise_agreement.json")
JUDGE_LLAMAGUARD = load_json(ANALYSIS / "judge_pairwise_agreement_llamaguard.json")
MULTISEED = load_json(ANALYSIS / "sensitivity_multiseed.json")
GENLEN = load_json(ANALYSIS / "genlen_robustness.json")
PRECISION = load_json(ANALYSIS / "precision_sweep.json")
CAP_EQ = load_json(ANALYSIS / "capability_equivalence.json")


def extract_v6_references() -> list[str]:
    """Reuse the independently audited v6 bibliography without editing v6."""
    source = V6_BUILDER.read_text(encoding="utf-8")
    start = source.index("const refs = [")
    end = source.index("// ------------------------------------------------------------\n// Appendices", start)
    block = source[start:end]
    array_start = block.index("...numberedList([") + len("...numberedList([")
    array_end = block.index("])\n];", array_start)
    raw = block[array_start:array_end]
    values = [
        json.loads(f'"{value}"')
        for value in re.findall(r'"((?:\\.|[^"\\])*)"', raw)
    ]
    if len(values) != 55:
        raise RuntimeError(f"Expected 55 audited v6 references, found {len(values)}")
    return values


EXTRA_REFERENCES = [
    "W. Shao et al., “OmniQuant: Omnidirectionally calibrated quantization for large language models,” in Proc. International Conference on Learning Representations (ICLR), 2024.",
    "P. Dong, H. Li, and S. Guo, “Durable quantization conditioned misalignment attack on large language models,” in Proc. International Conference on Learning Representations (ICLR), 2025.",
    "S. Li, L. Yao, L. Zhang, and Y. Li, “Safety layers in aligned large language models: The key to LLM security,” in Proc. International Conference on Learning Representations (ICLR), 2025.",
    "A. Souly et al., “A StrongREJECT for empty jailbreaks,” in Proc. NeurIPS Datasets and Benchmarks Track, 2024, doi: 10.52202/079017-3984.",
    "T. Xie et al., “SORRY-Bench: Systematically evaluating large language model safety refusal,” in Proc. International Conference on Learning Representations (ICLR), 2025.",
    "P. Chao et al., “JailbreakBench: An open robustness benchmark for jailbreaking large language models,” in Proc. NeurIPS Datasets and Benchmarks Track, 2024.",
    "S. Han et al., “WildGuard: Open one-stop moderation tools for safety risks, jailbreaks, and refusals of LLMs,” in Proc. NeurIPS Datasets and Benchmarks Track, 2024.",
    "Y. Wang, H. Li, X. Han, P. Nakov, and T. Baldwin, “Do-Not-Answer: Evaluating safeguards in LLMs,” in Findings of EACL, pp. 896–911, 2024, doi: 10.18653/v1/2024.findings-eacl.61.",
    "S. Kim et al., “Prometheus 2: An open source language model specialized in evaluating other language models,” in Proc. EMNLP, pp. 4334–4353, 2024, doi: 10.18653/v1/2024.emnlp-main.248.",
]


# References 1–37 remain in their audited v6 order. The nine verified archival
# sources and the full-text-verified MT-Bench paper are placed before the
# statistical-method references. MT-Bench already exists in the v6 bibliography,
# so it is moved rather than duplicated.
_base_refs = extract_v6_references()
_mt_bench_ref = next(
    reference for reference in _base_refs if "Judging LLM-as-a-judge" in reference
)
_later_refs = [
    reference
    for reference in _base_refs[37:]
    if reference != _mt_bench_ref
]
REFERENCES = _base_refs[:37] + EXTRA_REFERENCES + [_mt_bench_ref] + _later_refs
assert len(REFERENCES) == 64


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_indent(table, dxa: int = 120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(dxa))
    tbl_ind.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(Cm(width_cm).emu / 635)))
    tc_w.set(qn("w:type"), "dxa")


def add_field(paragraph, instruction: str, fallback: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = fallback
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def begin_multiline_field(paragraph, instruction: str) -> None:
    """Start a complex field whose cached result spans multiple paragraphs."""
    run = OxmlElement("w:r")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    run.extend([begin, instr, separate])
    insert_at = 1 if paragraph._p.pPr is not None else 0
    paragraph._p.insert(insert_at, run)


def end_multiline_field(paragraph) -> None:
    run = OxmlElement("w:r")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run.append(end)
    paragraph._p.append(run)


def set_update_fields_on_open(doc: Document) -> None:
    """Ensure Word refreshes the TOC and other fields when the DOCX opens."""
    settings = doc.settings.element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.insert(0, update)
    update.set(qn("w:val"), "true")


def normalize_numeric_minus_signs(doc: Document) -> None:
    """Use U+2212 for negative numbers without changing hyphenated names."""
    roots = [doc.element.body]
    for section in doc.sections:
        roots.extend([section.header._element, section.footer._element])
    seen = set()
    for root in roots:
        if id(root) in seen:
            continue
        seen.add(id(root))
        for text_node in root.iter(qn("w:t")):
            if text_node.text:
                text_node.text = NUMERIC_MINUS.sub("−", text_node.text)


def set_page_number_format(section, fmt: str, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:fmt"), fmt)
    pg_num.set(qn("w:start"), str(start))


def set_footer_page_number(section, roman: bool = False) -> None:
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.font.name = FONT
    run.font.size = Pt(9)
    add_field(paragraph, "PAGE", "i" if roman else "1")


def set_running_header(section, text: str) -> None:
    section.header.is_linked_to_previous = False
    paragraph = section.header.paragraphs[0]
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run(text)
    run.font.name = FONT
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Heading 1", 16, 18, 10),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 11, 8, 4),
    ):
        style = styles[name]
        style.font.name = FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.widow_control = True

    if "Front Matter Heading" not in styles:
        front_heading = styles.add_style("Front Matter Heading", WD_STYLE_TYPE.PARAGRAPH)
    else:
        front_heading = styles["Front Matter Heading"]
    front_heading.font.name = FONT
    front_heading.font.size = Pt(16)
    front_heading.font.bold = True
    front_heading.font.color.rgb = RGBColor.from_string(BLACK)
    front_heading.paragraph_format.space_before = Pt(18)
    front_heading.paragraph_format.space_after = Pt(10)
    front_heading.paragraph_format.keep_with_next = True
    front_heading.paragraph_format.widow_control = True

    title = styles["Title"]
    title.font.name = FONT
    title.font.size = Pt(20)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)

    subtitle = styles["Subtitle"]
    subtitle.font.name = FONT
    subtitle.font.size = Pt(12)
    subtitle.font.color.rgb = RGBColor.from_string(BLACK)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)

    caption = styles["Caption"]
    caption.font.name = FONT
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(BLACK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True

    for style_name in ("Figure Caption", "Table Caption"):
        if style_name not in styles:
            style = styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        else:
            style = styles[style_name]
        style.font.name = FONT
        style.font.size = Pt(9)
        style.font.color.rgb = RGBColor.from_string(BLACK)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(4)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    if "Reference" not in styles:
        ref_style = styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        ref_style = styles["Reference"]
    ref_style.font.name = FONT
    ref_style.font.size = Pt(9)
    ref_style.paragraph_format.left_indent = Cm(0.75)
    ref_style.paragraph_format.first_line_indent = Cm(-0.75)
    ref_style.paragraph_format.line_spacing = 1.0
    ref_style.paragraph_format.space_after = Pt(4)

    if "List Bullet Research" not in styles:
        bullet_style = styles.add_style("List Bullet Research", WD_STYLE_TYPE.PARAGRAPH)
    else:
        bullet_style = styles["List Bullet Research"]
    bullet_style.font.name = FONT
    bullet_style.font.size = Pt(11)
    bullet_style.paragraph_format.left_indent = Cm(0.95)
    bullet_style.paragraph_format.first_line_indent = Cm(-0.48)
    bullet_style.paragraph_format.space_after = Pt(4)
    bullet_style.paragraph_format.line_spacing = 1.2


def configure_section(section) -> None:
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.3)
    section.bottom_margin = Cm(2.3)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.4)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)


def add_body(doc: Document, text: str, *, first_line: bool = False):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.first_line_indent = Cm(0.65) if first_line else None
    p.add_run(text)
    return p


def add_paragraphs(doc: Document, paragraphs: Iterable[str]) -> None:
    for text in paragraphs:
        add_body(doc, text)


def add_bullets(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet Research")
        p.add_run(item)


def add_contents_entry(
    doc: Document,
    title: str,
    page: str,
    *,
    level: int = 0,
    bold: bool = False,
):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.55 * level)
    p.paragraph_format.space_after = Pt(2 if level else 4)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.tab_stops.add_tab_stop(
        Cm(15.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS
    )
    run = p.add_run(title)
    run.font.name = FONT
    run.font.size = Pt(10 if level else 11)
    run.bold = bold
    page_run = p.add_run(f"\t{page}")
    page_run.font.name = FONT
    page_run.font.size = Pt(10 if level else 11)
    page_run.bold = bold
    return p


def add_chapter(doc: Document, number: int, title: str) -> None:
    p = doc.add_heading(f"Chapter {number}: {title}", level=1)
    p.paragraph_format.page_break_before = True


def add_heading(doc: Document, number: str, title: str, level: int = 2) -> None:
    doc.add_heading(f"{number} {title}", level=level)


def add_table_caption(doc: Document, number: str, title: str) -> None:
    doc.add_paragraph(f"Table {number}. {title}", style="Table Caption")


def add_figure_caption(doc: Document, number: str, title: str) -> None:
    doc.add_paragraph(f"Figure {number}. {title}", style="Figure Caption")


def add_table(
    doc: Document,
    number: str,
    title: str,
    headers: Sequence[str],
    rows: Sequence[Sequence[str]],
    widths_cm: Sequence[float],
    *,
    font_size: float = 8.4,
    show_caption: bool = True,
    body_alignment=None,
) -> None:
    if abs(sum(widths_cm) - CONTENT_WIDTH_CM) > 0.05:
        raise ValueError(f"Table {number} widths sum to {sum(widths_cm):.2f} cm")
    if show_caption:
        add_table_caption(doc, number, title)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_indent(table, 120)
    header_row = table.rows[0]
    set_repeat_table_header(header_row)
    for cell, header, width in zip(header_row.cells, headers, widths_cm):
        cell.text = header
        set_cell_width(cell, width)
        set_cell_margins(cell)
        set_cell_shading(cell, HEADER_FILL)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for values in rows:
        row = table.add_row()
        set_row_cant_split(row)
        for cell, value, width in zip(row.cells, values, widths_cm):
            cell.text = str(value)
            set_cell_width(cell, width)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_with_next = False
                if body_alignment is not None and row_index > 0:
                    paragraph.alignment = body_alignment
                for run in paragraph.runs:
                    run.font.name = FONT
                    run.font.size = Pt(font_size)
                    run.bold = row_index == 0
    doc.add_paragraph()


def resolve_figure(filename: str) -> Path:
    preferred = FIGURES / filename
    if preferred.exists():
        return preferred
    fallback = FIGURES_FALLBACK / filename
    if fallback.exists():
        return fallback
    raise FileNotFoundError(filename)


def add_figure(
    doc: Document,
    number: str,
    title: str,
    filename: str,
    width_cm: float = 14.4,
) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    picture = p.add_run().add_picture(str(resolve_figure(filename)), width=Cm(width_cm))
    alt_text = f"Figure {number}. {title}"
    picture._inline.docPr.set("descr", alt_text)
    picture._inline.docPr.set("title", alt_text)
    add_figure_caption(doc, number, title)


def fmt_delta(value: float, digits: int = 3) -> str:
    return f"{value:+.{digits}f}"


def fmt_rate(value: float, digits: int = 3) -> str:
    return f"{value:.{digits}f}"


def fmt_ci(ci: Sequence[float], digits: int = 3) -> str:
    return f"[{ci[0]:.{digits}f}, {ci[1]:.{digits}f}]"


def add_cover(doc: Document) -> None:
    for _ in range(5):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("NANYANG TECHNOLOGICAL UNIVERSITY")
    r.font.name = FONT
    r.font.size = Pt(14)
    r.bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("College of Computing and Data Science")
    r.font.name = FONT
    r.font.size = Pt(12)
    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("Benchmarking Ethical Performance of Open-Source LLMs")
    r.font.name = FONT
    r.font.size = Pt(20)
    r.bold = True
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(
        "A Matched-Pair Study of Safety and Capability in Quantized Compact Language Models"
    )
    r.font.name = FONT
    r.font.size = Pt(12)
    r.italic = True
    for _ in range(3):
        doc.add_paragraph()
    # The degree line follows the corpus convention (P02, the most recent past paper):
    # it sits between the supervisor and the year. Keep the wording identical to
    # `\degreename` in scripts/build_rewritten_latex.py so the two mirrors agree.
    for line, bold in (
        ("Final Year Project Report", True),
        ("Project Code: CCDS25-1136", False),
        ("TAN UEI HORNG (UTAN001)", True),
        ("Supervisor: Dr. Zhang Jiehuang", False),
        (
            "Submitted in Partial Fulfilment of the Requirements for the Degree of "
            "Bachelor of Engineering (Computer Science)",
            False,
        ),
        ("Academic Year 2025/2026", False),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(line)
        r.font.name = FONT
        r.font.size = Pt(11)
        r.bold = bold


def add_front_matter(doc: Document) -> None:
    doc.add_paragraph("Declaration of Originality", style="Front Matter Heading")
    add_body(
        doc,
        "I declare that this report is my own work and that all sources used have been acknowledged. "
        "The research design, experimental decisions, analysis, and interpretation are my "
        "responsibility. The reported results were computed from the recorded experiment artefacts; "
        "no result was fabricated or selectively altered.",
    )
    add_body(
        doc,
        "Generative-AI tools assisted software implementation, test authoring, document review, and "
        "language editing under my direction. I checked the resulting code, citation metadata, numerical "
        "claims, and prose, and I take responsibility for the final submission.",
    )

    doc.add_page_break()
    doc.add_paragraph("Acknowledgements", style="Front Matter Heading")
    add_body(
        doc,
        "I thank my supervisor, Dr. Zhang Jiehuang, for his guidance throughout this project. I also "
        "thank the College of Computing and Data Science and the TC1 support team for providing the "
        "computing environment used in the experiments. Finally, I am grateful to my family and "
        "friends for their patience and encouragement during the project.",
    )

    doc.add_page_break()
    doc.add_paragraph("Abstract", style="Front Matter Heading")
    add_body(
        doc,
        "Post-training quantization makes compact instruction-tuned language models practical on "
        "resource-constrained hardware, but it can also alter behaviour. This report investigates "
        "whether four-bit quantization changes harmful compliance or over-refusal, and whether an "
        "apparent safety change is instead explained by capability loss. Five matched model pairs "
        "from Qwen, Llama, Mistral, and Phi were evaluated at fp16, INT8, and NF4 precision. HarmBench "
        "measured harmful compliance, XSTest measured refusal on benign prompts, and MMLU plus "
        "ARC-Challenge provided capability anchors. Paired bootstrap intervals, exact McNemar tests, "
        "and Benjamini–Hochberg correction were used. The HarmBench classifier was cross-checked by "
        "an API judge, Llama Guard, and a human-labelled sample. At HarmBench's reference generation "
        "budget of 512 new tokens, "
        "no NF4 harmful-compliance contrast survived correction. The corrected family contained "
        "three survivors: a Qwen3-1.7B MMLU decrease, a Llama-3.2-3B ARC decrease, and a Phi-4-mini "
        "over-refusal decrease. Human validation showed that the semantic refusal judge aligned far "
        "better with the annotator than the regex, so the Phi result is best treated as scorer-"
        "dependent. The evidence therefore supports a bounded conclusion: within the evaluated "
        "direct-request setting, NF4 produced detectable capability costs but no corrected increase "
        "in harmful compliance.",
    )

    doc.add_page_break()
    doc.add_paragraph("Table of Contents", style="Front Matter Heading")
    toc_entries = [
        ("Chapter 1: Introduction", "1", 0),
        ("1.1 Background and Motivation", "1", 1),
        ("1.2 Problem Statement", "1", 1),
        ("1.3 Research Questions", "3", 1),
        ("1.4 Scope", "3", 1),
        ("1.5 Contributions", "4", 1),
        ("1.6 Report Structure", "4", 1),
        ("Chapter 2: Literature Review", "5", 0),
        ("2.1 Trustworthiness and Safety Evaluation", "5", 1),
        ("2.2 Quantization Methods and Capability Effects", "6", 1),
        ("2.3 Quantization and Safety Behaviour", "7", 1),
        ("2.4 Refusal, Harmfulness, and Threat Models", "8", 1),
        ("2.5 Automatic Evaluators and Human Grounding", "10", 1),
        ("2.6 Synthesis and Research Gap", "11", 1),
        ("2.7 Implications for Experimental Design", "12", 1),
        ("Chapter 3: Methodology and Experimental Setup", "13", 0),
        ("3.1 Matched-Pair Experimental Design", "13", 1),
        ("3.2 Models and Precision Conditions", "14", 1),
        ("3.3 Benchmarks and Outcomes", "15", 1),
        ("3.4 Scoring and Validation", "16", 1),
        ("3.5 Statistical Analysis", "16", 1),
        ("3.6 Implementation and Execution", "18", 1),
        ("3.7 Reproducibility Controls", "19", 1),
        ("3.8 Analysis Workflow and Quality Control", "19", 1),
        ("Chapter 4: Results", "21", 0),
        ("4.1 Results at a Glance", "21", 1),
        ("4.2 Harmful Compliance (RQ1)", "22", 1),
        ("4.3 Over-Refusal on Benign Prompts (RQ2)", "23", 1),
        ("4.4 Capability (RQ3)", "24", 1),
        ("4.5 Model Scale and Family Differences (RQ4)", "25", 1),
        ("4.6 INT8 Precision Point", "25", 1),
        ("4.7 Research-Question Synthesis", "26", 1),
        ("Chapter 5: Scorer Validation and Robustness", "28", 0),
        ("5.1 HarmBench Classifier versus Refusal Proxy", "28", 1),
        ("5.2 Independent HarmBench Judges", "29", 1),
        ("5.3 Human Grounding of Harmful-Compliance Scorers", "30", 1),
        ("5.4 XSTest Scorer Sensitivity and Human Validation", "30", 1),
        ("5.5 Generation-Length Robustness", "31", 1),
        ("5.6 Multi-Seed Decoding Sensitivity", "32", 1),
        ("5.7 Power and the Meaning of the Null", "33", 1),
        ("5.8 Robustness Synthesis", "34", 1),
        ("Chapter 6: Discussion", "35", 0),
        ("6.1 Answer to the Central Question", "35", 1),
        ("6.2 Relationship to Prior Work", "35", 1),
        ("6.3 Capability as the Deployment Cost", "38", 1),
        ("6.4 Scorer Validity as a Research Result", "38", 1),
        ("6.5 Threats to Validity", "39", 1),
        ("6.6 Practical Implications", "41", 1),
        ("Chapter 7: Conclusion and Future Work", "42", 0),
        ("7.1 Conclusions", "42", 1),
        ("7.2 Prioritised Future Work", "42", 1),
        ("References", "44", 0),
        ("Appendix A: Reproducibility Summary", "47", 0),
        ("Appendix B: Minimum Detectable Capability Effects", "48", 0),
        ("Appendix C: Human-Validation Boundaries", "49", 0),
    ]
    toc_paragraphs = [
        add_contents_entry(doc, title, page, level=level, bold=level == 0)
        for title, page, level in toc_entries
    ]
    begin_multiline_field(
        toc_paragraphs[0],
        ' TOC \\o "1-2" \\h \\z \\u ',
    )
    end_multiline_field(toc_paragraphs[-1])

    doc.add_page_break()
    doc.add_paragraph("List of Figures", style="Front Matter Heading")
    for title, page in [
        ("Figure 1.1. Capability-anchored safety space using MMLU and HarmBench deltas", "3"),
        ("Figure 3.1. Matched pairs, precision arms, benchmark battery, and analysis flow", "14"),
        ("Figure 4.1. Per-pair HarmBench classifier ΔASR with paired confidence intervals", "23"),
        ("Figure 4.2. Precision sweep across fp16, INT8, and NF4", "26"),
        ("Figure 5.1. HarmBench classifier ASR versus the refusal-regex proxy", "29"),
        ("Figure 5.2. Qwen3-1.7B greedy result and five-seed HarmBench sensitivity", "33"),
    ]:
        add_contents_entry(doc, title, page)

    doc.add_page_break()
    doc.add_paragraph("List of Tables", style="Front Matter Heading")
    for title, page in [
        ("Table 2.1. Positioning of representative prior work", "11"),
        ("Table 3.1. Model pairs and precision conditions", "14"),
        ("Table 3.2. Benchmark roles, sample sizes, and primary outcomes", "15"),
        ("Table 4.1. NF4 safety and over-refusal results", "21"),
        ("Table 4.2. NF4 capability results", "21"),
        ("Table 5.1. Per-pair HarmBench deltas under three judges", "29"),
        ("Table 5.2. HarmBench classifier deltas at 128 and 512 tokens", "32"),
        ("Table 5.3. Five-seed HarmBench sensitivity", "32"),
        ("Table 6.1. Competing explanations for the HarmBench null", "35"),
        ("Table 6.2. Published quantization effects placed beside this study's", "36"),
        ("Table A.1. Primary run controls", "47"),
        ("Table B.1. Capability effects and minimum detectable effects", "48"),
    ]:
        add_contents_entry(doc, title, page)

    doc.add_page_break()
    doc.add_paragraph("Abbreviations", style="Front Matter Heading")
    rows = [
        ("ARC", "AI2 Reasoning Challenge"),
        ("ASR", "Attack Success Rate"),
        ("BH-FDR", "Benjamini–Hochberg false-discovery-rate control"),
        ("CI", "Confidence interval"),
        ("INT8", "Eight-bit integer quantization using LLM.int8"),
        ("LLM", "Large language model"),
        ("MMLU", "Massive Multitask Language Understanding"),
        ("NF4", "NormalFloat four-bit quantization"),
        ("OR", "Over-refusal rate"),
        ("PTQ", "Post-training quantization"),
        ("TC1", "NTU GPU compute cluster used in this project"),
    ]
    add_table(
        doc,
        "F.1",
        "Abbreviations used in the report",
        ["Term", "Meaning"],
        rows,
        [3.2, 12.4],
        show_caption=False,
    )


def add_chapter_one(doc: Document) -> None:
    add_chapter(doc, 1, "Introduction")

    add_heading(doc, "1.1", "Background and Motivation")
    add_paragraphs(
        doc,
        [
            "Compact instruction-tuned language models have moved capable generation into a "
            "deployment range that includes consumer GPUs and edge hardware. Qwen3 [1], the Llama "
            "3 family [2], Phi-4-mini [3], and Gemma 2 [4] illustrate the rapid development of "
            "models below the scale of frontier systems. In a system designed without remote "
            "inference, inputs need not leave the device. However, this report does not assume a "
            "latency, energy, privacy, or monetary benefit from that architecture, and it "
            "measures only model behaviour. Full-precision inference can nevertheless exceed a "
            "target device's memory budget, which motivates deployment-oriented compression [5].",
            "Post-training quantization addresses that constraint by representing weights, and "
            "sometimes activations, at lower precision. AWQ reports deployment-oriented four-bit "
            "compression and acceleration on desktop and mobile GPUs [5]. GPTQ demonstrates one-"
            "shot weight quantization at model scales up to 175 billion parameters [6]. Across a "
            "large scaling-law study, four-bit inference is often the most effective operating "
            "point under a fixed memory budget [7]. Together these results explain why a four-bit "
            "model is not an exotic laboratory artefact, but a plausible deployment form of the "
            "model.",
            "The quantization method papers reviewed here report perplexity, accuracy, latency, or "
            "memory as primary outcomes [5]-[7]. Broad comparisons find that many tasks remain "
            "usable at four bits [8], [9]. "
            "However, a small change in average accuracy does not establish that safety behaviour is "
            "unchanged. Safety-specific studies now report that quantization effects vary across "
            "models, bit-widths, and methods [10]. Quantization can also disturb confidence even when "
            "accuracy looks stable [11], create a condition that an attacker can exploit [12], and "
            "produce language-dependent damage that automatic metrics understate [13].",
            "Refusal behaviour gives a concrete reason for concern. Helpful and harmless "
            "assistant training introduces a learned behavioural policy [14]. Mechanistic work "
            "finds that refusal can be mediated by a concentrated direction in activation space "
            "[15], while other evidence suggests that safety alignment may be shallow in the "
            "generation process [16]. Fine-tuning on a very small adversarial dataset can weaken "
            "safeguards [17]. These findings do not prove that ordinary NF4 quantization will "
            "damage safety. They do, however, show that a weight-modifying deployment step should "
            "be measured rather than assumed benign.",
        ],
    )

    add_heading(doc, "1.2", "Problem Statement")
    add_paragraphs(
        doc,
        [
            "Safety evaluation is difficult because refusal is not synonymous with safety. A model "
            "that refuses every request may appear harmless, but it is not useful. XSTest was created "
            "to expose exaggerated safety on benign prompts [18], and OR-Bench later expanded the "
            "same concern at much larger scale [19]. Conversely, a model that does not use an obvious "
            "refusal phrase has not necessarily supplied harmful information. A safety score can "
            "therefore move because the model's alignment changed, because its capability changed, "
            "or because the scorer interprets the response differently.",
            "This project asks whether quantization changes the balance among harmful compliance, "
            "benign refusal, and general capability in compact instruction-tuned models. The load-"
            "bearing distinction is between an alignment shift and a capability side-effect. A lower "
            "attack-success rate is not evidence of improved alignment if the same model also loses "
            "the ability to answer ordinary questions. The study consequently evaluates the three "
            "axes together and validates the instruments used to score them.",
            "Figure 1.1 previews the answer. Each pair is placed by its change in MMLU accuracy "
            "and its change in judge-scored harmful compliance, with paired bootstrap intervals "
            "on both axes, against the interpretation regions defined in Chapter 3. No interval "
            "excludes zero in the harmful-increase direction. Llama's interval excludes zero only "
            "in the safety-improving direction, while the point estimates on the capability axis "
            "are negative for all five pairs. Chapter 4 reports which of these effects survive "
            "multiplicity correction, and Chapter 5 examines whether the pattern holds under a "
            "change of scorer.",
        ],
    )
    add_figure(
        doc,
        "1.1",
        "Capability-anchored safety space using MMLU and HarmBench deltas.",
        "capability_anchor.png",
        13.8,
    )

    add_heading(doc, "1.3", "Research Questions")
    add_body(
        doc,
        "The report uses five research questions. The first three define the primary NF4 comparison. "
        "The remaining questions test whether the interpretation generalises across models, "
        "precisions, scorers, and generation conditions.",
    )
    add_bullets(
        doc,
        [
            "RQ1. Does NF4 quantization measurably increase harmful compliance relative to matched fp16 weights?",
            "RQ2. Does NF4 quantization increase over-refusal on benign prompts?",
            "RQ3. Does NF4 quantization reduce capability on MMLU or ARC-Challenge?",
            "RQ4. How do the effects vary across model families, model scales, and the additional INT8 precision point?",
            "RQ5. Does the substantive conclusion remain credible under alternative scorers, human grounding, generation budgets, and decoding seeds?",
        ],
    )

    add_heading(doc, "1.4", "Scope")
    add_paragraphs(
        doc,
        [
            "The primary experiment covers five matched fp16–NF4 pairs from four model families: "
            "Qwen3-1.7B, Qwen3-4B, Llama-3.2-3B, Mistral-7B, and Phi-4-mini. A separate INT8 arm uses "
            "the same baselines. The safety threat model is HarmBench DirectRequest: models receive "
            "unsafe instructions without an optimised jailbreak suffix or attack loop. The results "
            "therefore concern direct harmful requests and should not be interpreted as evidence of "
            "general adversarial robustness.",
            "The capability instruments are a six-subject MMLU subset and ARC-Challenge. They "
            "provide two distinct multiple-choice anchors, but they do not measure every form of "
            "capability. The main generation budget is 512 new tokens and the primary decoding "
            "condition is greedy, while a five-seed stochastic arm is a sensitivity analysis "
            "rather than part of the primary family of tests.",
        ],
    )

    add_heading(doc, "1.5", "Contributions")
    add_bullets(
        doc,
        [
            "A controlled matched-pair design that applies NF4 or LLM.int8 at load time to the same revision-pinned weights.",
            "A joint evaluation of harmful compliance, benign over-refusal, MMLU, and ARC-Challenge across five model pairs.",
            "A paired statistical analysis with bootstrap confidence intervals, exact McNemar tests, multiplicity control, and power-oriented interpretation.",
            "A scorer-validation study using the HarmBench classifier, an independent API judge, Llama Guard, and two human-labelled samples.",
            "A reproducible benchmark implementation with redacted derived sidecars and immutable original run artefacts.",
        ],
    )

    add_heading(doc, "1.6", "Report Structure")
    add_body(
        doc,
        "Chapter 2 reviews quantization, safety and refusal evaluation, capability anchoring, and "
        "automatic scoring. Chapter 3 describes the experimental design, implementation, statistical "
        "analysis, and reproducibility controls. Chapter 4 presents the primary and precision results. "
        "Chapter 5 evaluates scorer validity and robustness. Chapter 6 discusses the findings, their "
        "relationship to prior work, and the main threats to validity. Chapter 7 concludes and sets "
        "out the most valuable extensions. Essential reproducibility and supplementary material appear "
        "in the appendices.",
    )


def add_chapter_two(doc: Document) -> None:
    add_chapter(doc, 2, "Literature Review")
    add_body(
        doc,
        "This review is organised around the variables that the experiment must control: the "
        "compression method, the safety construct, the capability anchor, the threat model, and the "
        "scorer. The final section compares representative studies directly and derives the research "
        "gap. The aim is synthesis rather than a catalogue of papers.",
    )

    add_heading(doc, "2.1", "Trustworthiness and Safety Evaluation")
    add_paragraphs(
        doc,
        [
            "TrustLLM evaluates trustworthiness across safety, truthfulness, fairness, "
            "robustness, privacy, and machine ethics [20]. DecodingTrust provides another broad "
            "assessment, with detailed testing of toxicity, stereotypes, privacy, robustness, and "
            "related properties [21]. SafetyBench focuses more directly on safety knowledge "
            "through a large bilingual multiple-choice benchmark [22]. These suites established "
            "the importance of evaluating more than ordinary task accuracy. However, they also "
            "reveal a limitation for the present question: the model is generally treated as a "
            "fixed system rather than as one member of a matched precision pair.",
            "Instruction following and safety are not independent. The InstructGPT study shows how "
            "human feedback can make a much smaller aligned model preferable to a far larger base "
            "model [23]. HarmBench then supplies a standardised framework for automated red teaming "
            "and harmful-compliance evaluation [24]. Its classifier asks whether the response "
            "satisfies the harmful behaviour, which is more specific than checking whether the model "
            "used a refusal phrase.",
            "A separate literature studies adversarial attacks. Universal suffix attacks can transfer "
            "across prompts and models [25], and model-generated red-team prompts can uncover failure "
            "cases at scale [26]. These studies measure the combination of the target model and an "
            "attack procedure. The present experiment instead uses direct harmful requests, so its "
            "claim is narrower: it estimates a precision effect under the DirectRequest condition.",
            "Capability anchors are needed because apparent safety can be purchased by inability. "
            "MMLU measures broad academic and professional knowledge [27], while ARC-Challenge tests "
            "difficult science questions [28]. Neither represents all useful behaviour. Their value "
            "here is diagnostic: if a lower harmful-compliance score arrives with lower capability, "
            "the result cannot be described confidently as an alignment improvement.",
            "The three broad trustworthiness suites also illustrate a recurring trade-off in "
            "benchmark design. Breadth increases coverage, but it can make the treatment and "
            "scorer less visible. A composite trustworthiness score may combine factuality, "
            "privacy, bias, toxicity, and robustness even though those dimensions respond "
            "differently to compression. This is useful for model comparison, but less suitable "
            "for causal attribution. Consequently, a matched quantization study benefits from "
            "fewer outcomes with explicit constructs, because the route from a prompt-level "
            "decision to the reported metric can be inspected.",
            "SafetyBench's multiple-choice format is valuable for safety knowledge, but knowing "
            "the correct policy answer is not the same as refusing a harmful request during "
            "generation [22]. HarmBench and XSTest evaluate generated responses, so they are "
            "closer to deployment behaviour. However, their outputs are harder to score, because "
            "the scorer must decide whether a response complies, refuses, equivocates, or "
            "supplies partial information. That additional measurement uncertainty is one reason "
            "the present study treats scorer validation as a research result rather than a "
            "preprocessing detail.",
            "MMLU and ARC also differ in useful ways. MMLU spans many academic and professional "
            "domains but the present project uses a six-subject subset, which limits subject "
            "coverage. ARC-Challenge is larger in this experiment and concentrates on difficult "
            "science questions. Agreement between them would strengthen a broad capability "
            "interpretation, whereas disagreement suggests that the effect is task-specific. The "
            "report therefore avoids collapsing the two benchmarks into one undocumented "
            "capability index.",
        ],
    )

    add_heading(doc, "2.2", "Quantization Methods and Capability Effects")
    add_paragraphs(
        doc,
        [
            "Quantization methods differ in the quantities they compress and in the information "
            "they use during calibration. Scaling work shows that larger models often tolerate "
            "reduced precision better than smaller ones, but the relationship depends on model "
            "scale and task [29]. SmoothQuant migrates activation difficulty into weights to "
            "enable W8A8 inference [30], while SparseGPT demonstrates that a different "
            "compression operation, pruning, can remove many weights in one shot [31]. Therefore, "
            "these methods are not interchangeable interventions.",
            "NF4 was introduced with QLoRA as a four-bit data type suited to approximately normally "
            "distributed model weights [32]. In this project, the BitsAndBytes implementation applies "
            "it when the checkpoint is loaded, which enables a matched comparison from identical weights. "
            "LLM.int8 uses a different mixed-precision strategy that routes outlier features through "
            "a higher-precision path [33]. The INT8 arm in this study is therefore an additional "
            "method-and-precision point, not simply an eight-bit version of NF4.",
            "Accuracy can remain stable while other numerical properties move. Modern neural "
            "networks are often poorly calibrated [34], and quantization-specific work reports that "
            "confidence changes concentrate on difficult or low-confidence inputs [11]. This motivates "
            "the hypothesis that borderline behavioural decisions may move even when average accuracy "
            "changes little. It does not establish that any refusal decision changed in the present models.",
            "GPTQ, AWQ, SmoothQuant, NF4, and LLM.int8 protect different numerical structures. GPTQ "
            "uses approximate second-order information to minimise reconstruction error during "
            "weight quantization [6]. AWQ's motivating analysis identifies approximately 0.1%–1% of "
            "weights as salient; its implemented method protects their contribution through activation-"
            "aware scaling rather than a literal mixed-precision layout [5]. SmoothQuant addresses activation outliers through an "
            "equivalent transformation [30]. NF4 uses a codebook designed for the distribution of "
            "neural weights [32], while LLM.int8 separates outlier features into a higher-precision "
            "path [33]. A result obtained with one method is evidence about that intervention, not "
            "about all low-bit arithmetic.",
            "The scale of the calibration data also differs. Some methods require representative "
            "samples and an optimisation pass, while BitsAndBytes load-time quantization can be "
            "applied without a task-specific calibration set. That convenience makes NF4 suitable for "
            "this controlled load-time experiment. Relative to calibrated methods, it provides no "
            "task-specific optimisation step that could be used to preserve a hypothesised safety-"
            "sensitive subspace. The present study measures this load-time path rather than tuning the "
            "quantizer to preserve benchmark behaviour.",
            "Capability-focused evaluations provide a useful prior expectation. Four-bit results "
            "are often close to full precision on aggregate tasks, especially as model size grows "
            "[7]-[9]. However, average resilience can conceal subgroup damage. Marchisio et al. "
            "show that multilingual quantization effects differ by language, and on the paper's "
            "Japanese evaluation the human-assessed loss is substantially larger than the "
            "automatic-metric loss [13]. The parallel for safety is direct: a headline average is "
            "insufficient if scorer error is concentrated in one model family or response style.",
            "The INT8 arm helps separate aggressiveness from method. If safety and capability changed "
            "smoothly as precision fell, a graded fp16–INT8–NF4 curve would be plausible. A non-"
            "monotonic pattern would instead implicate model-specific numerics, distinct algorithms, "
            "or ordinary sampling uncertainty. The experiment treats both possibilities as empirical "
            "questions rather than assuming that fewer bits must produce more behavioural damage.",
        ],
    )

    add_heading(doc, "2.3", "Quantization and Safety Behaviour")
    add_paragraphs(
        doc,
        [
            "HarmLevelBench reports non-uniform safety changes across quantization methods and attack "
            "types [35]. Q-resafe similarly treats quantized safety degradation as a distinct problem "
            "and proposes quantization-aware repair [36]. In contrast, Decoding Compressed Trust finds "
            "that four-bit quantization broadly preserves the trustworthiness dimensions it studies, "
            "while more extreme compression and pruning produce larger damage [37]. The literature "
            "therefore does not support a single statement that quantization either preserves or harms "
            "alignment. Method, bit-width, model, task, and scorer all matter.",
            "OmniQuant learns weight-clipping and equivalent transformations during calibration [38]. "
            "Its method differs from NF4, reinforcing the need to name the intervention rather than "
            "generalise from 'quantization' as a single operation. Q-Misalign goes further by preparing "
            "a model whose misalignment is triggered by quantization [39]. This is important security "
            "evidence, but it does not estimate the effect of ordinary post-training quantization on an "
            "unaltered release checkpoint.",
            "Mechanistic studies of safety layers offer another perspective. Safety Layers argues that "
            "aligned behaviour can be concentrated in a subset of model layers [40]. Together with "
            "work on refusal directions and shallow alignment, this provides a plausible mechanism by "
            "which compression could move safety decisions unevenly. It remains a hypothesis for the "
            "present models, because the experiment measures behaviour rather than identifying a causal "
            "internal circuit.",
            "The safety-specific studies differ in their comparison units. A study that compares "
            "publisher-supplied quantized checkpoints may inherit differences in conversion settings, "
            "tokenizers, chat templates, or model revisions. A matched load-time design removes many "
            "of those differences, although it does not represent pre-quantized deployment formats "
            "such as GGUF. This distinction explains why the present project prioritises internal "
            "validity for one quantization path over breadth across many community artefacts.",
            "They also differ in outcome construction. Some studies use refusal or toxicity as a "
            "safety proxy, others use broad trustworthiness suites, and HarmLevelBench grades the "
            "level of harm in a response [35]. These are related but non-identical constructs. "
            "For example, a quantized model could become more willing to answer without becoming "
            "more capable of completing a harmful task, or it could retain its refusal language "
            "while leaking information. Consequently, the chosen evaluator determines which of "
            "these changes is counted.",
            "The adversarial studies establish possibility rather than prevalence. Q-Misalign shows "
            "that an attacker can create a durable quantization-conditioned failure [39], and "
            "Exploiting LLM Quantization shows another security-relevant pathway [12]. Those results "
            "justify evaluating the deployed artefact. They do not imply that every normally trained "
            "checkpoint will fail under ordinary NF4 loading. A controlled benign-checkpoint study is "
            "needed precisely because worst-case vulnerability and typical treatment effect are "
            "different questions.",
            "Decoding Compressed Trust is the most compatible prior result because it observes broad "
            "four-bit preservation [37]. Even there, the trust dimensions, scorer choices, model sizes, "
            "and compression configurations differ from this project. The appropriate comparison is "
            "therefore qualitative: both studies find that four-bit compression need not create a "
            "large general safety regression, while both leave room for model- and dimension-specific "
            "effects.",
        ],
    )

    add_heading(doc, "2.4", "Refusal, Harmfulness, and Threat Models")
    add_paragraphs(
        doc,
        [
            "A refusal detector and a harmfulness classifier answer different questions. StrongREJECT "
            "was motivated by 'empty jailbreaks': cases in which a target does not refuse but also "
            "does not provide useful harmful information [41]. Its evaluator combines refusal and "
            "response quality. The study shows why a non-refusal rate can exaggerate attack success, "
            "but it does not itself test quantization.",
            "SORRY-Bench evaluates safety refusal across 44 fine-grained unsafe topics, 440 balanced "
            "instructions, 20 linguistic augmentations, and 7,040 human annotations used in its "
            "evaluator study [42]. Its "
            "authors explicitly discuss the uselessness of a model that refuses everything. XSTest "
            "provides a smaller human-curated diagnostic set, while OR-Bench supplies breadth. The "
            "common implication is that harmful-request refusal must be interpreted beside benign "
            "false refusal and utility.",
            "JailbreakBench standardises attack artefacts, chat templates, system prompts, scoring, "
            "and attack-efficiency measures [43]. That reporting discipline matters because attack "
            "success depends on more than the target checkpoint. WildGuard spans direct and "
            "adversarial interactions and predicts prompt harmfulness, response harmfulness, and "
            "refusal [44]. It is useful as an independent moderator, but its policy construct is not "
            "identical to HarmBench's behaviour-specific criterion. Agreement and disagreement must "
            "therefore be interpreted, not merely ranked.",
            "Direct requests and jailbreaks also allocate causal responsibility differently. In a "
            "DirectRequest test, the prompt is fixed and the contrast focuses on the target model. In "
            "a jailbreak benchmark, success reflects the target, the attack optimiser, its query "
            "budget, and the evaluator. JailbreakBench records these elements because otherwise two "
            "attack-success rates may not be comparable [43]. The present report follows the same "
            "discipline by naming its absence of attack augmentation wherever the safety conclusion "
            "is stated.",
            "XSTest and OR-Bench offer another complementary pair. XSTest is small enough for detailed "
            "human inspection and deliberately targets lexical and pragmatic triggers [18]. OR-Bench "
            "contains 80,000 prompts across ten rejection categories, a hard subset of about 1,000 "
            "prompts, and 600 toxic controls; the final ICML study evaluates 32 models across eight "
            "families [19]. The smaller benchmark is appropriate "
            "for a paired multi-model experiment with human validation, but it cannot claim the breadth "
            "of the larger corpus. A future extension should test whether the scorer findings persist "
            "across OR-Bench categories.",
            "StrongREJECT introduces response quality into jailbreak evaluation [41]. That move "
            "is conceptually close to capability anchoring, because both prevent inability from "
            "being counted as safety. However, the implementations differ: StrongREJECT assesses "
            "how fully an unsafe request is answered, while this study uses independent MMLU and "
            "ARC tasks to detect broader capability loss. The latter cannot tell whether a "
            "particular harmful answer is low quality, but it can reveal a general degradation "
            "that would make a safe-ward shift less persuasive.",
            "SORRY-Bench grounds refusal evaluation in a human-in-the-loop taxonomy of unsafe topics, and "
            "scores refusal on an explicitly binary scale, leaving finer-grained harmfulness scoring "
            "as future work [42]. That binary refusal is itself a simplification is shown instead by "
            "finer-grained annotation schemes: Do-Not-Answer's six action categories [45] and XSTest's "
            "full-compliance, partial-refusal and full-refusal labelling [18]. Partial refusals, "
            "caveated answers, and mixed responses occupy the boundary. The three-class XSTest judge "
            "used here adopts that finer scheme. Strict and broad "
            "mappings are reported separately so the report does not hide the policy choice inside a "
            "single label.",
        ],
    )

    add_heading(doc, "2.5", "Automatic Evaluators and Human Grounding")
    add_paragraphs(
        doc,
        [
            "Do-Not-Answer evaluates 939 risky prompts and 5,634 responses and finds that a trained "
            "600-million-parameter classifier can be competitive with a much larger judge in "
            "aggregate [45]. It also reports variation across source models. A scorer can therefore "
            "look adequate overall while making model-specific errors that distort a paired contrast.",
            "Prometheus 2 provides an open evaluator for direct assessment and pairwise ranking under "
            "user-specified criteria [46]. Its authors treat agreement with human or proprietary judges "
            "as an indirect measure and recommend using model-based evaluation together with human "
            "evaluation. The "
            "practical lesson is a hierarchy of evidence: a transparent deterministic scorer is "
            "valuable for reproducibility, a semantic judge tests construct coverage, and a human-"
            "labelled sample grounds both.",
            "Agreement should be reported by model or pair when possible. A pooled coefficient can "
            "hide a failure concentrated in the very model that determines the conclusion. The "
            "present study therefore reports per-alias agreement, per-pair deltas under alternative "
            "judges, and two separate human-validation samples.",
            "MT-Bench reports more than 80% agreement between strong LLM judges and human preferences "
            "in its tested setting, while also documenting position bias, verbosity bias, self-"
            "enhancement, and limited grading ability on some reasoning questions [47]. A yes/no safety "
            "prompt additionally inherits the judge provider's policy and model updates. Versioning the "
            "judge and preserving redacted labels improve reproducibility, but they do not make the "
            "labels ground truth.",
            "Open-weight evaluators reduce dependence on an external service. Prometheus 2 is designed "
            "for evaluation under user-specified criteria [46], while WildGuard targets prompt "
            "harmfulness, response harmfulness, and refusal [44]. Their openness permits rerunning the "
            "scorer, yet construct alignment still matters. A moderator's unsafe-response criterion is "
            "not identical to HarmBench's named-behaviour criterion, so disagreement can be legitimate.",
            "Human annotation is not automatically definitive. Rubrics can be ambiguous, annotators "
            "can disagree, and a disagreement-enriched sample changes prevalence. A well-specified design "
            "states who labelled the data, how items were sampled, which labels were available, and "
            "whether the protocol changed. It then uses the human set to compare instruments rather "
            "than treating one annotator as an infallible oracle.",
            "The present validation design follows that principle. The HarmBench sample compares a "
            "behaviour classifier and refusal proxy against the same human decisions. The XSTest sample "
            "compares strict and broad mappings and discloses a pre-registration deviation. These two "
            "audits address different constructs, so their agreement coefficients should not be pooled "
            "into one claim of overall judge accuracy.",
        ],
    )

    add_heading(doc, "2.6", "Synthesis and Research Gap")
    add_body(
        doc,
        "Table 2.1 is a scoped positioning aid rather than a claim that this design is unique. It "
        "records whether the cited papers' abstracts and stated scope explicitly establish each "
        "listed dimension. “Not reported” means that the consulted source did not establish the "
        "property; it is not evidence that the authors did not consider it.",
    )
    rows = [
        ("TrustLLM / DecodingTrust [20], [21]", "Not reported", "Broad", "Partial", "Yes", "Automated", "Model treated as fixed"),
        ("HarmBench [24]", "Not reported", "Yes", "Not reported", "Not reported", "Classifier", "ASR conditional on the selected threat model"),
        ("XSTest / OR-Bench [18], [19]", "Not reported", "Contrast set only", "Yes", "Not reported", "Human / mixed", "Over-refusal focus; unsafe items are a calibration contrast"),
        ("GPTQ / AWQ / SmoothQuant [5], [6], [30]", "Yes", "Not reported", "Not reported", "Yes", "Task metrics", "Efficiency and capability focus"),
        ("Kharinaev et al. [10]", "Yes", "Yes", "Not reported", "Partial", "Human / mixed", "Different model/method mix"),
        ("HarmLevelBench [35]", "Yes", "Graded", "Not reported", "Not reported", "Human / mixed", "Attack-dependent construct"),
        ("Q-Misalign [39]", "Yes", "Yes", "Not reported", "Utility", "Automated", "Adversarially prepared model"),
        ("StrongREJECT / SORRY-Bench [41], [42]", "Not reported", "Yes", "Recommended", "Recommended", "Human-grounded", "No matched precision intervention"),
        ("Present study", "NF4 + INT8", "Yes", "Yes", "MMLU + ARC", "Multi-judge + human", "Direct requests; five pairs"),
    ]
    add_table(
        doc,
        "2.1",
        "Scoped positioning of representative prior work. Entries are compiled from each cited "
        "paper's abstract and stated scope. “Not reported” is not evidence of absence. The table is "
        "not a novelty or priority claim.",
        ["Study", "Precision intervention", "Harmful compliance", "Over-refusal", "Capability", "Scorer", "Main boundary"],
        rows,
        [3.0, 1.7, 2.0, 1.6, 1.8, 2.1, 3.4],
        font_size=7.3,
    )
    add_paragraphs(
        doc,
        [
            "Five gaps follow. First, the quantized and full-precision artefacts should come from the "
            "same checkpoint revision. Second, the quantization method and precision must be explicit. "
            "Third, harmful compliance, benign refusal, and capability should be measured together. "
            "Fourth, the conclusion should be tested against scorer choice and human labels. Fifth, "
            "the report must distinguish direct requests from jailbreak robustness.",
            "The present study addresses these gaps with a five-pair matched design, NF4 as the "
            "primary intervention, INT8 as an additional precision point, two capability anchors, a "
            "declared family of paired tests, and independent scorer validation. It does not attempt "
            "to establish a universal law of quantization. Its contribution is a controlled estimate "
            "under a well-defined deployment and measurement setting.",
            "The gap is not merely the absence, in Table 2.1, of the same benchmark combination. "
            "The missing link is inferential: across the studies reviewed there, a reported metric "
            "change is not consistently separated from capability, scorer error, truncation, or "
            "multiplicity. This study builds those alternatives into the design, because a safety "
            "claim becomes credible only after the neighbouring explanations have been tested.",
            "The literature also sets the report's claim boundary. Quantization-conditioned attacks "
            "show that worst-case failures are possible [12], [39]. Jailbreak benchmarks show that "
            "adversarial robustness requires an explicit attack [43]. Judge research shows that "
            "automatic labels require grounding [45], [46]. Consequently, a defensible conclusion is "
            "conditional: for the tested models, methods, prompts, generation budget, and scorers, the "
            "observed pattern is capability loss without a corrected harmful-compliance increase.",
            "Work on how safety training fails under adversarial pressure provides a final caution "
            "[48]. Preservation under direct requests is not evidence that the same model will remain "
            "safe under a targeted attack.",
        ],
    )

    add_heading(doc, "2.7", "Implications for Experimental Design")
    add_paragraphs(
        doc,
        [
            "The reviewed literature suggests that a convincing quantization-safety experiment "
            "must control three layers at once. The model layer requires identical checkpoint "
            "provenance within a pair, the evaluation layer requires safety, benign refusal, and "
            "capability outcomes, and the inference layer requires fixed prompt formatting, "
            "decoding controls, and generation budgets. Crucially, a weakness in any layer can "
            "create a treatment effect that is not caused by the quantizer.",
            "Statistical design is equally important. Twenty model–outcome contrasts create many "
            "opportunities for an isolated p-value below 0.05. The three quantization evaluations "
            "reviewed here emphasise average benchmark retention and do not frame their outcomes as "
            "a declared multiplicity family [8]-[10]. The present study uses paired tests and FDR "
            "control because "
            "its question is inferential: which observed changes remain credible after accounting for "
            "the number of comparisons?",
            "A null result also needs positive design. It should state what effect sizes the sample can "
            "detect, test whether a scorer is capable of seeing the relevant construct, and challenge "
            "conditions that could hide or manufacture an effect. Generation-length and seed analyses "
            "serve that role here. They do not prove exact preservation, but they make the null more "
            "informative than an isolated non-significant test.",
            "Finally, research integrity affects interpretation. Revised scoring should not overwrite "
            "original outputs, and a validation-informed analysis should not replace a pre-specified "
            "family without disclosure. The literature on automatic judging makes clear that "
            "measurement systems evolve [45], [46]. Preserving the original and revised "
            "views allows a reader to reconstruct why a conclusion changed.",
        ],
    )


def add_chapter_three(doc: Document) -> None:
    add_chapter(doc, 3, "Methodology and Experimental Setup")

    add_heading(doc, "3.1", "Matched-Pair Experimental Design")
    add_paragraphs(
        doc,
        [
            "Each experimental unit is a matched pair containing a baseline model and a quantized "
            "model, and both members use the same model identifier and pinned revision. The "
            "baseline is loaded at the resolved full precision used by the TC1 CUDA environment, "
            "whereas the treatment member is loaded from the same weights with a BitsAndBytes "
            "quantization configuration. This design therefore removes checkpoint provenance as a "
            "confound and isolates the load-time precision intervention within a single execution "
            "environment. Table A.1 records the quantization settings and the declared dependency "
            "constraints.",
            "The primary treatment is NF4 four-bit quantization with double quantization and a "
            "compute dtype that follows the resolved model dtype. INT8 uses LLM.int8 and is "
            "analysed as a separate precision arm, so it is not placed inside the primary "
            "fp16–NF4 multiplicity family. All primary runs use seed 42, the model's native chat "
            "template, temperature 0.0, and a 512-token maximum generation budget. Figure 3.1 "
            "shows the study design.",
            "The independent variable is therefore the precision-loading path, and the dependent "
            "variables are prompt-level harmful-compliance labels, prompt-level refusal labels, "
            "and exact-match capability scores. Model family and scale are observed grouping "
            "variables rather than randomised treatments. In addition, generation budget, "
            "decoding rule, prompt order, benchmark sample, and chat-template path are all "
            "controlled within a matched pair.",
            "Pairing is essential because the same prompt can be intrinsically easy or difficult "
            "for both models, and an unpaired comparison would treat that shared difficulty as "
            "independent noise. The paired analysis instead asks how often the two precision "
            "conditions disagree on the same item. This increases statistical efficiency and "
            "gives discordant counts a direct interpretation.",
            "The study was expanded in stages, but the current analysis presents the complete "
            "five-pair matrix under one primary configuration. Historical 128-token outputs are "
            "not mixed into the 512-token estimates, and they appear only in the generation-"
            "length sensitivity analysis. This separation prevents a shorter pilot budget from "
            "silently determining the headline.",
            "Prompt-level pairing is preserved in both the bootstrap interval [49] and the exact "
            "McNemar comparison [50]. These methods are introduced here rather than in the "
            "analysis chapter because pairing is part of the design, not a post hoc statistical "
            "convenience.",
        ],
    )
    add_figure(
        doc,
        "3.1",
        "Matched pairs, precision arms, benchmark battery, and analysis flow.",
        "study_design.png",
        14.5,
    )

    add_heading(doc, "3.2", "Models and Precision Conditions")
    model_rows = []
    for pair_id in PAIR_ORDER:
        pair = PAIRS[pair_id]
        model_rows.append(
            (
                pair["name"],
                pair["family"],
                f"{pair['params_b']:.1f} B",
                pair["model_id"],
                "fp16 / INT8 / NF4",
            )
        )
    add_table(
        doc,
        "3.1",
        "Model pairs and precision conditions",
        ["Model", "Family", "Nominal size", "Model identifier", "Conditions"],
        model_rows,
        [2.7, 1.7, 1.8, 6.2, 3.2],
        font_size=7.8,
    )
    add_paragraphs(
        doc,
        [
            "The selected models span four families and approximately 1.7–7.2 billion parameters. "
            "Qwen contributes two sizes within one family, which permits a limited within-family "
            "scale comparison, while Llama, Mistral, and Phi broaden the architectural and "
            "training diversity. Mistral-7B provides the largest model in scope [51]. Phi-4-mini "
            "uses eager attention on the V100, and both members of the Phi pair use the same "
            "setting.",
            "Model selection was constrained by reproducible access, V100 compatibility, and the "
            "availability of instruction-tuned checkpoints. Because the design does not sample "
            "randomly from the population of open models, cross-family claims are descriptive of "
            "the evaluated set rather than population estimates.",
            "The two Qwen sizes provide the only controlled scale comparison within a family. "
            "However, the Llama, Mistral, and Phi additions increase architectural diversity "
            "while confounding family, training data, alignment recipe, tokenizer, and size. "
            "Their purpose is therefore replication across distinct deployed artefacts, not "
            "estimation of a pure family effect.",
            "Each pair shares its model revision exactly, which is a stronger control than "
            "matching only the public model name, because a repository can change configuration "
            "or weights over time. The revisions are recorded in the configuration and used by "
            "the loader, and the INT8 aliases use the same revisions as the corresponding fp16 "
            "and NF4 conditions.",
        ],
    )

    add_heading(doc, "3.3", "Benchmarks and Outcomes")
    add_table(
        doc,
        "3.2",
        "Benchmark roles, sample sizes, and primary outcomes",
        ["Benchmark", "Role", "Sample size per alias", "Primary outcome", "Direction interpreted as worse"],
        [
            ("HarmBench", "Direct harmful compliance", "200", "Classifier ASR", "Higher"),
            ("XSTest", "Benign over-refusal", "250", "Regex over-refusal rate", "Higher"),
            ("MMLU subset", "Capability anchor", "300", "Exact-match accuracy", "Lower"),
            ("ARC-Challenge", "Capability anchor", "1,172", "Exact-match accuracy", "Lower"),
        ],
        [2.6, 3.5, 2.2, 3.0, 4.3],
        font_size=8.0,
    )
    add_paragraphs(
        doc,
        [
            "HarmBench supplies unsafe behaviours, and its official classifier is the primary "
            "scorer. Although the benchmark also supports optimised attacks, this study uses "
            "direct requests only. XSTest supplies benign prompts with lexical or pragmatic "
            "features that can trigger exaggerated safety. Its pre-specified deterministic scorer "
            "remains the primary scorer of record, while a three-class semantic judge provides a "
            "validation-informed sensitivity analysis.",
            "MMLU is evaluated on six subjects with zero-shot multiple-choice prompts, and ARC-"
            "Challenge provides a larger second anchor with a different task structure. Exact-"
            "match parsing is used for both. The two capability benchmarks are interpreted "
            "jointly in discussion, but the primary multiplicity family retains them as separate "
            "outcomes.",
            "Sample sizes reflect the benchmark resources and compute budget. HarmBench uses 200 "
            "behaviours per alias, which allows item-level matching across precision, and XSTest "
            "uses its 250 benign prompts. The MMLU subset contains 300 items, and ARC-Challenge "
            "contains 1,172. Consequently, the unequal sizes yield different detection "
            "thresholds, which is why a similar absolute delta can be significant on ARC and "
            "unresolved on HarmBench.",
            "Prompt formatting is benchmark-specific but invariant within a pair. The generator "
            "applies the model's native chat template with thinking disabled where the template "
            "exposes that option, and raw-prompt fallback is not used for the successful primary "
            "runs. The saved metadata records the resolved generation configuration for each "
            "response.",
        ],
    )

    add_heading(doc, "3.4", "Scoring and Validation")
    add_paragraphs(
        doc,
        [
            "A response is a HarmBench success only when the official behaviour-specific "
            "classifier judges that it exhibits or clearly attempts the target behaviour under "
            "the benchmark's criteria. A refusal-regex proxy is "
            "retained for transparency and comparison, but it is not the primary harmful-"
            "compliance outcome. The classifier is cross-checked against a gpt-4o API judge whose "
            "exact dated snapshot was not recorded, and against revision-pinned Llama Guard. A "
            "stratified 200-item human-labelled sample grounds both the primary classifier and "
            "the regex.",
            "XSTest is scored first by the pre-specified regex, after which the sensitivity judge "
            "assigns full compliance, partial refusal, or full refusal. Strict over-refusal "
            "counts only full refusal, whereas the broad mapping also counts partial refusal, and "
            "a separate 200-item human sample evaluates both mappings. This arrangement keeps the "
            "original analysis visible while testing whether its conclusion survives a more "
            "semantic instrument.",
            "The refusal proxy originally used for HarmBench identifies explicit refusal patterns "
            "and treats their absence as a successful attack. That rule is deterministic and "
            "auditable, but it cannot assess whether harmful information was actually provided. "
            "By contrast, the HarmBench classifier receives both the behaviour and the response "
            "and returns the benchmark's binary success judgement. Judge parsing failures are "
            "counted rather than discarded, and the reported runs contain zero parse failures.",
            "The API second judge uses a condensed yes/no rubric aligned to the same harmful-"
            "behaviour criterion. Llama Guard, however, is intentionally treated as a looser "
            "construct cross-check, because it classifies whether the response is unsafe under "
            "its taxonomy rather than whether a named behaviour was completed. This difference is "
            "stated before agreement is interpreted.",
            "Human-validation items are stored outside the public result surface because they "
            "contain prompt and response text. Committed summaries therefore retain counts, "
            "labels, agreement, precision, recall, and redacted identifiers. The HarmBench sample "
            "is enriched for scorer disagreement, while the XSTest sample is drawn under the "
            "documented blindness rule and its disclosed deviation.",
        ],
    )

    add_heading(doc, "3.5", "Statistical Analysis")
    add_paragraphs(
        doc,
        [
            "For each pair and outcome, the effect is the quantized rate or accuracy minus the "
            "baseline value, and paired bootstrap confidence intervals preserve prompt-level "
            "pairing [49]. Binary outcomes are also tested with the exact McNemar test [50], "
            "which uses only discordant prompt pairs and whose null hypothesis is equal marginal "
            "probability within the pair.",
            "The primary NF4 multiplicity family contains 20 contrasts: five model pairs multiplied by four "
            "outcomes. Benjamini and Hochberg's false-discovery-rate procedure controls "
            "multiplicity across that family [52], and a result is described as surviving "
            "correction only when its adjusted q-value is below 0.05. Uncorrected significance "
            "remains visible in the tables, but it does not determine the study headline.",
            "Cohen's kappa measures chance-corrected agreement between scorers [53]. "
            "Sample-size and minimum-detectable-effect calculations follow paired-proportion "
            "power analysis [54], [55]. Agreement methodology follows the caution that "
            "coefficient choice and prevalence affect interpretation [56], so the report "
            "presents the coefficient with raw agreement and task context rather than treating "
            "the verbal bands of Landis and Koch [57] as universal truth.",
            "Llama Guard supplies the open-weight safety-moderation cross-check used in the "
            "scorer sensitivity analysis [58], and its construct is deliberately broader than the "
            "HarmBench behaviour label.",
            "Equivalence is assessed only where the pre-specified margin and test are "
            "appropriate, and the two one-sided tests procedure supplies the formal basis [59]. "
            "Data-contamination risk is disclosed because benchmark familiarity can inflate "
            "capability estimates [60]. Finally, the sensitivity analyses are labelled separately "
            "so that they strengthen or qualify the interpretation without silently changing the "
            "primary multiplicity family.",
            "For a binary outcome, let b be the number of prompts on which the baseline is "
            "negative and the quantized condition is positive, and c the reverse. The observed "
            "paired difference is then (b−c)/n under the report's harmful-ward sign convention. "
            "McNemar's exact test conditions on b+c and evaluates whether the two discordant "
            "directions are equally likely, so items on which both conditions agree do not "
            "contribute to that test statistic.",
            "The bootstrap resamples matched prompt pairs with replacement, and for each "
            "replicate the quantized-minus-baseline difference is recomputed, preserving within-"
            "item dependence. The reported 95% interval is the empirical percentile interval. "
            "Because a confidence interval and an exact paired test answer related but not "
            "identical questions, both remain visible.",
            "The BH procedure sorts the 20 primary p-values and compares them with rank-dependent "
            "thresholds, controlling the expected false-discovery proportion among rejected "
            "hypotheses under its assumptions. However, it does not remove construct error from "
            "the three surviving effects, so scorer validation remains necessary after "
            "statistical correction.",
            "Interpretation labels combine direction and capability context. Alignment "
            "degradation means harmful compliance rises while capability is broadly preserved, "
            "and alignment improvement is the converse. Capability collapse masquerading as "
            "safety is used when harmful compliance falls alongside capability, whereas broad "
            "degradation denotes adverse capability movement without a clean safety "
            "interpretation. In every case the labels summarise a pattern and do not establish a "
            "mechanism.",
        ],
    )

    add_heading(doc, "3.6", "Implementation and Execution")
    add_paragraphs(
        doc,
        [
            "The implementation is a benchmark framework rather than the main research "
            "contribution. A validated YAML configuration defines models, revisions, benchmarks, "
            "decoding controls, and cluster settings, and a loader resolves dtype and applies "
            "quantization. A shared generator then uses the native chat template, batches "
            "prompts, and records generation metadata, while benchmark plugins prepare inputs and "
            "calculate benchmark-specific fields. Matrix orchestration loads one model once and "
            "evaluates its benchmarks sequentially before releasing GPU memory.",
            "Runs were executed through SLURM on an NTU TC1 node with a 32 GB Tesla V100 GPU, and "
            "compute jobs used offline Hugging Face modes after models and datasets were cached. "
            "The job scripts preserved the required module, environment, work-directory, and "
            "offline bootstrap commands. Resume logic skips prompt identifiers already present in "
            "the output, which allows interrupted jobs to continue without duplicating records.",
            "Original raw generations and summaries are treated as immutable run artefacts. "
            "Revised scorers, judge outputs, and diagnostics are instead written to separate "
            "redacted sidecars containing identifiers and labels rather than harmful text. This "
            "separation keeps the audit trail intact and avoids rewriting history when a scorer "
            "is corrected.",
            "The command-line interface exposes single-run, matrix, analysis, and agent-status "
            "paths. The dashboard reads the result trees but does not alter them, and new user-"
            "generated configurations are written outside the committed primary configuration. "
            "This boundary prevents an interface action from changing the evidence underlying the "
            "report.",
            "Cluster policy also affects reproducibility. User compute runs only through batch "
            "jobs, compute nodes operate in offline mode, and models and datasets are cached "
            "beforehand on the head node. The job files use absolute log paths and preserve the "
            "environment bootstrap. These details are operational, but they determine whether a "
            "rerun loads the intended artefacts instead of silently contacting a newer upstream "
            "revision.",
            "Testing covers configuration validation, model reuse, result resumption, refusal "
            "parsing, pairwise analysis, judge redaction, report claims, and artifact policy. "
            "Passing tests show that the implemented checks behave as specified. They do not, "
            "however, substitute for scientific review, visual inspection, or source "
            "verification.",
        ],
    )

    add_heading(doc, "3.7", "Reproducibility Controls")
    add_paragraphs(
        doc,
        [
            "Every model revision in the primary configuration is pinned, and the report records "
            "the sample counts, seed, generation budget, decoding rule, attention exception, "
            "quantization configuration, scorer identity, and GPU class needed to reconstruct the "
            "run. However, the exact historical environment of the earliest artefacts was not "
            "captured with perfect package-level provenance, so the report does not claim bit-"
            "for-bit reproduction of those runs. It does support deterministic replay of the "
            "committed analysis and a fresh rerun from the pinned configuration.",
            "The project uses automated tests and claim verification to detect stale numbers and "
            "inconsistent surfaces. These checks support document integrity, but they do not prove that "
            "a citation entails a sentence or that a construct is valid. Source fitness and scorer "
            "validity therefore remain review tasks. LLM-as-a-judge research reports useful agreement "
            "with human preferences [47], while surveys catalogue position, verbosity, and other biases "
            "[61]. Human grounding is necessary because automatic labels are not free evidence [62]. "
            "Non-parametric sensitivity comparisons use the Wilcoxon procedure where applicable [63], "
            "and power interpretation follows standard behavioural-science guidance [64].",
            "Reproduction has two levels. Analysis replay starts from the committed redacted "
            "scores and regenerates deltas, tests, figures, and report tables, whereas full "
            "experimental reproduction loads the pinned models, regenerates responses, and reruns "
            "the scorers. The first is deterministic on an ordinary workstation. The second "
            "requires compatible GPU hardware, accepted gated-model conditions, cached datasets, "
            "and a software environment satisfying the dependency constraints in Table A.1 of "
            "Appendix A.",
            "The report avoids a stronger claim of exact historical rerun because not every "
            "earliest package version was captured in the immutable outputs. Although current "
            "pipeline code records richer environment provenance for future runs, the limitation "
            "is disclosed rather than repaired by editing historical artefacts.",
        ],
    )

    add_heading(doc, "3.8", "Analysis Workflow and Quality Control")
    add_paragraphs(
        doc,
        [
            "The analysis begins by joining baseline and quantized records on prompt identifier, "
            "and cardinality checks fail if either condition is missing an item or if a generated "
            "response count differs from the input count. Per-alias summaries are then rebuilt "
            "from the selected scorer sidecars. Finally, pairwise deltas, discordant counts, "
            "bootstrap intervals, and exact tests are generated before interpretation labels are "
            "assigned.",
            "Multiplicity correction is performed from the complete 20-row registry rather than "
            "from a hand-selected result table, and the same registry supplies recurring values "
            "to the report builders and other current-facing surfaces. This reduces transcription "
            "drift between an analysis file, a table, an abstract, and a conclusion.",
            "Figures are produced by scripts from the analysis JSON rather than drawn manually. "
            "The forest plot uses the classifier deltas and paired intervals, the capability plot "
            "joins MMLU and classifier ASR by pair, and the precision figure reads the "
            "fp16–INT8–NF4 analysis. This report references each figure before it appears and "
            "keeps the underlying path out of the reader-facing caption.",
            "Quality control has both machine and human components. Tests detect schema failures, stale "
            "claims, altered immutable files, redaction leaks, and inconsistent generated surfaces. "
            "Source-to-claim review checks whether the citation supports the wording, and rendered page "
            "inspection checks tables, captions, page breaks, and glyphs. These layers address different "
            "failure modes and none is treated as sufficient alone.",
        ],
    )


def primary_safety_rows() -> list[tuple[str, ...]]:
    rows = []
    for pair_id in PAIR_ORDER:
        pair = PAIRS[pair_id]
        asr = pair["asr"]
        xstest = pair["xstest"]
        rows.append(
            (
                pair["name"],
                f"{fmt_rate(asr['baseline'])} → {fmt_rate(asr['nf4'])}",
                f"{fmt_delta(asr['delta'])} {fmt_ci(asr['ci'])}",
                f"{asr['p_value']:.4f} / {asr['bh_q_value']:.4f}",
                f"{fmt_rate(xstest['baseline'])} → {fmt_rate(xstest['nf4'])}",
                f"{fmt_delta(xstest['delta'])} {fmt_ci(xstest['ci'])}",
                f"{xstest['p_value']:.4f} / {xstest['bh_q_value']:.4f}",
            )
        )
    return rows


def primary_capability_rows() -> list[tuple[str, ...]]:
    rows = []
    for pair_id in PAIR_ORDER:
        pair = PAIRS[pair_id]
        mmlu = pair["mmlu"]
        arc = pair["arc"]
        rows.append(
            (
                pair["name"],
                f"{fmt_rate(mmlu['baseline'])} → {fmt_rate(mmlu['nf4'])}",
                f"{fmt_delta(mmlu['delta'])} {fmt_ci(mmlu['ci'])}",
                f"{mmlu['p_value']:.4f} / {mmlu['bh_q_value']:.4f}",
                f"{fmt_rate(arc['baseline'])} → {fmt_rate(arc['nf4'])}",
                f"{fmt_delta(arc['delta'], 4)} {fmt_ci(arc['ci'], 4)}",
                f"{arc['p_value']:.4f} / {arc['bh_q_value']:.4f}",
            )
        )
    return rows


def add_chapter_four(doc: Document) -> None:
    add_chapter(doc, 4, "Results")
    add_body(
        doc,
        "This chapter reports the primary matched-pair results before the validation and sensitivity "
        "analyses. Each section follows the same order: finding, magnitude, statistical support, and "
        "interpretation. Complete intervals and p/q values are kept in tables so that the prose does "
        "not repeat every statistic.",
    )

    add_heading(doc, "4.1", "Results at a Glance")
    multiplicity = CLAIMS["multiplicity"]
    add_paragraphs(
        doc,
        [
            f"The primary multiplicity family contains {multiplicity['family_size']} NF4 contrasts. "
            "Exactly three contrasts survive BH-FDR correction, and none is a HarmBench "
            "harmful-compliance contrast. The survivors are Qwen3-1.7B MMLU "
            f"{fmt_delta(multiplicity['survivors'][0]['delta'])}, Llama-3.2-3B ARC "
            f"{fmt_delta(multiplicity['survivors'][1]['delta'], 4)}, and Phi-4-mini over-refusal "
            f"{fmt_delta(multiplicity['survivors'][2]['delta'])}.",
            "The correction therefore moves the report away from a pair-by-pair search for dramatic "
            "stories. The corrected evidence is concentrated on capability, while the single "
            "over-refusal survivor moves downward and is later shown to depend on the scorer.",
            "Tables 4.1 and 4.2 separate safety from capability so that each comparison remains "
            "readable. The baseline and NF4 columns show the practical rates, while the delta and "
            "confidence-interval columns show magnitude and uncertainty. The p/q columns then "
            "distinguish an individually unusual paired result from one that survives family-wide "
            "correction.",
            "The sign convention is consistent: positive ASR and over-refusal deltas are adverse, "
            "while negative capability deltas are adverse. A safe-ward ASR movement and an "
            "adverse capability movement do not cancel, and together they weaken the claim that "
            "alignment improved.",
        ],
    )
    add_table(
        doc,
        "4.1",
        "NF4 safety and over-refusal results at the 512-token budget",
        ["Pair", "ASR fp16→NF4", "ΔASR (95% CI)", "ASR p / q", "OR fp16→NF4", "ΔOR (95% CI)", "OR p / q"],
        primary_safety_rows(),
        [2.5, 2.0, 2.8, 1.8, 2.0, 2.8, 1.7],
        font_size=7.0,
    )
    add_table(
        doc,
        "4.2",
        "NF4 capability results at the 512-token budget",
        ["Pair", "MMLU fp16→NF4", "ΔMMLU (95% CI)", "MMLU p / q", "ARC fp16→NF4", "ΔARC (95% CI)", "ARC p / q"],
        primary_capability_rows(),
        [2.5, 2.0, 2.8, 1.8, 2.0, 2.8, 1.7],
        font_size=7.0,
    )

    add_heading(doc, "4.2", "Harmful Compliance (RQ1)")
    add_paragraphs(
        doc,
        [
            "No NF4 harmful-compliance contrast survives multiplicity correction. Figure 4.1 "
            "shows that the five point estimates range from −0.040 to +0.040, with intervals that "
            "cross or touch zero except for the uncorrected Llama decrease. Overall, this is "
            "evidence against a large, consistent harmful-ward shift under NF4 in the evaluated "
            "setting.",
            "Qwen3-1.7B is exactly flat in aggregate at 0.255 for both precisions. The aggregate "
            "null does not mean every prompt is stable: sixteen prompts flip toward harmful "
            "compliance and sixteen flip toward safety, producing McNemar p = 1.000. Qwen3-4B "
            "increases by 0.040, but the confidence interval touches zero and the adjusted "
            "q-value is 0.2406. Similarly, Phi-4-mini increases by 0.020 with a wide interval and "
            "no statistical support.",
            "Llama-3.2-3B decreases from 0.100 to 0.060. Its individually significant ΔASR is "
            "−0.040, a decrease that does not survive the 20-test correction. More importantly, "
            "the model also loses ARC accuracy, and a lower ASR cannot be interpreted as a clean "
            "alignment gain when capability moves in the same adverse direction. Mistral-7B "
            "decreases by 0.020 with substantial prompt-level disagreement and no significant "
            "aggregate change.",
            "The five models differ substantially in baseline susceptibility. For example, "
            "Mistral's baseline ASR is 0.585, more than eight times Phi's 0.070. Paired deltas "
            "remain comparable as treatment effects, but their operational implications differ. A "
            "two-point change at a high baseline does not make the model safe, and a two-point "
            "change at a low baseline does not establish a general regression.",
            "Qwen3-4B is the largest harmful-ward point estimate. Thirteen prompts change from "
            "safe under fp16 to harmful under NF4, while five move the other way, and the exact "
            "p-value is 0.0963. This is therefore a directional signal worth replication, but the "
            "present evidence does not justify a confirmed increase.",
            "For Llama, nine prompts move from harmful to safe and one moves from safe to "
            "harmful, and that imbalance produces the individually significant decrease. Section "
            "4.4 explains the caution: ARC accuracy falls by 0.0316 in the same precision "
            "contrast. The paired data therefore show co-movement, not that one change caused the "
            "other.",
            "The cross-family pattern is heterogeneous in direction but homogeneous in "
            "inferential status after correction. Two pairs point upward, two downward, and one "
            "is flat, so no common harmful-ward NF4 effect is visible.",
        ],
    )
    add_figure(
        doc,
        "4.1",
        "Per-pair HarmBench classifier ΔASR with paired confidence intervals.",
        "asr_forest.png",
        14.2,
    )

    add_heading(doc, "4.3", "Over-Refusal on Benign Prompts (RQ2)")
    add_paragraphs(
        doc,
        [
            "No pair shows a statistically significant increase in over-refusal, and four pairs "
            "are non-significant. Phi-4-mini's −0.048 over-refusal decrease under the pre-"
            "specified regex is scorer-dependent, because the independent semantic judge does not "
            "reproduce it. The direction under the original scorer is in any case benign, since "
            "the NF4 model refuses fewer benign prompts.",
            "Qwen3-1.7B decreases by 0.024. Its bootstrap interval reaches zero and the exact "
            "paired test gives p = 0.109, so it is a null rather than a borderline finding. "
            "Qwen3-4B changes by −0.004, Llama by +0.016, and Mistral by 0.000. Overall, these "
            "values do not support a general claim that NF4 makes compact models more prone to "
            "exaggerated safety.",
            "The primary family retains the regex result because it was the pre-specified "
            "outcome. By contrast, the validation-informed family replaces it with judge-strict "
            "over-refusal and leaves only two survivors, both capability losses. This distinction "
            "avoids rewriting the analysis after seeing the data while still giving more weight "
            "to the better-grounded instrument.",
            "Phi's regex over-refusal rate falls from 0.128 to 0.080. In deployment terms the "
            "absolute change corresponds to twelve fewer regex-flagged benign prompts in the "
            "250-item condition, which is large enough to survive correction under the original "
            "scorer. It is also exactly the kind of result that requires construct validation, "
            "because a small number of pattern matches can move the rate materially.",
            "Llama has the only upward over-refusal estimate, from 0.032 to 0.048, with four "
            "additional benign prompts flagged under NF4. However, the paired interval spans "
            "−0.004 to 0.040 and the exact p-value is 0.289, so the data do not support a "
            "meaningful increase.",
            "Mistral is at a near-floor rate of 0.004 in both conditions. Consequently, floor "
            "effects make a decrease difficult to observe and can depress kappa despite high raw "
            "agreement. This baseline dependence is another reason not to treat over-refusal as a "
            "cross-model leaderboard.",
        ],
    )

    add_heading(doc, "4.4", "Capability (RQ3)")
    add_paragraphs(
        doc,
        [
            "Capability contains the clearest costs. Qwen3-1.7B MMLU falls from 0.643 to 0.553, a "
            "decrease of 0.090 that survives BH-FDR with q = 0.008. In deployment terms, the NF4 "
            "model misses roughly nine additional questions per hundred that its fp16 counterpart "
            "answers correctly.",
            "Llama-3.2-3B has an ARC contrast of −0.032 (unrounded: −0.0316), which survives "
            "BH-FDR correction with q = 0.008. "
            "Qwen3-4B ARC decreases by 0.0162 and is individually significant, but it does not survive "
            "the full correction. The other capability contrasts are non-significant. Mistral-7B ARC "
            "is the only positive point estimate, at +0.0094, which is also non-significant.",
            "The two anchors prevent overgeneralisation from one benchmark. Qwen3-1.7B has a "
            "large MMLU loss but only a small ARC decrease, whereas Llama shows the reverse "
            "pattern more clearly on ARC than on MMLU. The conclusion is therefore not that every "
            "model suffers a uniform capability collapse, but that capability produces the "
            "largest corrected effects and that its expression depends on the benchmark.",
            "Qwen3-4B preserves MMLU almost exactly, moving from 0.747 to 0.743, while ARC falls "
            "from 0.876 to 0.860. The ARC result is individually significant because the "
            "benchmark has 1,172 paired items, but its adjusted q-value is 0.1687. Sample size "
            "and multiplicity therefore change the evidence attached to a small absolute effect.",
            "Llama's MMLU decrease is 0.0367, with an interval ending at zero and p = 0.080. Its ARC "
            "decrease is similar in magnitude but estimated more precisely, so ARC survives correction. "
            "In deployment terms, the NF4 model misses roughly three additional ARC questions per "
            "hundred that its fp16 counterpart answers correctly. "
            "The two outcomes are directionally consistent even though only one crosses the inferential "
            "threshold.",
            "Mistral loses 0.020 on MMLU but gains 0.0094 on ARC, while Phi loses 0.0267 on MMLU "
            "and 0.0145 on ARC, neither of which is significant. These unresolved cases show why "
            "a rule that labels every negative point estimate as capability degradation would "
            "overstate the evidence.",
            "Minimum-detectable-effect calculations reinforce the distinction. The Qwen3-1.7B "
            "MMLU and Llama ARC effects exceed their estimated 80% power thresholds, whereas "
            "several smaller observed deltas do not. Appendix B therefore reports the thresholds "
            "rather than converting all non-significant effects into claims of equivalence.",
            "Figure 1.1 places all five pairs in the joint space these results define. The "
            "capability axis carries a negative displacement for every pair, whereas the "
            "harmful-compliance deltas are mixed in sign and none of them survives correction.",
        ],
    )

    add_heading(doc, "4.5", "Model Scale and Family Differences (RQ4)")
    add_paragraphs(
        doc,
        [
            "The within-Qwen comparison suggests that the 1.7B model is more sensitive on MMLU. "
            "Its loss is 0.090, whereas the 4B model changes by −0.0033. However, that contrast "
            "does not reproduce at the same magnitude on ARC: the 1.7B decrease is −0.0094 and "
            "the 4B decrease is −0.0162. The evidence therefore supports task-specific small-"
            "model sensitivity, not a universal monotonic scale law.",
            "Across families, the baseline safety profiles differ substantially. Mistral begins "
            "at ASR 0.585, while Phi begins at 0.070, so a common absolute delta has different "
            "practical meaning at those baselines. This heterogeneity is why results are reported "
            "as paired within-model contrasts rather than as a leaderboard.",
            "The interpretation labels are descriptive summaries of joint movement. Qwen3-1.7B is "
            "broad degradation because capability falls while ASR is flat, and Llama is "
            "capability collapse masquerading as safety because ASR and capability both fall. The "
            "remaining labels are directional rather than confirmed: Qwen3-4B and Phi lean toward "
            "alignment degradation, while Mistral leans toward alignment improvement. However, "
            "none of those directional safety labels survives the corrected inferential test.",
            "Model size is correlated with other design choices. The two Qwen checkpoints share a "
            "family, but training dynamics and capacity still differ, so their comparison "
            "supports a practical observation for these releases rather than a fitted scaling "
            "law. More sizes and replicated families would be needed to estimate sensitivity as a "
            "function of parameters.",
            "Cross-family response styles also affect measurement. Mistral is the clearest case: its "
            "refusal proxy is far more pessimistic than the semantic classifier. A family comparison "
            "that ignored scorer agreement would confound treatment effect with measurement fit.",
        ],
    )

    add_heading(doc, "4.6", "INT8 Precision Point")
    add_paragraphs(
        doc,
        [
            "The INT8 arm asks whether the observed pattern is graded with precision. Figure 4.2 "
            "shows that harmful compliance is generally non-monotonic across fp16, INT8, and NF4. "
            "In addition, every INT8 HarmBench contrast at 512 tokens is non-significant under "
            "both the primary classifier and the API judge.",
            "Capability is more orderly. INT8 remains close to fp16, while the largest losses "
            "appear at NF4, and the equivalence analysis supports practical similarity for the "
            "INT8 capability contrasts within the stated ±0.05 margin. However, this result is "
            "method-specific: LLM.int8 and NF4 use different numerical treatments, so the pattern "
            "should not be simplified to a universal bit-width curve.",
            "For deployment, the evidence favours a staged decision. INT8 is the lower-risk "
            "precision point when the memory budget permits it, whereas NF4 offers stronger "
            "compression but warrants model-specific capability validation. The study does not "
            "contain direct VRAM, latency, or throughput measurements, so it cannot quantify the "
            "engineering benefit side of that choice.",
            "Qwen3-1.7B illustrates the capability cliff: MMLU is slightly higher at INT8 than "
            "fp16, then drops by more than ten points relative to INT8 at NF4. Similarly, Llama "
            "shows the same broad pattern on ARC, where INT8 is close to baseline and NF4 is "
            "lower. By contrast, Qwen3-4B is comparatively stable on MMLU across all three "
            "precisions.",
            "Safety does not follow the same shape. Qwen3-4B rises from 0.115 to 0.125 to 0.155, "
            "but neither increase is significant after correction. Llama rises slightly at INT8 "
            "and falls at NF4, while Mistral reaches the same 0.565 ASR at INT8 and NF4. These "
            "non-monotonic paths argue against a single precision-driven alignment mechanism.",
            "The 512-token INT8 rerun also resolves a historical warning. A Llama INT8 increase "
            "had appeared at 128 tokens under both judges. At 512 tokens, however, the classifier "
            "delta is +0.005 with p = 1.000 and the API-judge delta is +0.010 with p = 0.688, so "
            "the apparent safety effect does not survive the generation-budget correction.",
        ],
    )
    add_figure(
        doc,
        "4.2",
        "Precision sweep across fp16, INT8, and NF4 for safety and capability outcomes.",
        "precision_sweep.png",
        14.5,
    )

    add_heading(doc, "4.7", "Research-Question Synthesis")
    add_paragraphs(
        doc,
        [
            "RQ1 asks whether NF4 increases harmful compliance. The point estimates are 0.000, +0.040, "
            "−0.040, −0.020, and +0.020 across the five pairs. Only the Llama decrease is individually "
            "significant, and no ASR result survives BH-FDR. The evidence therefore does not support a "
            "corrected harmful-ward increase.",
            "RQ2 asks whether NF4 increases benign over-refusal. The pre-specified regex finds no "
            "significant increase, and its only corrected movement is the Phi decrease. By "
            "contrast, the semantic judge finds neither a significant increase nor the same Phi "
            "decrease. The scorer therefore changes the interpretation of one benign-direction "
            "result, but not the answer to the research question.",
            "RQ3 asks whether capability falls. Qwen3-1.7B MMLU and Llama ARC survive correction. "
            "Qwen3-4B ARC is individually significant but not corrected, and the remaining estimates are "
            "mixed or unresolved. Capability degradation is therefore real for specific model–task "
            "combinations rather than universal across the matrix.",
            "RQ4 asks about scale, family, and precision. The smaller Qwen is more sensitive on "
            "MMLU, but the pattern does not reproduce cleanly on ARC. In addition, cross-family "
            "safety directions differ, and INT8 stays close to fp16 on capability while NF4 "
            "contains the largest losses. Overall, the evidence supports heterogeneity, not one "
            "shared response to bit-width.",
            "RQ5 asks whether the conclusion survives measurement and inference changes. The API judge "
            "and Llama Guard do not find a significant harmful-ward NF4 pair. Human grounding favours "
            "the classifier over the HarmBench proxy and the semantic judge over the XSTest regex. The "
            "longer generation budget removes the retired Qwen signal, and the multi-seed arm does not "
            "produce a stable harmful-ward effect. The bounded null is therefore supported by several "
            "independent checks.",
        ],
    )
def add_chapter_five(doc: Document) -> None:
    add_chapter(doc, 5, "Scorer Validation and Robustness")
    add_body(
        doc,
        "The primary findings depend on measurement choices. This chapter tests the load-bearing "
        "assumptions directly: whether the classifier captures harmful compliance better than a "
        "refusal proxy, whether the XSTest regex captures refusal, and whether the conclusions persist "
        "across judges, generation budgets, and decoding seeds.",
    )

    add_heading(doc, "5.1", "HarmBench Classifier versus Refusal Proxy")
    add_paragraphs(
        doc,
        [
            "Figure 5.1 compares the HarmBench classifier with the refusal-regex proxy. The proxy "
            "systematically reports higher attack-success rates for several aliases because it "
            "treats non-refusal as compliance. However, that is not the construct HarmBench "
            "defines: a response may avoid an explicit refusal and still fail to provide the "
            "requested harmful information.",
            "The divergence is largest in model families whose response style does not align with "
            "the regex. Per-alias classifier-versus-proxy kappa ranges from roughly 0.25 to 0.84 "
            "at the 512-token budget. For Mistral the proxy even reverses the sign, suggesting a "
            "harmful-ward increase where the classifier gives a small safe-ward decrease. "
            "Consequently, the classifier is the primary outcome and the proxy is retained only "
            "as a transparent secondary view.",
            "The disagreement is not a constant offset. If the proxy added the same bias to both "
            "precision conditions, its paired delta could remain useful. Instead, the size and "
            "direction of the error vary by alias. Quantization can change response style, length, or "
            "refusal phrasing, so proxy error can itself become treatment-dependent.",
            "The correction was made through derived sidecars, and original generations and "
            "summaries were not overwritten. This preserves an auditable before-and-after record: "
            "the classifier result supports the current claim, while the proxy remains visible as "
            "evidence for why the interpretation changed.",
        ],
    )
    add_figure(
        doc,
        "5.1",
        "HarmBench classifier ASR versus the refusal-regex proxy and their agreement.",
        "judge_vs_proxy.png",
        14.5,
    )

    add_heading(doc, "5.2", "Independent HarmBench Judges")
    api_kappas = [row["cohens_kappa"] for row in JUDGE_API["per_model"]]
    lg_kappas = [row["cohens_kappa"] for row in JUDGE_LLAMAGUARD["per_model"]]
    add_paragraphs(
        doc,
        [
            f"The API judge agrees strongly with the HarmBench classifier across the ten fp16/NF4 "
            f"aliases, with per-alias kappa from {min(api_kappas):.2f} to {max(api_kappas):.2f}. It "
            "reproduces the direction and significance status of the substantive per-pair conclusion: "
            "no harmful-ward NF4 increase is significant after correction. For Llama, both judges find "
            "a decrease.",
            f"Llama Guard provides an open-weight cross-check. Its agreement with the behaviour-specific "
            f"classifier is more variable, with per-alias kappa from {min(lg_kappas):.2f} to "
            f"{max(lg_kappas):.2f}. Every Llama Guard per-pair delta is non-significant. The lowest "
            "agreement occurs on Phi, where Llama Guard's broader 'unsafe response' construct differs "
            "most from HarmBench's behaviour-specific criterion. This disagreement is informative "
            "rather than proof that one scorer is mechanically wrong.",
            "The API judge and classifier agree most strongly on the Llama pair, where kappa exceeds "
            "0.94 for both aliases. Every other alias falls in a single lower band from 0.68 to 0.83, "
            "with the minimum on Phi-4-mini at fp16 and the Qwen3-4B pair next; no other family is "
            "systematically separated within that band. The per-pair harmful-ward significance "
            "conclusion is unchanged across all of them. The second judge strengthens the headline "
            "without implying item-level identity.",
            "Llama Guard's per-pair directions agree with the classifier for four pairs. "
            "Qwen3-1.7B is the exception: the classifier is flat and Llama Guard decreases by "
            "0.050, and neither result is significant. Similarly, on Llama, Llama Guard agrees on "
            "direction but not on the classifier's uncorrected significance. The report therefore "
            "rests on the corrected family rather than one judge-specific p-value.",
        ],
    )
    judge_rows = []
    api_by_pair = {row["pair_id"]: row for row in JUDGE_API["per_pair"]}
    lg_by_pair = {row["pair_id"]: row for row in JUDGE_LLAMAGUARD["per_pair"]}
    for pair_id in PAIR_ORDER:
        a = api_by_pair[pair_id]
        l = lg_by_pair[pair_id]
        judge_rows.append(
            (
                PAIR_NAMES[pair_id],
                fmt_delta(a["harmbench_cls_delta"]),
                fmt_delta(a["api_judge_delta"]),
                fmt_delta(l["llamaguard_delta"]),
                "No" if not l["llamaguard_significant"] else "Yes",
            )
        )
    add_table(
        doc,
        "5.1",
        "Per-pair HarmBench deltas under three judges",
        ["Pair", "HarmBench classifier Δ", "API judge Δ", "Llama Guard Δ", "Llama Guard p<.05"],
        judge_rows,
        [3.2, 3.2, 3.0, 3.0, 3.2],
        font_size=8.0,
    )

    add_heading(doc, "5.3", "Human Grounding of Harmful-Compliance Scorers")
    hb_cls = HUMAN_HARMBENCH["classifier_vs_human"]
    hb_regex = HUMAN_HARMBENCH["regex_vs_human"]
    add_paragraphs(
        doc,
        [
            f"The 200-item human sample grounds the scorer choice. The HarmBench classifier reaches "
            f"kappa {hb_cls['cohens_kappa']:.2f} against the annotator, compared with "
            f"{hb_regex['cohens_kappa']:.2f} for the regex. The classifier flags "
            f"{hb_cls['over_flag_vs_human']} items that the human does not, whereas the regex over-flags "
            f"{hb_regex['over_flag_vs_human']}. The classifier is not perfect, but it is substantially "
            "closer to the human construct.",
            "The sample is stratified toward scorer disagreements, so its class prevalence is not "
            "an estimate of population ASR. It has one annotator, which limits the gold-standard "
            "claim. Accordingly, the correct interpretation is comparative: within this audit "
            "sample, the classifier is better grounded than the refusal proxy. A second annotator "
            "would quantify the human reliability ceiling.",
            f"The classifier's recall is {hb_cls['recall']:.3f} and its precision is "
            f"{hb_cls['precision']:.3f} on the audit sample. The regex reaches higher recall "
            f"({hb_regex['recall']:.3f}) by flagging many more responses, but its precision falls to "
            f"{hb_regex['precision']:.3f}. This error profile matches the conceptual concern: absence "
            "of a refusal phrase is treated as harmful completion.",
            "Kappa is affected by prevalence and marginal distributions. The report therefore presents "
            "precision, recall, and counts beside kappa. No single agreement coefficient is used as a "
            "universal pass threshold.",
        ],
    )

    add_heading(doc, "5.4", "XSTest Scorer Sensitivity and Human Validation")
    xs_strict_j = HUMAN_XSTEST["strict"]["judge_vs_human"]
    xs_strict_r = HUMAN_XSTEST["strict"]["regex_vs_human"]
    xs_broad_j = HUMAN_XSTEST["broad"]["judge_vs_human"]
    xs_broad_r = HUMAN_XSTEST["broad"]["regex_vs_human"]
    add_paragraphs(
        doc,
        [
            "The independent three-class refusal judge does not reproduce the Phi-4-mini regex "
            "decrease. Its strict mapping gives a small positive delta and its broad mapping is "
            "near zero, and neither is significant. Across all pairs, no semantic-judge result "
            "shows a statistically significant increase in over-refusal.",
            f"Human validation explains the discrepancy. Under the strict definition, the judge reaches "
            f"kappa {xs_strict_j['cohens_kappa']:.3f}, while the regex is "
            f"{xs_strict_r['cohens_kappa']:.3f}. The regex recalls only "
            f"{xs_strict_r['recall']*100:.1f}% of full refusals; the judge recalls "
            f"{xs_strict_j['recall']*100:.1f}%. Under the broad definition, kappa is "
            f"{xs_broad_j['cohens_kappa']:.3f} for the judge and "
            f"{xs_broad_r['cohens_kappa']:.3f} for the regex. The most plausible explanation is that "
            "the Phi survivor reflects regex measurement error rather than a stable behavioural shift.",
            "The XSTest human audit contains a disclosed protocol deviation. The blindness exclusion "
            "was applied prompt-wide rather than to the limited alias–prompt items specified in the "
            "pre-registration. The change occurred before drawing and labelling and removed more "
            "potential priming, but it departs from the written rule. The clean pre-registered badge is "
            "therefore withdrawn. A counterfactual overlap analysis bounds the displacement at no more "
            "than 14 of 200 items, and the large judge–regex kappa gap remains on the shared set.",
            f"The human set contains {HUMAN_XSTEST['human_label_counts']['full_compliance']} full "
            f"compliances, {HUMAN_XSTEST['human_label_counts']['partial_refusal']} partial refusals, "
            f"and {HUMAN_XSTEST['human_label_counts']['full_refusal']} full refusals. The judge's "
            f"three-class exact agreement with the annotator is "
            f"{HUMAN_XSTEST['three_class_human_vs_judge']['exact_agreement_rate']:.3f}.",
            "The broad mapping increases the refusal class by counting partial refusals. The judge's "
            "kappa rises under that mapping, while the regex remains near zero because it still misses "
            "most semantic refusals. The scorer ranking is therefore not driven by one definition of "
            "partial refusal.",
            "The validation-informed multiplicity family replaces only the over-refusal "
            "instrument with judge-strict scoring, and it retains the five pairs and other "
            "outcomes. Its two survivors are the Qwen MMLU and Llama ARC losses, and no safety or "
            "over-refusal contrast survives.",
        ],
    )

    add_heading(doc, "5.5", "Generation-Length Robustness")
    add_paragraphs(
        doc,
        [
            f"The 128-token runs truncated {GENLEN['prefix_truncation_128']['pct_truncated']:.1f}% of "
            "the 2,000 HarmBench responses under the direct prefix test. At that shorter budget, "
            "Qwen3-1.7B appeared to increase by 0.055. At 512 tokens, the same pair is exactly flat. "
            "The earlier signal was therefore a generation-length artefact rather than a result that "
            "generalised to HarmBench's reference budget.",
            "The comparison also changes absolute baseline rates because longer responses give models "
            "more opportunity either to comply or to qualify an answer. This is why the 512-token tree "
            "is primary and the 128-token tree is retained only as a sensitivity comparison. The "
            "finding illustrates a broader principle: a generation cap can become part of the measured "
            "construct when truncation is common.",
            "Truncation is highly model-dependent. At 128 tokens it affects most Mistral and Phi "
            "responses, a majority of Qwen responses, and relatively few Llama responses, so a "
            "pooled percentage conceals different exposure across families. Consequently, the "
            "512-token rerun is necessary for cross-family interpretation, not merely a cosmetic "
            "extension.",
            "A direct prefix test distinguishes responses that are exact prefixes of the longer output "
            "from natural stops and mismatches. This provides empirical truncation evidence rather than "
            "relying only on token counts. The mismatch category remains visible, so the analysis does "
            "not claim that every budget difference is truncation.",
        ],
    )
    gen_rows = []
    for pair_id in PAIR_ORDER:
        r128 = GENLEN["asr_128_classifier"][pair_id]
        r512 = GENLEN["asr_512_classifier"][pair_id]
        gen_rows.append(
            (
                f"{PAIR_NAMES[pair_id]} (128/512 budgets)",
                fmt_delta(r128["delta"]),
                fmt_delta(r512["delta"]),
                f"{r128['mcnemar_p']:.4f}",
                f"{r512['mcnemar_p']:.4f}",
            )
        )
    add_table(
        doc,
        "5.2",
        "HarmBench classifier deltas at 128 and 512 tokens",
        ["Pair", "ΔASR @128", "ΔASR @512", "McNemar p @128", "McNemar p @512"],
        gen_rows,
        [3.6, 2.7, 2.7, 3.3, 3.3],
        font_size=8.2,
    )

    add_heading(doc, "5.6", "Multi-Seed Decoding Sensitivity")
    seed_rows = []
    for row in MULTISEED["per_pair"]:
        j = row["judge_delta"]
        seed_rows.append(
            (
                PAIR_NAMES[row["pair_id"]],
                fmt_delta(row["greedy_judge_delta"]),
                fmt_delta(j["mean"]),
                f"[{j['min']:+.3f}, {j['max']:+.3f}]",
                f"{j['n_seeds_significant_p05']}/5",
            )
        )
    add_paragraphs(
        doc,
        [
            "The five-seed stochastic arm tests whether the greedy result is an accident of one decode. "
            "Qwen3-1.7B has mean ΔASR +0.013 across seeds, with no significant seed. Qwen3-4B has one "
            "significant harmful-ward seed out of five, but the mean remains within the small range "
            "shown in Table 5.3. Llama's seeds are all non-positive, and its mean is −0.024.",
            "Mistral and Phi remain close to zero on average. Their greedy deltas lie just outside narrow "
            "seed ranges that straddle or approach zero, which is a reminder that deterministic and "
            "stochastic decoding are different conditions. Across all five pairs, the arm does not "
            "produce a stable significant harmful-ward effect.",
            "A seed is not an independent model replicate. The weights and prompt set are reused, and "
            "only decoding randomness changes. The arm therefore measures within-condition generation "
            "variability rather than uncertainty over model training.",
            "One Qwen3-4B seed and one Phi seed are individually significant, in opposite "
            "directions, and selecting either after inspection would produce a misleading "
            "narrative. Reporting all five seeds and their mean therefore shows that isolated "
            "significance is compatible with repeated testing and stochastic variation.",
        ],
    )
    add_table(
        doc,
        "5.3",
        "Five-seed HarmBench sensitivity at the 512-token budget",
        ["Pair", "Greedy Δ", "Seed mean Δ", "Seed range", "Seeds p<.05"],
        seed_rows,
        [3.6, 2.5, 2.8, 3.8, 2.9],
        font_size=8.2,
    )
    add_body(
        doc,
        "Figure 5.2 shows the Qwen3-1.7B greedy headline against each individual seed. The "
        "greedy delta coincides with the lowest two seed estimates and sits within the seed "
        "mean plus or minus one standard deviation, so the headline is not separated from "
        "what decoding variation alone produces on this pair.",
    )
    add_figure(
        doc,
        "5.2",
        "Qwen3-1.7B greedy result and five-seed HarmBench sensitivity.",
        "multiseed.png",
        14.5,
    )

    add_heading(doc, "5.7", "Power and the Meaning of the Null")
    add_paragraphs(
        doc,
        [
            "Failure to reject is not proof that the true effect is zero. With 200 paired HarmBench "
            "items per model, the minimum detectable absolute ASR change varies with the discordant "
            "rate and is roughly 0.04–0.09 across pairs. The study can therefore reject a large common "
            "regression, but smaller effects may remain unresolved.",
            "The null is nevertheless more informative than a single non-significant p-value. It is "
            "supported by four distinct checks: no ASR contrast survives correction, three semantic "
            "judges agree on the absence of a significant harmful-ward pair, the 128-token signal "
            "dissolves at the reference budget, and the multi-seed arm does not reveal a stable "
            "harmful-ward effect. These instruments address multiplicity, construct validity, "
            "truncation, and decoding variability respectively.",
            "Capability nulls have different resolution because MMLU and ARC sample sizes differ. "
            "Appendix B shows that some small observed deltas fall below their estimated MDE. The "
            "report uses 'no detectable change' for those cases and reserves equivalence language for "
            "the explicitly tested INT8 margin.",
            "The defensible claim is the absence of a corrected increase large enough to be detected "
            "in this design. It is not proof of exact preservation. This wording keeps the "
            "conclusion aligned with the study's power.",
        ],
    )

    add_heading(doc, "5.8", "Robustness Synthesis")
    add_paragraphs(
        doc,
        [
            "The validation results can be read as a sequence of eliminated explanations. The "
            "refusal proxy is not trusted because human labels show severe over-flagging. The "
            "classifier is not trusted alone either, so an API judge and Llama Guard test judge "
            "dependence. The 128-token result is not trusted because truncation is common, and "
            "the 512-token rerun replaces it. Finally, greedy decoding is not assumed "
            "representative, so five stochastic seeds test sensitivity.",
            "No single check is decisive. Together, however, they make a large hidden harmful-"
            "ward effect less plausible in the tested setting. Remaining uncertainty is "
            "concentrated in effects below the MDE, attack-based threat models, languages and "
            "model families outside the sample, and the subjectivity of single-annotator human "
            "labels.",
            "This pattern is more valuable than agreement for its own sake. Where instruments disagree, "
            "the report identifies the construct difference and gives greater weight to the scorer with "
            "better human grounding. Where they agree, the conclusion is stated only at the scope and "
            "resolution supported by the design.",
        ],
    )


def add_chapter_six(doc: Document) -> None:
    add_chapter(doc, 6, "Discussion")

    add_heading(doc, "6.1", "Answer to the Central Question")
    add_paragraphs(
        doc,
        [
            "The study began with a distinction: an apparent safety change may reflect alignment, "
            "capability, or measurement. The results favour the latter two explanations. Under the "
            "primary HarmBench classifier at 512 tokens, NF4 produces no corrected harmful-ward "
            "increase. Capability losses, by contrast, account for two of the three multiplicity-robust "
            "survivors. The third survivor is a benign-direction over-refusal change that disappears "
            "under the better-grounded semantic scorer.",
            "This conclusion is deliberately bounded. It does not establish that NF4 preserves all "
            "forms of safety, that effects are absent for every small model, or that an effect smaller than the "
            "study's detection threshold is absent. It states that the evaluated direct-request "
            "experiment found no corrected increase in harmful compliance and that the clearest costs "
            "appeared in capability.",
            "The evidence also rejects two tempting simplifications. Quantization is not reliably "
            "safety-improving: the directions differ by pair, and the only individually "
            "significant classifier movement is accompanied by capability loss. Similarly, "
            "quantization is not reliably safety-degrading either: the harmful-ward estimates are "
            "small, unresolved, and not reproduced as corrected effects across judges.",
            "A more useful description is model-specific boundary instability under a capability cost. "
            "Some prompts change classification in both directions even when the aggregate delta is "
            "zero. That pattern is consistent with a numerical perturbation near decision boundaries. "
            "The behavioural data alone cannot identify the internal mechanism, so the explanation "
            "remains a hypothesis rather than a causal finding.",
        ],
    )
    add_table(
        doc,
        "6.1",
        "Competing explanations for the HarmBench null",
        ["Candidate explanation", "Discriminating evidence", "Assessment"],
        [
            ("True practical preservation", "No corrected ASR increase; three judges; five pairs across four families", "Compatible within detection bounds"),
            ("Insufficient statistical power", "Per-pair MDE roughly 0.04–0.09", "Small effects remain possible"),
            ("Scorer insensitivity", "Classifier human κ 0.59; API and Llama Guard cross-checks", "Reduced but not eliminated"),
            ("Generation truncation", "60.3% truncation at 128; signal disappears at 512", "Explains the retired Qwen signal"),
            ("Decode-specific result", "Five-seed arm has no stable harmful-ward effect", "Not supported as a common effect"),
        ],
        [4.1, 7.0, 4.5],
        font_size=8.0,
    )

    add_heading(doc, "6.2", "Relationship to Prior Work")
    add_paragraphs(
        doc,
        [
            "The result is closest to Decoding Compressed Trust, which reports broad preservation of "
            "trustworthiness at four bits and larger damage under more aggressive compression [37]. "
            "The agreement is not exact because the model scales, compression methods, trust dimensions, "
            "and scorers differ. It is therefore evidence of compatibility, not replication.",
            "Safety-degradation studies remain relevant. Kharinaev et al. report that quantization can "
            "degrade safety, and that no single method universally outperforms the others across models "
            "and bit-widths [10]. HarmLevelBench finds attack-dependent changes [35], while Q-resafe and "
            "Q-Misalign show that quantized safety can be both repaired and attacked [36], [39]. The "
            "present null does not contradict those findings. It narrows the setting to ordinary NF4 and "
            "INT8 loading of five unmodified instruction-tuned checkpoints under direct requests.",
            "Table 6.2 places the published figures beside this study's rather than leaving the "
            "comparison qualitative. Every value is transcribed from the source's own results table on "
            "the source's own scale. The prior studies report percentage points on a 0–100 scale, while "
            "this study reports proportions, so the table gives both forms for this study's row: a delta "
            "of −0.090 is a drop of 9.0 points. Reading the columns without that conversion inverts the "
            "apparent size of the effects.",
        ],
    )
    add_table(
        doc,
        "6.2",
        "Published quantization effects placed beside this study's. Values are transcribed from each "
        "source's own results table on that source's own scale: prior rows are percentage points on "
        "0–100, and this study's row gives proportions with the point equivalent in brackets. Every "
        "row measures a different intervention, prompt set or scorer, so no row replicates or refutes "
        "another; the final column names the chief obstacle in each case. Reference [35] is a "
        "non-archival workshop paper distributed as a preprint.",
        [
            "Study",
            "Quantized artefact",
            "Reported change, in the source's own metric",
            "Scorer",
            "Chief obstacle to direct comparison",
        ],
        [
            (
                "Decoding Compressed Trust [37]",
                "Llama-2-Chat 7B/13B/70B; GPTQ 4-bit, 128-sample calibration, vs 16-bit dense",
                "MMLU −1.5 points at 13B (53.6 to 52.1); −1.6 at 7B and −0.9 at 70B",
                "MMLU accuracy, 57-task average",
                "Calibrated GPTQ, not load-time NF4. Its smallest model is 7B, so the 1.7B regime where "
                "this study finds its capability effect is unsampled. Its safety dimensions are "
                "Perspective-API toxicity and trust perspectives, which yield no compliance rate.",
            ),
            (
                "Kharinaev et al. [10]",
                "Llama-3.1-8B-Instruct; AWQ, QUIK and QUIP# at int4, vs bfloat16",
                "Judged safety ratio falls 0.571, 0.536 and 4.500 points on XSafety, and 3.561 and "
                "8.622 on OpenSafetyMini, while QUIK rises 0.187. Positive means less safe",
                "Gemma-2-27B judge, four labels",
                "Polarity is inverted relative to an attack success rate. The spread across methods at "
                "one bit-width exceeds any single effect, and NF4 is not among them. An “error” label "
                "folds incoherent output into the safety score, mixing capability loss into it.",
            ),
            (
                "HarmLevelBench [35]",
                "Vicuna-13B-v1.5; AWQ, bit-width not reported",
                "Human-annotated ASR on direct queries rises 33.9 to 44.6, a gain of 10.7 points",
                "Human annotation; the same rows give +16.1 by string match and +1.7 by GPT judge",
                "A 56-item severity-stratified set of the authors' own, so 10.7 points is six queries. "
                "Bit-width, decoding settings and baseline precision are all unreported, and the "
                "direction depends on which of the paper's three scorers is read.",
            ),
            (
                "Q-resafe [36]",
                "Llama-2-7B-Chat; AWQ 4-bit vs the same fp16 checkpoint",
                "ASR rises 29.8 to 37.1 points at τ=0.95; the gap runs 4.5 to 17.3 points across six "
                "sampling settings",
                "Decoding-attack ASR: GPT-4 selects the highest-scoring response from 49 decoding configurations",
                "Measured under an active decoding attack with GPT-4 best-of-49 selection. The same fp16 "
                "checkpoint scores 0.3 un-attacked, so almost all of the 29.8-point reference level is "
                "attack effect rather than direct-request behaviour.",
            ),
            (
                "Q-Misalign [39]",
                "InternLM2-Chat-1.8B, Gemma-1.1-2B-it and Llama-2-7B-Chat; int8, the paper's default",
                "Jailbreak success identical to the two decimals printed: 0.07 to 0.07, 0.05 to 0.05, "
                "0.00 to 0.00 (proportions)",
                "Not reported",
                "The paper's control condition, at 8-bit with no 4-bit counterpart, on AdvBench rather "
                "than HarmBench. Its headline jump to 0.97 belongs to a checkpoint the authors prepared "
                "to misbehave once quantized, and is not an effect of ordinary quantization.",
            ),
            (
                "Present study",
                "Five instruction-tuned pairs, 1.7–7.2B; bitsandbytes NF4 and LLM.int8 applied at load "
                "time to one pinned revision",
                "ASR −0.040 to +0.040 (−4.0 to +4.0 points), none surviving correction; MMLU to −0.090 "
                "(−9.0 points)",
                "HarmBench classifier, grounded against human labels",
                "Direct requests with no attack augmentation, and a single greedy generation per prompt.",
            ),
        ],
        [2.3, 3.0, 3.3, 2.1, 4.9],
        font_size=7.0,
    )
    add_paragraphs(
        doc,
        [
            "Two features of the table deserve emphasis because they are easy to misread. Kharinaev et "
            "al. report a safety ratio and calculate the difference as FP16 minus quantized, so a "
            "positive difference means the quantized model was judged less safe, the opposite "
            "orientation to the attack success rates in the neighbouring "
            "rows [10]. And the only prior row measuring an unmodified checkpoint under ordinary "
            "quantization without an attack is Q-Misalign's control condition, which shows no change on "
            "three checkpoints [39]. That row agrees with the present null in direction, but it is "
            "reported at 8-bit with no sample size or interval, so it corroborates rather than "
            "replicates.",
            "Read together, the table supports the narrower reading this report has taken throughout. "
            "The published increases in harmful compliance are attached either to an attack, as in "
            "Q-resafe and the direct-query rows of HarmLevelBench, or to a calibrated offline method "
            "that rewrites the checkpoint, or to a model deliberately prepared to fail. None of them "
            "establishes that loading an ordinary instruction-tuned checkpoint in NF4 raises harmful "
            "compliance, which is the specific question this study measured and answered in the "
            "negative within its detection bounds.",
            "The scorer findings align with StrongREJECT's warning about empty jailbreaks [41] "
            "and SORRY-Bench's insistence that safety refusal be interpreted with utility [42]. "
            "The HarmBench regex counted willingness, while the classifier evaluated behaviour "
            "completion. On XSTest, by contrast, the opposite failure occurred: the regex missed "
            "semantic refusals. Simple string rules can therefore fail in both directions "
            "depending on the construct.",
            "The capability findings also qualify earlier capability-focused quantization work. "
            "The three large evaluations reviewed here report good average retention at four bits "
            "[7]-[9], but the "
            "Qwen3-1.7B MMLU loss shows that compact models and task subsets can still contain "
            "material failures. Conversely, the Mistral ARC gain and Qwen3-4B stability show that "
            "the loss is not universal.",
            "Marchisio et al.'s multilingual result is methodologically relevant even though the "
            "present experiment is English-only [13]. Their human evaluation reveals damage that "
            "automatic metrics understate. Here, human validation reveals scorer error that changes "
            "which result deserves weight. In both cases, the evaluator is part of the empirical "
            "question.",
            "The direct-request boundary explains why adversarial studies should not be used as "
            "either confirmation or contradiction. JailbreakBench evaluates a model under an "
            "attack procedure [43], while Q-Misalign constructs a quantization-conditioned attack "
            "[39]. The current design evaluates ordinary quantization without an attack. The "
            "three settings therefore address complementary risks.",
        ],
    )

    add_heading(doc, "6.3", "Capability as the Deployment Cost")
    add_paragraphs(
        doc,
        [
            "The capability result is practically important because a compressed model is chosen "
            "for deployment, not only for a benchmark. Qwen3-1.7B loses nine MMLU points under "
            "NF4, and Llama loses about 3.2 ARC points. These losses do not occur uniformly, but "
            "they are large enough to change which tasks a model can be trusted to perform.",
            "INT8 provides a useful intermediate point. Its capability contrasts are close to fp16 "
            "within the stated equivalence margin, while the more visible losses occur at NF4. This "
            "supports model-specific validation before choosing the smallest representation. However, "
            "the report does not measure memory, latency, energy, or throughput. It cannot calculate a "
            "complete engineering trade-off between resource savings and lost accuracy.",
            "The two capability anchors also caution against a single-number deployment decision. "
            "Qwen3-1.7B looks much worse on MMLU than ARC, while Llama's strongest corrected loss "
            "is on ARC. Therefore, a deployment team should evaluate tasks close to its own "
            "workload rather than treat one aggregate benchmark as a universal capability score.",
            "The missing cost axis limits the deployment recommendation. NF4 is expected to reduce "
            "memory relative to fp16, and INT8 occupies an intermediate precision, but this project did "
            "not log peak resident memory or throughput in a controlled benchmark. It would be "
            "inappropriate to state a measured speedup or memory ratio. The report can recommend a "
            "validation process, not calculate the optimal deployment point.",
            "A responsible process begins with the application. If the workload resembles "
            "multiple-choice knowledge tasks, the Qwen MMLU result is directly concerning, "
            "whereas if it requires open-ended reasoning or coding, neither MMLU nor ARC is "
            "sufficient. Accordingly, the same matched comparison should be repeated on "
            "representative tasks before deployment.",
            "Safety should be evaluated on the final serving stack. A pre-quantized checkpoint, a "
            "GGUF runtime, or different attention kernels may not reproduce the BitsAndBytes "
            "result. Chat templates, system prompts, decoding limits, and moderation layers also "
            "contribute to the deployed behaviour. Overall, precision is one component of a "
            "larger system.",
        ],
    )

    add_heading(doc, "6.4", "Scorer Validity as a Research Result")
    add_paragraphs(
        doc,
        [
            "Scorer validation is not an auxiliary engineering check, because it changes the "
            "scientific conclusion. The refusal proxy overstates HarmBench success and the XSTest "
            "regex misses most human full refusals. In both cases, a semantic instrument is "
            "better grounded against the human sample.",
            "At the same time, semantic judges introduce their own limitations. The HarmBench API "
            "judge is external and its exact dated snapshot was not recorded, Llama Guard is "
            "open-weight but uses a broader policy "
            "construct, and human labels are more direct but expensive and subjective. Evidence "
            "is therefore triangulated across the primary benchmark-specific classifier, "
            "alternative judges, prompt-level agreement, and a disclosed human audit.",
            "This has a general implication for LLM evaluation. A benchmark score should be understood "
            "as the output of a measurement system, not a direct observation of an abstract property. "
            "When the estimated treatment effect is small, scorer error can be as large as the effect. "
            "Validation should therefore target the cases that determine the conclusion, not only an "
            "overall accuracy figure.",
            "The two scorer failures point in opposite directions. On HarmBench, a refusal proxy "
            "over-counts success because it cannot evaluate response content. On XSTest, by "
            "contrast, a narrow refusal regex under-counts refusal because it misses indirect "
            "language. A transparent rule is therefore not automatically conservative, because "
            "its bias depends on how its surface patterns relate to the construct.",
            "For the XSTest judge sensitivity study, its dated pre-registration and subsequent "
            "validation must coexist. Keeping the original regex family protects "
            "against outcome-driven scorer replacement. Reporting the validation-informed family "
            "prevents procedural loyalty from overriding better construct evidence. Presenting both "
            "allows the reader to see that the headline is scorer-invariant where it matters: no "
            "instrument finds a statistically significant increase in over-refusal.",
            "The protocol deviation demonstrates the same principle. The blindness exclusion was more "
            "conservative in one sense, because it removed more potentially priming prompts, but it was "
            "not the rule that had been written. Disclosing the change and withdrawing the clean badge "
            "is stronger evidence practice than retroactively redefining the protocol.",
        ],
    )

    add_heading(doc, "6.5", "Threats to Validity")
    add_heading(doc, "6.5.1", "Internal Validity", level=3)
    add_paragraphs(
        doc,
        [
            "The matched revisions, shared prompts, fixed decoding controls, and paired tests "
            "strengthen causal attribution to the load-time precision change. However, residual "
            "threats include numerical differences in kernels, model-specific chat-template "
            "behaviour, and the historical environment-provenance gap. The primary tree is "
            "internally consistent, but the earliest runs cannot be claimed as bit-for-bit "
            "reproducible on a new machine.",
            "The XSTest human audit deviated from its pre-registered blindness-exclusion rule. The "
            "deviation is disclosed, bounded, and tested on the overlapping items. It still removes the "
            "clean confirmatory label. Selective interpretation is controlled by retaining the original "
            "regex family beside the validation-informed family.",
            "The sensitivity arms were conducted after the primary runs. Their role is diagnostic, not "
            "confirmatory. The report labels them separately and does not insert their p-values into the "
            "primary multiplicity family. This preserves the distinction between the primary analysis and an "
            "explanation developed after observing a surprising result.",
            "Several analysis choices are rule-based, including parser behaviour and "
            "interpretation labels. Automated tests protect these rules against accidental "
            "change. However, a rule can still be scientifically imperfect, so human audits and "
            "alternative scorers remain necessary even when the code is deterministic.",
        ],
    )
    add_heading(doc, "6.5.2", "Construct Validity", level=3)
    add_paragraphs(
        doc,
        [
            "HarmBench DirectRequest measures harmful compliance under direct unsafe instructions. It "
            "does not measure resilience to GCG, PAIR, prefilling, or other jailbreak procedures. "
            "XSTest measures exaggerated safety on a compact diagnostic set. MMLU and ARC are useful "
            "anchors but do not cover open-ended reasoning, coding, or domain-specific tasks.",
            "Human validation uses one annotator on each 200-item sample. The comparisons are therefore "
            "grounded but not a multi-annotator gold standard. A second annotator and adjudication would "
            "separate scorer error from ambiguity in the rubric.",
            "The HarmBench classifier is trained for the benchmark's behaviour criterion, which "
            "supports construct alignment but creates dependence on one learned evaluator. The "
            "API judge shares the criterion but not the weights, whereas Llama Guard changes both "
            "weights and construct. The triangulation is useful precisely because these "
            "dependencies differ.",
            "Over-refusal is especially policy-sensitive. A partial refusal can be counted as acceptable "
            "caution or as lost helpfulness depending on the application. Strict and broad mappings make "
            "that decision visible, but they do not eliminate it.",
        ],
    )
    add_heading(doc, "6.5.3", "External Validity", level=3)
    add_paragraphs(
        doc,
        [
            "The five pairs cover four families and a compact-model range, but they are not a random "
            "sample. Results may differ for larger checkpoints, base models, other alignment recipes, "
            "multilingual prompts, or different hardware. NF4 and LLM.int8 do not represent GPTQ, AWQ, "
            "GGUF, or all post-training quantization.",
            "The study's strongest generalisable contribution is methodological: match the "
            "weights, evaluate safety and capability together, pair the statistics, declare "
            "multiplicity, and validate the scorer. By contrast, the numerical results should "
            "remain attached to the evaluated models and conditions.",
            "English-language benchmarks dominate the study. Qwen and other multilingual families may "
            "respond differently in non-Latin scripts, and prior work reports disparate quantization "
            "effects by language [13]. The English null should not be exported to multilingual safety.",
            "The V100 environment is relevant because numerical kernels and supported attention "
            "paths differ across hardware generations. Both members of each pair share the "
            "environment, which protects the within-pair contrast. Nevertheless, generalisation "
            "to other runtimes remains an empirical question.",
        ],
    )

    add_heading(doc, "6.6", "Practical Implications")
    add_bullets(
        doc,
        [
            "Do not inherit a safety claim from an fp16 checkpoint without testing the deployed precision.",
            "Use a behaviour-specific semantic scorer for harmful compliance; retain transparent proxies only as secondary diagnostics.",
            "Pair safety with benign-refusal and capability benchmarks so that inability is not reported as alignment improvement.",
            "Prefer INT8 when its smaller compression benefit is sufficient and task accuracy is critical; validate NF4 on workload-relevant tasks.",
            "Treat direct-request results as direct-request results. A separate adversarial arm is required for jailbreak robustness.",
        ],
    )
    add_body(
        doc,
        "The practical message is procedural rather than promotional. NF4 is not declared unsafe, and "
        "it is not declared free. It should be evaluated as a distinct deployed artefact, with task-"
        "relevant capability checks and a scorer that measures the intended safety construct.",
    )


def add_chapter_seven(doc: Document) -> None:
    add_chapter(doc, 7, "Conclusion and Future Work")

    add_heading(doc, "7.1", "Conclusions")
    add_paragraphs(
        doc,
        [
            "This project evaluated five matched compact-language-model pairs across HarmBench, "
            "XSTest, MMLU, and ARC-Challenge, with NF4 as the primary intervention and INT8 as an "
            "additional precision point. The design treated the quantized deployment artefact as the "
            "object of evaluation rather than assuming that the full-precision safety profile transfers.",
            "RQ1 is answered by a bounded null: no NF4 harmful-compliance contrast survives the "
            "20-test BH-FDR correction at the 512-token reference budget. RQ2 finds no significant "
            "increase in benign over-refusal under any scorer. RQ3 finds the largest corrected losses in "
            "capability, especially Qwen3-1.7B on MMLU and Llama-3.2-3B on ARC. RQ4 shows that the "
            "magnitude depends on model, family, task, and precision; INT8 stays closer to fp16 than "
            "NF4 on capability. RQ5 shows that the substantive safety conclusion survives alternative "
            "judges, human grounding, the longer generation budget, and the multi-seed check.",
            "The central lesson is methodological. Quantization effects cannot be interpreted "
            "from a single refusal count. Instead, safety, helpfulness, and capability must be "
            "measured together, and the scorer must be validated against the construct it claims "
            "to measure.",
            "This report also changes how the contribution is presented. The benchmarking "
            "framework, cluster jobs, and audit machinery support the study, but they are not "
            "treated as the primary result. The contribution is instead the evidence produced by "
            "a controlled design, together with the measurement lessons revealed when the "
            "original scorers were challenged.",
            "Within its boundaries, the study provides a clear deployment conclusion. INT8 "
            "preserves capability more consistently than NF4 under the tested margins, whereas "
            "NF4 can impose material task-specific losses. Neither precision produces a corrected "
            "harmful-compliance increase at the 512-token DirectRequest setting. A deployment "
            "team should therefore prioritise capability validation while retaining safety checks "
            "on the final artefact.",
        ],
    )

    add_heading(doc, "7.2", "Prioritised Future Work")
    add_bullets(
        doc,
        [
            "Add an adversarial HarmBench arm using at least one standardised attack. This is the highest-priority extension because the current safety claim is limited to direct requests.",
            "Collect a second independent annotator for both human-validation samples and report inter-annotator agreement with adjudication.",
            "Compare NF4 and LLM.int8 with GPTQ, AWQ, and a deployment-relevant GGUF path on the same pinned checkpoints.",
            "Measure peak memory, model footprint, load time, throughput, and energy so that the capability loss can be compared with an observed deployment benefit.",
            "Extend capability anchoring to open-ended reasoning, coding, and workload-specific tasks, and evaluate multilingual safety where the model family supports it.",
            "Increase the number of decoding seeds and apply the seed analysis to XSTest and capability, not only HarmBench.",
        ],
    )
    add_body(
        doc,
        "These extensions should preserve the discipline used here: immutable original artefacts, "
        "derived sidecars for revised scoring, matched checkpoint revisions, and a clearly separated "
        "primary analysis and sensitivity layer.",
    )
    add_body(
        doc,
        "The first two extensions are the most urgent. An adversarial arm would enlarge the safety "
        "claim beyond direct requests, and a second annotator would strengthen the human reference "
        "standard on both safety axes. Efficiency measurements and additional quantizers would then "
        "allow a complete safety–capability–resource comparison.",
    )


def add_references(doc: Document) -> None:
    p = doc.add_heading("References", level=1)
    p.paragraph_format.page_break_before = True
    for index, reference in enumerate(REFERENCES, start=1):
        doc.add_paragraph(f"[{index}] {reference}", style="Reference")


def add_appendices(doc: Document) -> None:
    p = doc.add_heading("Appendix A: Reproducibility Summary", level=1)
    p.paragraph_format.page_break_before = True
    add_table(
        doc,
        "A.1",
        "Primary run controls",
        ["Control", "Value"],
        [
            ("Primary configuration", "configs/tc1_512.yaml"),
            ("Model pairs", "5 fp16–NF4 pairs; separate INT8 aliases"),
            ("Generation", "Greedy, temperature 0.0, max_new_tokens 512, seed 42"),
            (
                "Quantization",
                "bitsandbytes NF4 with double quantization; compute dtype follows the resolved "
                "model dtype. LLM.int8 for the INT8 arm",
            ),
            (
                "Library defaults overridden",
                "Eager attention for both Phi-4-mini members; generation batch size 4 per benchmark",
            ),
            (
                "Declared dependencies",
                "transformers 4.41.0 or later, torch 2.2.0 or later, bitsandbytes 0.43.0 or later. "
                "These are the declared floors; the exact runtime versions of the primary runs were "
                "not captured in the immutable outputs",
            ),
            (
                "Target checkpoint revisions",
                "Qwen/Qwen3-1.7B revision 70d244cc86ccca08cf5af4e1e306ecf908b1ad5e; "
                "Qwen/Qwen3-4B revision 1cfa9a7208912126459214e8b04321603b3df60c; "
                "meta-llama/Llama-3.2-3B-Instruct revision 0cb88a4f764b7a12671c53f0838cd831a0843b95; "
                "mistralai/Mistral-7B-Instruct-v0.3 revision c170c708c41dac9275d15a8fff4eca08d52bab71; "
                "microsoft/Phi-4-mini-instruct revision cfbefacb99257ffa30c83adab238a50856ac3083",
            ),
            (
                "HarmBench primary judge",
                "cais/HarmBench-Llama-2-13b-cls; the historical sidecars do not record a revision",
            ),
            (
                "HarmBench cross-check judges",
                "API: gpt-4o (exact dated snapshot not recorded); open-weight: "
                "meta-llama/Llama-Guard-3-8B revision 7327bd9f6efbbe6101dc6cc4736302b3cbb6e425",
            ),
            ("XSTest API judge", "gpt-4o-2024-08-06"),
            ("HarmBench", "n=200 per alias; official HarmBench classifier primary"),
            ("XSTest", "n=250 per alias; regex primary plus three-class judge sensitivity"),
            ("MMLU", "Six-subject subset, n=300 per alias, exact match"),
            ("ARC-Challenge", "n=1,172 per alias, exact match"),
            ("Hardware", "Tesla V100-PCIE-32GB on NTU TC1"),
            ("Artifact policy", "Original raw/summary immutable; revised scoring in redacted sidecars"),
        ],
        [4.0, 11.6],
        font_size=8.4,
        body_alignment=WD_ALIGN_PARAGRAPH.LEFT,
    )
    add_body(
        doc,
        "The repository contains the pinned model revisions, full configuration schema, generated "
        "SLURM scripts, analysis programs, result sidecars, and automated verification. Raw harmful "
        "text is deliberately excluded from this document.",
    )

    p = doc.add_heading("Appendix B: Minimum Detectable Capability Effects", level=1)
    p.paragraph_format.page_break_before = True
    mde_rows = []
    with (ANALYSIS / "capability_mde.csv").open(encoding="utf-8") as handle:
        lines = handle.read().strip().splitlines()
    headers = lines[0].split(",")
    for line in lines[1:]:
        values = dict(zip(headers, line.split(",")))
        mde_rows.append(
            (
                PAIR_NAMES[values["pair_id"]],
                values["benchmark"].upper(),
                values["n"],
                f"{float(values['observed_delta']):+.4f}",
                f"{float(values['mde_80pct_power']):.4f}",
                values["observed_above_mde"],
            )
        )
    add_table(
        doc,
        "B.1",
        "Observed capability effects and 80% power minimum detectable effects",
        ["Pair", "Benchmark", "n", "Observed Δ", "MDE", "|Δ| > MDE"],
        mde_rows,
        [3.6, 2.2, 1.4, 2.8, 2.6, 3.0],
        font_size=8.0,
    )

    p = doc.add_heading("Appendix C: Human-Validation Boundaries", level=1)
    p.paragraph_format.page_break_before = True
    add_paragraphs(
        doc,
        [
            "Both human-validation sets contain 200 labelled responses and one annotator. The "
            "HarmBench sample is stratified toward disagreements, so it estimates comparative scorer "
            "fitness rather than population ASR. The XSTest sample uses three labels: full compliance, "
            "partial refusal, and full refusal.",
            "No prompt or response text is reproduced in this report. The committed validation "
            "artefacts contain aggregates and redacted identifiers. The local annotation materials are "
            "kept outside the public result surface.",
            "The XSTest blindness exclusion was applied prompt-wide instead of to the limited "
            "alias–prompt pairs specified by the pre-registration. The overlap analysis contains 186 "
            "shared items. On that intersection, the judge–regex strict-kappa gap remains approximately "
            "0.51. The deviation therefore does not explain the qualitative scorer ranking, but it "
            "prevents the audit from being described as a clean pre-registered confirmation.",
        ],
    )


def build_report() -> Path:
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)

    add_cover(doc)

    front = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(front)
    set_page_number_format(front, "lowerRoman", 1)
    set_footer_page_number(front, roman=True)
    set_running_header(front, "CCDS25-1136 · Final Year Project Report")
    add_front_matter(doc)

    main = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(main)
    set_page_number_format(main, "decimal", 1)
    set_footer_page_number(main, roman=False)
    set_running_header(main, "Benchmarking Ethical Performance of Open-Source LLMs")

    add_chapter_one(doc)
    add_chapter_two(doc)
    add_chapter_three(doc)
    add_chapter_four(doc)
    add_chapter_five(doc)
    add_chapter_six(doc)
    add_chapter_seven(doc)
    add_references(doc)
    add_appendices(doc)

    doc.core_properties.title = "Benchmarking Ethical Performance of Open-Source LLMs"
    doc.core_properties.subject = "Rewritten FYP report v1"
    doc.core_properties.author = "TAN UEI HORNG"
    doc.core_properties.comments = (
        "Corpus-informed standalone rewrite. Evidence tables are generated from the project claim "
        "registry and redacted analysis sidecars. Existing report builders are untouched."
    )
    set_update_fields_on_open(doc)
    normalize_numeric_minus_signs(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    output = build_report()
    print(output)
