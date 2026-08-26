#!/usr/bin/env python3
"""Build the user-directed interim: v4 Chapters 2–4 plus a progress chapter."""

from __future__ import annotations

import importlib.util
import sys
import tempfile
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
V4_BUILDER = ROOT / "scripts" / "build_rewritten_fyp_report_v4.py"
V1_BUILDER = ROOT / "scripts" / "build_rewritten_fyp_report_v1.py"
COMPACT_BUILDER = ROOT / "scripts" / "build_rewritten_fyp_interim_compact.py"
SOURCE = ROOT / "docs" / "rewritten_fyp_interim_v4.docx"
OUT = ROOT / "docs" / "rewritten_fyp_interim_progress.docx"
PHYSICAL_PAGE_COUNT = 33


TOC_PAGE_NUMBERS = {
    "Chapter 1: Introduction": 1,
    "1.1 Background and Motivation": 1,
    "1.2 Problem Statement": 2,
    "1.3 Research Questions": 3,
    "1.4 Scope": 3,
    "1.5 Contributions": 4,
    "Chapter 2: Literature Review": 5,
    "2.1 Trustworthiness and Safety Evaluation": 5,
    "2.2 Quantization Methods and Capability Effects": 6,
    "2.3 Quantization and Safety Behaviour": 7,
    "2.4 Refusal, Harmfulness, and Threat Models": 8,
    "2.5 Automatic Evaluators and Human Validation": 10,
    "2.6 Synthesis and Research Gap": 11,
    "Chapter 3: Methodology and Experimental Setup": 13,
    "3.1 Matched-Pair Experimental Design": 13,
    "3.2 Models and Precision Conditions": 14,
    "3.3 Benchmarks and Outcomes": 15,
    "3.4 Scoring and Validation": 16,
    "3.5 Statistical Analysis": 16,
    "3.6 Implementation and Execution": 18,
    "Chapter 4: Results": 19,
    "4.1 Results at a Glance": 19,
    "4.2 Harmful Compliance (RQ1)": 20,
    "4.3 Over-Refusal on Benign Prompts (RQ2)": 21,
    "4.4 Capability (RQ3)": 22,
    "4.5 Model Scale and Family Differences (RQ4)": 23,
    "4.6 INT8 Precision Point": 24,
    "4.7 Research-Question Synthesis": 25,
    "Chapter 5: Progress and Remaining Work": 26,
    "5.1 Progress Achieved": 26,
    "5.2 Remaining Work": 26,
    "References": 28,
}


def _load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _remove_paragraph(paragraph) -> None:
    paragraph._element.getparent().remove(paragraph._element)


def _replace_chapter_one(doc: Document, compact) -> None:
    compact._set_text(
        compact._paragraph(
            doc,
            "This project studies five matched model pairs drawn from four compact open-weight model families: "
            "Qwen3, the Llama 3 family, Phi-4-mini, and Mistral-7B [1]–[4]. Instruction-tuned models at this "
            "scale are increasingly practical to run locally, because they fit on consumer GPUs and because "
            "low-bit compression has been developed explicitly for on-device and mobile inference [5]–[7]. "
            "Local inference can keep prompts and generated responses on the user's device when no remote "
            "service is involved, but it does not by itself guarantee safe or reliable behaviour. Quantization "
            "reduces memory demand by representing model weights at lower precision, and in this study NF4 and "
            "LLM.int8 are the two lower-precision conditions. These transformations can affect more than "
            "storage: a model may preserve some task accuracy while changing how it answers harmful or benign "
            "requests. Safety and capability therefore need to be measured together rather than inferred from "
            "model size or memory savings alone.",
        ),
        "Compact open-weight language models are becoming increasingly feasible for local deployment, as "
        "low-bit quantization reduces memory requirements and enables inference on consumer GPUs, mobile "
        "devices, and other resource-constrained hardware [5]–[7]. This study examines five matched model pairs "
        "from four such families: Qwen3, Llama 3, Phi-4-mini, and Mistral-7B [1]–[4]. Local inference can keep "
        "prompts and generated responses on the user's device when no remote service is involved, but it does "
        "not by itself guarantee safe or reliable behaviour. Quantization reduces memory demand by representing "
        "model weights at lower precision, and in this study NF4 and LLM.int8 are the two lower-precision "
        "conditions. However, its effects are not limited to memory savings. A model may still retain similar "
        "task accuracy while changing how it responds to harmful or benign requests. Safety and capability "
        "therefore need to be evaluated together rather than inferred from model size or memory savings alone.",
    )

    compact._set_text(
        compact._paragraph(
            doc,
            "Quantization evaluations commonly report perplexity or benchmark accuracy [8], [9], but those "
            "measures alone do not reveal whether an apparent reduction in harmful compliance is a genuine "
            "alignment change or a loss of capability. Safety evaluation has a second complication: a model may "
            "comply less with harmful prompts while also refusing benign prompts more often, or it may appear "
            "safer because its responses have become less capable. Prior work also shows that behavioural and "
            "trustworthiness effects can depend on both the model and the quantization method [8]–[10]. This "
            "project therefore treats harmful compliance and benign over-refusal as distinct outcomes and "
            "evaluates them alongside general capability. HarmBench and XSTest provide complementary ways to "
            "measure these behaviours [11], [12]. The interpretive problem is not only to detect a change, but "
            "to determine what the change means. A controlled comparison must hold the model revision, prompts, "
            "decoding settings, and sample pairing constant. Without such controls, aggregate safety scores may "
            "be mistaken for improved alignment when they instead reflect degraded task competence or scorer "
            "behaviour. This creates a practical methodological gap for developers who need evidence that is "
            "reproducible and interpretable before deployment.",
        ),
        "Quantization evaluations commonly report perplexity or benchmark accuracy [8], [9], but those measures "
        "alone do not reveal whether an apparent reduction in harmful compliance is a genuine alignment change "
        "or a loss of capability. Safety evaluation is further complicated when a model complies less with "
        "harmful prompts while also refusing benign prompts more often or appears safer because its responses "
        "have become less capable. Prior work also shows that behavioural and trustworthiness effects can depend "
        "on both the model and the quantization method [8]–[10]. This project therefore treats harmful compliance "
        "and benign over-refusal as distinct outcomes and evaluates them alongside general capability. HarmBench "
        "and XSTest provide complementary ways to measure these behaviours [11], [12]. The interpretive problem "
        "is not only to detect a change, but to determine what the change means. A controlled comparison must "
        "hold the model revision, prompts, decoding settings, and sample pairing constant. Without such controls, "
        "aggregate safety scores may be mistaken for improved alignment when they instead reflect degraded task "
        "competence or scorer behaviour. This creates a practical methodological gap for developers who need "
        "evidence that is reproducible and interpretable before deployment.",
    )

    compact._set_text(
        compact._paragraph(
            doc,
            "To address this gap, this study evaluates matched fp16–NF4 versions of five revision-pinned model "
            "pairs, with a separate INT8 comparison at the same baseline revisions. The primary comparison "
            "applies NF4 at load time to the same weights as the fp16 baseline; INT8 supplies an additional "
            "precision point rather than being treated as the same quantization method. Outputs are evaluated on "
            "four tasks: HarmBench harmful compliance, XSTest benign over-refusal, a six-subject MMLU subset, and "
            "ARC-Challenge capability. The analysis uses paired bootstrap confidence intervals, exact McNemar "
            "tests, Benjamini–Hochberg false-discovery-rate control, and power-oriented interpretation, and scorer "
            "validity is examined with the official HarmBench classifier, an independent API judge, Llama Guard, "
            "and human-labelled samples. Section 1.4 states the boundaries of that design. The aim is to determine "
            "whether quantization-related safety changes remain when capability and scorer validity are considered "
            "together, providing a transparent basis for interpreting the deployment risks of compact open-weight "
            "models.",
        ),
        "To address this gap, this study evaluates matched fp16–NF4 versions of five revision-pinned model pairs, "
        "with a separate INT8 comparison at the same baseline revisions. The primary comparison applies NF4 at "
        "load time to the same weights as the fp16 baseline, while INT8 supplies an additional precision point "
        "rather than being treated as the same quantization method. Outputs are evaluated on four tasks: HarmBench "
        "harmful compliance, XSTest benign over-refusal, a six-subject MMLU subset, and ARC-Challenge capability. "
        "The analysis uses paired bootstrap confidence intervals, exact McNemar tests, Benjamini–Hochberg "
        "false-discovery-rate control, and power-oriented interpretation, and scorer validity is examined with the "
        "official HarmBench classifier, an independent API judge, Llama Guard, and human-labelled samples. The aim "
        "is to determine whether quantization-related safety changes remain when capability and scorer validity "
        "are considered together, thus providing a transparent basis for interpreting the deployment risks of "
        "compact open-weight models.",
    )

    compact._set_text(
        compact._paragraph(
            doc,
            "The problem addressed by this study is interpretive rather than merely operational: a change in "
            "harmful compliance cannot be labelled an alignment shift until it is read alongside benign "
            "over-refusal, capability benchmarks, and evidence that the scorer measures what it is taken to "
            "measure. Each of those three can move a safety number on its own, and none of them is visible in an "
            "attack-success rate reported by itself.",
        ),
        "The problem addressed by this study is interpretive rather than merely operational. A change in harmful "
        "compliance cannot be labelled an alignment shift until it is read alongside benign over-refusal, "
        "capability benchmarks, and evidence that the scorer measures what it is taken to measure. Each of those "
        "three can move a safety number on its own, and none of them is visible in an attack-success rate reported "
        "by itself.",
    )
    compact._set_text(
        compact._paragraph(
            doc,
            "This study therefore measures all four outcomes on the same prompts under a single controlled "
            "precision contrast, and treats the scorer as an object of study rather than as a fixed instrument. "
            "Chapter 4 reports the outcome measurements and Chapter 5 reports the scorer evidence.",
        ),
        "This study therefore measures four outcomes under a single controlled precision contrast: harmful "
        "compliance, benign over-refusal, MMLU accuracy, and ARC-Challenge accuracy. Prompts are paired between "
        "precision conditions within each benchmark, and the scorer is treated as an object of study rather than "
        "as a fixed instrument.",
    )
    compact._set_text(
        compact._paragraph(
            doc,
            "Figure 1.1 places each model pair's change in MMLU accuracy alongside its change in harmful "
            "compliance scored by the official HarmBench classifier, against the interpretation regions defined "
            "in Chapter 3. Paired bootstrap confidence intervals are shown for both measures. None of the five "
            "intervals provides evidence of an increase in harmful compliance. The Llama pair interval excludes "
            "zero only in the direction of decreased harmful compliance, while all five MMLU point estimates show "
            "capability declines. Chapter 4 reports the results after multiple-comparison correction, and Chapter "
            "5 examines whether the interpretation changes under independent scoring methods.",
        ),
        "Figure 1.1 places each model pair's change in MMLU accuracy alongside its change in harmful compliance "
        "scored by the official HarmBench classifier, against the interpretation regions defined in Chapter 3. "
        "Paired bootstrap confidence intervals are shown for both measures. None of the five intervals provides "
        "evidence of an increase in harmful compliance. The Llama pair interval excludes zero only in the direction "
        "of decreased harmful compliance, while all five MMLU point estimates show capability declines. Chapter 4 "
        "reports the results after multiple-comparison correction and summarizes whether the interpretation changes "
        "under independent scoring methods.",
    )

    compact._set_text(
        compact._paragraph(doc, "The report uses five research questions. The first three define the primary NF4 comparison. The remaining questions test whether the interpretation generalises across models, precisions, scorers, and generation conditions."),
        "The report uses five research questions. The first three define the primary NF4 comparison. The remaining questions test whether the interpretation generalises across models, precisions, scorers, and generation conditions.",
    )
    compact._set_text(compact._paragraph(doc, "RQ1. Does NF4 quantization measurably increase harmful compliance relative to matched fp16 weights?"), "RQ1. Does NF4 quantization measurably increase harmful compliance relative to matched fp16 weights?")

    paired = compact._paragraph(doc, "Paired statistical analysis. The study combines bootstrap confidence intervals, exact McNemar tests, multiplicity control, and power-oriented interpretation.")
    scorer = compact._paragraph(doc, "Scorer-validation study. HarmBench results are cross-checked using the HarmBench classifier, an independent API judge, Llama Guard, and two human-labelled samples.")
    reproducible = compact._paragraph(doc, "Reproducible benchmarking pipeline. The implementation preserves immutable raw model outputs and stores derived scorer classifications in separate, verifiable audit sidecars.")
    compact._set_text(paired, "Scorer-validation study. HarmBench results are cross-checked using the HarmBench classifier, an independent API judge, Llama Guard, and two human-labelled samples.")
    compact._set_text(scorer, "Paired statistical analysis. The study combines bootstrap confidence intervals, exact McNemar tests, multiplicity control, and power-oriented interpretation.")
    _remove_paragraph(reproducible)
    compact._remove_range(doc, "1.6 Report Structure", "Chapter 2: Literature Review")


def _apply_user_scope_cuts(doc: Document, compact) -> None:
    """Reproduce the subsections and implementation detail removed in Word."""
    for text in (
        "RQ1. Does NF4 quantization measurably increase harmful compliance relative to matched fp16 weights?",
        "RQ2. Does NF4 quantization increase over-refusal on benign prompts?",
        "RQ3. Does NF4 quantization reduce capability on MMLU or ARC-Challenge?",
        "RQ4. How do the effects vary across model families, model scales, and the additional INT8 precision condition?",
        "RQ5. Does the substantive conclusion remain credible under alternative scorers, comparison with human annotations, generation budgets, and decoding seeds?",
        "The command-line interface exposes single-run, matrix, analysis, and agent-status paths. The dashboard reads the stored results but does not alter them, and configurations created through the dashboard are written outside the primary configuration. This boundary prevents an interface action from changing the results used in the report.",
        "Testing covers configuration validation, model reuse, result resumption, refusal parsing, pairwise analysis, judge redaction, report claims, and artifact policy. Passing tests show that the implemented checks behave as specified. They do not, however, substitute for scientific review, visual inspection, or source verification.",
    ):
        _remove_paragraph(compact._paragraph(doc, text))

    compact._remove_range(
        doc,
        "2.7 Implications for Experimental Design",
        "Chapter 3: Methodology and Experimental Setup",
    )
    compact._remove_range(doc, "3.7 Reproducibility Controls", "Chapter 4: Results")


def _insert_progress_chapter(doc: Document, v1, compact, evidence: dict) -> None:
    chapter = compact._paragraph(doc, "Chapter 7: Conclusion and Future Work")
    compact._set_text(chapter, "Chapter 5: Progress and Remaining Work")
    references = compact._paragraph(doc, "References")
    current = chapter._element.getnext()
    while current is not None and current is not references._element:
        next_el = current.getnext()
        current.getparent().remove(current)
        current = next_el
    if current is None:
        raise AssertionError("References heading did not follow the former Chapter 7")

    references.insert_paragraph_before("5.1 Progress Achieved", style="Heading 2")
    references.insert_paragraph_before(
        "At the interim report stage, the literature review, methodology, experimental setup, and primary "
        "results chapters have been completed. The experimental matrix is also complete.",
        style="Normal",
    )
    references.insert_paragraph_before(
        "The primary evidence is sufficiently developed for the interim report to state its central result. At "
        "the 512-token reference budget, no harmful-compliance contrast survives Benjamini–Hochberg "
        "false-discovery-rate correction. The corrected family contains three survivors: the Qwen3-1.7B MMLU "
        "decrease, the Llama-3.2-3B ARC decrease, and the Phi-4-mini regex-scored over-refusal decrease. The last "
        "result is scorer-dependent because it is not reproduced by the semantic refusal judge and the human "
        "audit favours that judge over the regex. In the harmful-compliance human-validation sample, the official "
        "classifier reaches Cohen's κ = 0.59, compared with κ = 0.11 for the refusal regex.",
        style="Normal",
    )
    compact._add_table_before(
        doc,
        references,
        v1,
        "5.1",
        "Completed work and repository evidence",
        ["Workstream", "Status", "Primary repository evidence"],
        [
            ("Literature review", "Complete", "Retained Chapter 2 through Section 2.6 and source-entailment audit"),
            ("Methodology and setup", "Complete", "Revision-pinned configuration; Chapters 3.1–3.6"),
            ("Experiment matrix", "Complete", "results_512/<alias>/<benchmark>/summary.json (60 summaries)"),
            ("Primary paired analysis", "Complete", "multiple_comparisons.json; pairwise_deltas.json"),
            ("Validation and robustness evidence", "Complete", "judge, human-validation, generation-length, precision and multi-seed sidecars"),
        ],
        [3.6, 2.5, 9.5],
        font_size=7.6,
    )

    references.insert_paragraph_before("5.2 Remaining Work", style="Heading 2")
    references.insert_paragraph_before(
        "The remaining work is the detailed final-report writing for two chapters whose underlying analyses are "
        "already complete. The Scorer Validation and Robustness chapter will organise the classifier-versus-proxy "
        "comparison, independent API and Llama Guard judgements, both human-validation studies, generation-length "
        "analysis, multi-seed sensitivity, precision comparison, and power interpretation into one coherent "
        "account.",
        style="Normal",
    )
    references.insert_paragraph_before(
        "The Discussion chapter will then connect the primary results to the literature, explain the capability "
        "and deployment implications, assess the validity of the scoring methods, state the internal, construct, "
        "and external validity boundaries, and answer the central interpretive question. Its purpose will be to "
        "integrate the evidence already reported and validated, rather than introduce a new result after the "
        "interim stage.",
        style="Normal",
    )
    compact._add_table_before(
        doc,
        references,
        v1,
        "5.2",
        "Remaining final-report writing",
        ["Planned chapter", "Evidence already complete", "Writing still required"],
        [
            (
                "Scorer Validation and Robustness",
                "Alternative judges, human validation, budget, seed, precision and power analyses",
                "Organise and interpret the completed validation and sensitivity evidence",
            ),
            (
                "Discussion",
                "Primary deltas, corrected tests, capability results and validation findings",
                "Synthesize the central answer, prior-work relationship, validity boundaries and implications",
            ),
        ],
        [4.0, 6.2, 5.4],
        font_size=7.6,
    )


def _bookmark_headings(doc: Document) -> dict[str, str]:
    bookmarks = {}
    next_id = 1000
    for paragraph in doc.paragraphs:
        if not paragraph.style or paragraph.style.name not in {"Heading 1", "Heading 2"}:
            continue
        title = paragraph.text.strip()
        name = f"toc_{next_id}"
        start = OxmlElement("w:bookmarkStart")
        start.set(qn("w:id"), str(next_id))
        start.set(qn("w:name"), name)
        end = OxmlElement("w:bookmarkEnd")
        end.set(qn("w:id"), str(next_id))
        insert_at = 1 if paragraph._p.pPr is not None else 0
        paragraph._p.insert(insert_at, start)
        paragraph._p.append(end)
        bookmarks[title] = name
        next_id += 1
    return bookmarks


def _remove_existing_toc_bookmarks(doc: Document) -> None:
    ids = set()
    for start in list(doc.element.body.iter(qn("w:bookmarkStart"))):
        if (start.get(qn("w:name")) or "").startswith("toc_"):
            ids.add(start.get(qn("w:id")))
            start.getparent().remove(start)
    for end in list(doc.element.body.iter(qn("w:bookmarkEnd"))):
        if end.get(qn("w:id")) in ids:
            end.getparent().remove(end)


def _rebuild_linked_toc(doc: Document, v1, compact) -> None:
    toc_heading = compact._paragraph(doc, "Table of Contents")
    current = toc_heading._element.getnext()
    boundary = None
    while current is not None:
        if current.tag == qn("w:p") and current.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            boundary = current
            break
        next_el = current.getnext()
        current.getparent().remove(current)
        current = next_el
    if boundary is None:
        raise AssertionError("Could not locate the section break after the TOC")

    update_fields = doc.settings.element.find(qn("w:updateFields"))
    if update_fields is not None:
        doc.settings.element.remove(update_fields)
    _remove_existing_toc_bookmarks(doc)
    bookmarks = _bookmark_headings(doc)

    for heading in [p for p in doc.paragraphs if p.style and p.style.name in {"Heading 1", "Heading 2"}]:
        title = heading.text.strip()
        if title not in TOC_PAGE_NUMBERS:
            raise AssertionError(f"No TOC page recorded for {title!r}")
        level = 0 if heading.style.name == "Heading 1" else 1
        before = list(doc.element.body.iterchildren())
        paragraph = v1.add_contents_entry(
            doc, title, str(TOC_PAGE_NUMBERS[title]), level=level, bold=level == 0
        )
        first_run = paragraph.runs[0]._r
        paragraph._p.remove(first_run)
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("w:anchor"), bookmarks[title])
        hyperlink.set(qn("w:history"), "1")
        hyperlink.append(first_run)
        insert_at = 1 if paragraph._p.pPr is not None else 0
        paragraph._p.insert(insert_at, hyperlink)
        old = set(before)
        added = [
            child
            for child in doc.element.body.iterchildren()
            if child not in old and child.tag != qn("w:sectPr")
        ]
        if len(added) != 1:
            raise AssertionError(f"Expected one TOC entry block, found {len(added)}")
        boundary.addprevious(added[0])


def _assert_structure(doc: Document) -> None:
    headings = [p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    required = {
        "Chapter 1: Introduction",
        "1.5 Contributions",
        "Chapter 2: Literature Review",
        "Chapter 3: Methodology and Experimental Setup",
        "3.6 Implementation and Execution",
        "Chapter 4: Results",
        "4.7 Research-Question Synthesis",
        "Chapter 5: Progress and Remaining Work",
        "5.1 Progress Achieved",
        "5.2 Remaining Work",
        "References",
    }
    missing = sorted(required - set(headings))
    if missing:
        raise AssertionError(f"Missing required headings: {missing}")
    forbidden = {
        "1.6 Report Structure",
        "Chapter 5: Scorer Validation and Robustness",
        "Chapter 6: Discussion",
        "Chapter 7: Conclusion and Future Work",
        "2.7 Implications for Experimental Design",
        "3.7 Reproducibility Controls",
        "3.8 Analysis Workflow and Quality Control",
        "Appendix A: Reproducibility Summary",
    }
    residue = sorted(forbidden & set(headings))
    if residue:
        raise AssertionError(f"Removed chapters remain: {residue}")
    text = "\n".join(p.text for p in doc.paragraphs)
    for stale in ("Chapter 5 shows", "Appendix A", "Appendix B", "Appendix C", "Table A.", "Table 5.3", "Table 6.1"):
        if stale in text:
            raise AssertionError(f"Dangling removed-section reference: {stale!r}")


def _mark_table_header_rows(doc: Document) -> None:
    """Expose every table's first row as a header to Word and screen readers."""
    for table in doc.tables:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:tblHeader")) is None:
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)


def repair_current_document(path: Path = OUT) -> Path:
    """Repair navigation and citation identity after the student's Word cuts."""
    v4 = _load_module(V4_BUILDER, "build_rewritten_fyp_report_v4_for_progress_repair")
    v1 = _load_module(V1_BUILDER, "build_rewritten_fyp_report_v1_for_progress_repair")
    compact = _load_module(COMPACT_BUILDER, "build_rewritten_fyp_interim_compact_for_progress_repair")
    doc = Document(path)

    repairs = (
        (
            "GPTQ / AWQ / SmoothQuant [5], [6], [22]",
            "GPTQ / AWQ / SmoothQuant [2], [1], [22]",
        ),
        (
            "HarmBench results are cross-checked using the HarmBench classifier, an independent API judge, Llama Guard, and two human-labelled samples.",
            "HarmBench and XSTest scoring are cross-checked using the HarmBench classifier, an independent API judge, Llama Guard, and benchmark-specific human-labelled samples.",
        ),
        (
            "Full v4 Chapter 2 and source-entailment audit",
            "Retained Chapter 2 through Section 2.6 and source-entailment audit",
        ),
        (
            "Revision-pinned configuration; Chapters 3.1–3.8",
            "Revision-pinned configuration; Chapters 3.1–3.6",
        ),
        (
            "harmful complicance relative to matched fp16 weights?",
            "harmful compliance relative to matched fp16 weights?",
        ),
        (
            "Table A.1 records the quantization settings and the declared dependency constraints.",
            "The quantization settings and declared dependency constraints are specified in the "
            "versioned configuration and requirements files used for every run.",
        ),
    )
    for old, new in repairs:
        hits = compact._replace_all_text(doc, old, new)
        if hits > 1:
            raise AssertionError(f"Expected at most one manual-document repair, found {hits}: {old!r}")
        if hits == 0:
            paragraphs = list(doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        paragraphs.extend(cell.paragraphs)
            if not any(new in paragraph.text for paragraph in paragraphs) and old.startswith(
                ("GPTQ / AWQ", "HarmBench results", "harmful complicance")
            ):
                raise AssertionError(f"Manual-document repair target is absent in both forms: {old!r}")

    _rebuild_linked_toc(doc, v1, compact)
    _mark_table_header_rows(doc)
    _assert_structure(doc)
    citation_summary = v4._renumber_citations(doc)
    if citation_summary["dropped"] not in ([58, 59, 60, 61], []):
        raise AssertionError(
            f"Expected the four newly uncited tail references to be dropped, got {citation_summary['dropped']}"
        )
    v4._assert_ieee_order(doc)
    doc.core_properties.modified = datetime(2026, 8, 25)
    doc.core_properties.last_modified_by = "TAN UEI HORNG"
    doc.core_properties.comments = (
        "Student-edited interim with deleted subsections preserved. TOC and citation identities repaired "
        "post-edit; the removed-appendix Table A.1 cross-reference now points to the versioned config "
        "and requirements files."
    )
    doc.save(path)
    compact.PHYSICAL_PAGE_COUNT = PHYSICAL_PAGE_COUNT
    compact._refresh_package_metadata(path, doc)
    return path


def build_interim() -> Path:
    v4 = _load_module(V4_BUILDER, "build_rewritten_fyp_report_v4_for_progress")
    v1 = _load_module(V1_BUILDER, "build_rewritten_fyp_report_v1_for_progress")
    compact = _load_module(COMPACT_BUILDER, "build_rewritten_fyp_interim_compact_for_progress")
    evidence = compact._load_progress_evidence()
    source_hash = compact._sha256(SOURCE)

    with tempfile.TemporaryDirectory(prefix="fyp_interim_progress_") as temporary:
        generated = Path(temporary) / "v4_report.docx"
        original_out = v4.OUT
        try:
            v4.OUT = generated
            v4.build_report()
        finally:
            v4.OUT = original_out

        doc = Document(generated)
        compact._retitle(doc)
        compact._remove_range(doc, "Acknowledgements", "Abstract")
        compact._remove_front_lists(doc)
        _replace_chapter_one(doc, compact)
        _apply_user_scope_cuts(doc, compact)

        compact._remove_range(
            doc,
            "Chapter 5: Scorer Validation and Robustness",
            "Chapter 7: Conclusion and Future Work",
        )
        _insert_progress_chapter(doc, v1, compact, evidence)
        compact._remove_appendices(doc)

        replacements = (
            (
                "This chapter reports the primary matched-pair results before the validation and sensitivity analyses.",
                "This chapter reports the primary matched-pair results and summarizes the completed validation and sensitivity evidence in Section 4.7.",
            ),
            (
                "The sole surviving over-refusal effect is a decrease for Phi-4-mini, which Chapter 5 shows to be most plausibly an artifact of keyword regex scoring, as it does not survive semantic evaluation.",
                "The sole surviving over-refusal effect is a decrease for Phi-4-mini. The scorer-sensitivity evidence summarized in Section 4.7 identifies it as most plausibly a keyword-regex measurement artifact because it does not survive semantic evaluation.",
            ),
            (
                "The second requires compatible GPU hardware, accepted gated-model conditions, cached datasets, and a software environment satisfying the dependency constraints in Table A.1 of Appendix A.",
                "The second requires compatible GPU hardware, accepted gated-model conditions, cached datasets, and a software environment satisfying the revision-pinned configuration and recorded dependency constraints.",
            ),
            (
                "Appendix B therefore reports the thresholds rather than converting all non-significant effects into claims of equivalence.",
                "The power analysis therefore treats smaller non-significant effects as unresolved rather than converting them into claims of equivalence.",
            ),
            (
                "Table A.1 records the quantization settings and the declared dependency constraints.",
                "The quantization settings and declared dependency constraints are specified in the "
                "versioned configuration and requirements files used for every run.",
            ),
        )
        for old, new in replacements:
            hits = compact._replace_all_text(doc, old, new)
            if hits != 1:
                raise AssertionError(f"Expected one cross-reference repair, found {hits}: {old!r}")

        citation_summary = v4._renumber_citations(doc)
        table_citation_hits = compact._replace_all_text(
            doc,
            "GPTQ / AWQ / SmoothQuant [1], [2], [22]",
            "GPTQ / AWQ / SmoothQuant [2], [1], [22]",
        )
        if table_citation_hits != 1:
            raise AssertionError(
                f"Expected one GPTQ/AWQ/SmoothQuant table citation repair, found {table_citation_hits}"
            )
        validation_wording_hits = compact._replace_all_text(
            doc,
            "HarmBench results are cross-checked using the HarmBench classifier, an independent API judge, Llama Guard, and two human-labelled samples.",
            "HarmBench and XSTest scoring are cross-checked using the HarmBench classifier, an independent API judge, Llama Guard, and benchmark-specific human-labelled samples.",
        )
        if validation_wording_hits != 1:
            raise AssertionError(
                f"Expected one human-validation scope repair, found {validation_wording_hits}"
            )
        v4._assert_ieee_order(doc)
        _rebuild_linked_toc(doc, v1, compact)
        _assert_structure(doc)

        doc.core_properties.title = "Benchmarking Ethical Performance of Open-Source LLMs"
        doc.core_properties.subject = "Interim report: completed Chapters 1–4 and remaining report writing"
        doc.core_properties.author = "TAN UEI HORNG"
        doc.core_properties.last_modified_by = "TAN UEI HORNG"
        doc.core_properties.created = datetime(2026, 8, 24)
        doc.core_properties.modified = datetime(2026, 8, 24)
        doc.core_properties.comments = (
            "User-directed interim derived non-destructively from the verified v4 report. Chapters 2–4 are "
            "retained from v4; Chapter 5 records completed progress and the remaining writing of the detailed "
            "validation/robustness and discussion chapters."
        )
        OUT.parent.mkdir(parents=True, exist_ok=True)
        doc.save(OUT)
        compact.PHYSICAL_PAGE_COUNT = PHYSICAL_PAGE_COUNT
        compact._refresh_package_metadata(OUT, doc)

    if compact._sha256(SOURCE) != source_hash:
        raise AssertionError("The verified v4 interim changed during the progress build")
    print(
        f"    progress interim citations: {citation_summary['kept']} references; "
        f"dropped uncited references: {citation_summary['dropped'] or 'none'}"
    )
    return OUT


if __name__ == "__main__":
    print(build_interim())
