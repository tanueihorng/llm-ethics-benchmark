#!/usr/bin/env python3
"""Build a compact, progress-centred interim report from the verified v4 line.

The comprehensive v4 report remains untouched.  This builder creates a separate
interim document by retaining the load-bearing introduction, method, results,
validation and limitations blocks, removing final-report detail, and replacing
the conclusion with a repository-grounded ``Progress and Remaining Work``
chapter.  Citations are re-enumerated after compression so the retained
bibliography remains IEEE first-use ordered.
"""

from __future__ import annotations

import hashlib
import importlib.util
import json
import re
import sys
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
V4_BUILDER = ROOT / "scripts" / "build_rewritten_fyp_report_v4.py"
V1_BUILDER = ROOT / "scripts" / "build_rewritten_fyp_report_v1.py"
PROJECT_LOG = ROOT / "docs" / "PROJECT_LOG.md"
OUT = ROOT / "docs" / "rewritten_fyp_interim_compact.docx"
PHYSICAL_PAGE_COUNT = 36

COVER_LINE = ("Final Year Project Report", "Interim Report")
RUNNING_HEADER = (
    "CCDS25-1136 · Final Year Project Report",
    "CCDS25-1136 · Interim Report",
)

# Verified numbered pages for the compact line. Keep the compact TOC static:
# Word mis-resolves the generated multiline TOC field and replaces every entry
# with page 1 when fields are refreshed. The stable page map is regenerated from
# source and checked against the rendered artifact instead.
TOC_PAGE_NUMBERS = {
    "Chapter 1: Introduction": 1,
    "1.1 Background and Motivation": 1,
    "1.2 Problem Statement": 2,
    "1.3 Research Questions": 3,
    "1.4 Scope": 3,
    "Chapter 2: Literature Review": 5,
    "2.1 Quantization Methods and Capability Effects": 5,
    "2.2 Quantization and Safety Behaviour": 6,
    "2.3 Synthesis and Research Gap": 7,
    "Chapter 3: Methodology and Experimental Setup": 9,
    "3.1 Matched-Pair Experimental Design": 9,
    "3.2 Models and Precision Conditions": 10,
    "3.3 Benchmarks and Outcomes": 10,
    "3.4 Scoring and Validation": 11,
    "3.5 Statistical Analysis": 12,
    "Chapter 4: Results": 14,
    "4.1 Results at a Glance": 14,
    "4.2 Harmful Compliance (RQ1)": 15,
    "4.3 Over-Refusal on Benign Prompts (RQ2)": 16,
    "4.4 Capability (RQ3)": 17,
    "4.5 Model Scale and Family Differences (RQ4)": 18,
    "4.6 INT8 Precision Point": 19,
    "4.7 Research-Question Synthesis": 20,
    "Chapter 5: Scorer Validation and Robustness": 21,
    "5.1 HarmBench Classifier versus Refusal Proxy": 21,
    "5.2 Independent HarmBench Judges": 21,
    "5.3 Human Validation of Harmful-Compliance Scorers": 22,
    "5.4 XSTest Scorer Sensitivity and Human Validation": 23,
    "5.5 Generation-Length Robustness": 24,
    "5.6 Power and the Meaning of the Null": 24,
    "5.7 Robustness Synthesis": 25,
    "Chapter 6: Discussion": 26,
    "6.1 Answer to the Central Question": 26,
    "6.2 Threats to Validity": 26,
    "6.3 Practical Implications": 28,
    "Chapter 7: Progress and Remaining Work": 29,
    "7.1 Progress Achieved": 29,
    "7.2 Problems Encountered and Resolutions": 29,
    "7.3 Remaining Work": 30,
    "7.4 Submission Timeline": 30,
    "References": 31,
}


def _load_module(path: Path, name: str):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _paragraph(doc: Document, exact: str):
    matches = [paragraph for paragraph in doc.paragraphs if paragraph.text == exact]
    if len(matches) != 1:
        raise AssertionError(f"Expected one paragraph {exact!r}, found {len(matches)}")
    return matches[0]


def _set_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def _remove_range(doc: Document, start: str, end: str, *, include_start: bool = True) -> None:
    """Remove body blocks from ``start`` up to but not including ``end``."""
    start_el = _paragraph(doc, start)._element
    end_el = _paragraph(doc, end)._element
    current = start_el if include_start else start_el.getnext()
    removed = 0
    while current is not None and current is not end_el:
        next_el = current.getnext()
        current.getparent().remove(current)
        current = next_el
        removed += 1
    if current is None:
        raise AssertionError(f"End marker {end!r} did not follow {start!r}")
    if removed == 0:
        raise AssertionError(f"Compression range {start!r} -> {end!r} removed nothing")


def _remove_front_lists(doc: Document) -> None:
    """Drop figures/tables/abbreviations while preserving the main-section break."""
    start = _paragraph(doc, "List of Figures")._element
    current = start
    removed = 0
    while current is not None:
        if current.tag == qn("w:p") and current.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            break
        next_el = current.getnext()
        current.getparent().remove(current)
        current = next_el
        removed += 1
    if current is None or removed < 3:
        raise AssertionError("Could not preserve the main-document section break while removing front lists")


def _remove_appendices(doc: Document) -> None:
    current = _paragraph(doc, "Appendix A: Reproducibility Summary")._element
    removed = 0
    while current is not None and current.tag != qn("w:sectPr"):
        next_el = current.getnext()
        current.getparent().remove(current)
        current = next_el
        removed += 1
    if removed == 0:
        raise AssertionError("Appendix removal matched no body blocks")


def _remove_figure(doc: Document, caption_text: str) -> None:
    """Remove a figure drawing paragraph and its following caption."""
    caption = _paragraph(doc, caption_text)._element
    drawing = caption.getprevious()
    if drawing is None or drawing.find(".//" + qn("w:drawing")) is None:
        raise AssertionError(f"Expected a drawing immediately before {caption_text!r}")
    drawing.getparent().remove(drawing)
    caption.getparent().remove(caption)


def _rename_heading(doc: Document, old: str, new: str) -> None:
    paragraph = _paragraph(doc, old)
    if not paragraph.style.name.startswith("Heading"):
        raise AssertionError(f"Refusing to rename non-heading paragraph {old!r}")
    _set_text(paragraph, new)


def _replace_all_text(doc: Document, old: str, new: str) -> int:
    hits = 0
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for paragraph in paragraphs:
        if old in paragraph.text:
            _set_text(paragraph, paragraph.text.replace(old, new))
            hits += 1
    return hits


def _retitle(doc: Document) -> None:
    cover_hits = 0
    header_hits = 0
    for paragraph in doc.paragraphs:
        if COVER_LINE[0] in paragraph.text:
            _set_text(paragraph, paragraph.text.replace(*COVER_LINE))
            cover_hits += 1
    for section in doc.sections:
        for header in (section.header, section.first_page_header, section.even_page_header):
            if header is None:
                continue
            for paragraph in header.paragraphs:
                if RUNNING_HEADER[0] in paragraph.text:
                    _set_text(paragraph, paragraph.text.replace(*RUNNING_HEADER))
                    header_hits += 1
    if cover_hits != 1 or header_hits < 1:
        raise AssertionError(f"Retitle failed: cover={cover_hits}, headers={header_hits}")


def _load_progress_evidence() -> dict:
    summaries = sorted(ROOT.glob("results_512/*/*/summary.json"))
    if len(summaries) != 60:
        raise AssertionError(f"Expected 60 primary benchmark summaries, found {len(summaries)}")

    comparisons_path = ROOT / "results_512/analysis/multiple_comparisons.json"
    comparisons = json.loads(comparisons_path.read_text(encoding="utf-8"))
    tests = comparisons.get("tests", comparisons if isinstance(comparisons, list) else [])
    if not isinstance(tests, list):
        # The current artifact stores the rows under ``comparisons``.
        tests = comparisons.get("comparisons", [])
    survivors = [row for row in tests if row.get("bh_significant") or row.get("significant_bh")]
    if not survivors:
        # Fail closed on schema drift rather than silently inventing a count.
        expected = comparisons.get("n_bh_significant_q05")
        if expected is None and isinstance(comparisons.get("bh_survivors"), list):
            expected = len(comparisons["bh_survivors"])
        if expected != 3:
            raise AssertionError("Could not derive the three BH-FDR survivors from multiple_comparisons.json")
        survivor_count = expected
    else:
        survivor_count = len(survivors)
    if survivor_count != 3:
        raise AssertionError(f"Expected three BH-FDR survivors, found {survivor_count}")

    required = [
        ROOT / "results_512/analysis/pairwise_deltas.json",
        ROOT / "results_512/analysis/human_validation.json",
        ROOT / "results_512/analysis/xstest_human_validation.json",
        ROOT / "results_512/analysis/judge_pairwise_agreement.json",
        ROOT / "results_512/analysis/judge_pairwise_agreement_llamaguard.json",
        ROOT / "results_512/analysis/genlen_robustness.json",
        ROOT / "results_512/analysis/sensitivity_multiseed.json",
        ROOT / "results_512/analysis/precision_sweep.json",
        ROOT / "results_512/analysis/refusal_margin.json",
    ]
    missing = [str(path.relative_to(ROOT)) for path in required if not path.exists()]
    if missing:
        raise AssertionError(f"Progress evidence is missing: {missing}")

    multiseed = json.loads((ROOT / "results_512/analysis/sensitivity_multiseed.json").read_text(encoding="utf-8"))
    per_pair = {row["pair_id"]: row for row in multiseed["per_pair"]}
    qwen4 = per_pair["qwen_4b"]["judge_delta"]
    llama = per_pair["llama_3_2_3b"]["judge_delta"]
    qwen4_positive = sum(row["delta"] > 0 for row in qwen4["per_seed"])
    llama_negative = sum(row["delta"] < 0 for row in llama["per_seed"])
    if (qwen4_positive, qwen4["n_seeds_significant_p05"]) != (5, 1):
        raise AssertionError("Qwen3-4B multi-seed evidence no longer supports 5/5 positive, 1/5 p<.05")
    if (llama_negative, llama["n_seeds_significant_p05"]) != (5, 2):
        raise AssertionError("Llama multi-seed evidence no longer supports 5/5 negative, 2/5 p<.05")

    genlen = json.loads((ROOT / "results_512/analysis/genlen_robustness.json").read_text(encoding="utf-8"))
    per_model = genlen["prefix_truncation_128"]["per_model"]
    llama_truncation = sorted(
        per_model[alias]["pct_truncated"]
        for alias in ("llama_3_2_3b_base", "llama_3_2_3b_4bit")
    )
    if llama_truncation != [3.0, 4.0]:
        raise AssertionError(f"Llama truncation evidence changed: {llama_truncation}")

    refusal_margin = json.loads((ROOT / "results_512/analysis/refusal_margin.json").read_text(encoding="utf-8"))
    pooled_flips = refusal_margin["pooled"]
    if (pooled_flips["n_harmful_ward"], pooled_flips["n_safe_ward"]) != (50, 42):
        raise AssertionError("Refusal-margin evidence no longer supports 50 harmful-ward / 42 safe-ward flips")

    project_log = PROJECT_LOG.read_text(encoding="utf-8")
    required_log_claims = [
        "T3. Run `MyTCinfo` on TC1",
        "T61. Capture the TC1 conda environment once",
        "T30b, optional strengthening",
        "T33. Residual pre-submission code hardening",
        "DUE Mon 19 Oct 2026",
        "DUE Fri 13 Nov 2026",
        "Oral presentation — Fri 4 / Mon 7 / Tue 8 Dec 2026",
    ]
    absent = [claim for claim in required_log_claims if claim not in project_log]
    if absent:
        raise AssertionError(f"PROJECT_LOG no longer supports progress/timeline claims: {absent}")

    return {
        "summary_count": len(summaries),
        "survivor_count": survivor_count,
        "qwen4_positive_seeds": qwen4_positive,
        "qwen4_significant_seeds": qwen4["n_seeds_significant_p05"],
        "llama_negative_seeds": llama_negative,
        "llama_significant_seeds": llama["n_seeds_significant_p05"],
        "llama_truncation_low": int(llama_truncation[0]),
        "llama_truncation_high": int(llama_truncation[1]),
        "harmful_ward_flips": pooled_flips["n_harmful_ward"],
        "safe_ward_flips": pooled_flips["n_safe_ward"],
    }


def _apply_correctness_audit_repairs(doc: Document, evidence: dict) -> None:
    """Apply narrow, fail-closed repairs found by the compact-artifact audit."""
    repairs = (
        (
            "), which is what a truncation artifact would predict.",
            "). This budget-sensitive non-replication is attributed to cross-run greedy divergence rather than "
            f"truncation: Llama has the study's lowest 128-token truncation rate "
            f"({evidence['llama_truncation_low']}–{evidence['llama_truncation_high']}%) and the largest "
            "prefix-mismatch count.",
        ),
        (
            "Human ground-truth annotation clarifies this discrepancy.",
            "Human annotation clarifies this discrepancy.",
        ),
        (
            "It reproduces the direction and statistical conclusion for each model pair: no NF4 increase "
            "in harmful compliance is significant after correction.",
            "It reproduces the statistical conclusion for every pair and the direction for the four "
            "non-flat classifier contrasts; on Qwen3-1.7B the classifier is flat while the API delta is "
            "+0.005. No NF4 increase in harmful compliance is significant after correction.",
        ),
        (
            "The broad refusal mapping includes partial refusals alongside complete rejections. While this "
            "increases the semantic judge's agreement with human annotators, the keyword regex remains near "
            "zero because it cannot recognise non-formulaic refusals. The semantic judge is therefore "
            "substantially better aligned with human annotators under both evaluation standards.",
            "The broad refusal mapping includes partial refusals alongside complete rejections. While this "
            "increases the semantic judge's agreement with the single annotator, the keyword regex remains "
            "near zero because it cannot recognise non-formulaic refusals. The semantic judge is therefore "
            "substantially better aligned with this annotator under both evaluation standards.",
        ),
        (
            "The regex captures only 3.2% of true full refusals compared to 96.8% for the semantic judge, "
            "missing unconventional refusal phrasing that the judge accurately identifies.",
            "Against the single annotator, the regex recalls 3.2% of labelled full refusals, compared with "
            "96.8% for the semantic judge; the difference is consistent with the regex missing refusal "
            "phrasing outside its keyword patterns.",
        ),
        (
            "and no decoding seed in the multi-seed arm shows a consistent rise in harmful compliance.",
            f"and the five-seed arm shows pair-specific rather than common movement: Qwen3-4B is positive "
            f"in {evidence['qwen4_positive_seeds']}/5 seeds "
            f"({evidence['qwen4_significant_seeds']} individually significant), whereas Llama is negative "
            f"in {evidence['llama_negative_seeds']}/5 "
            f"({evidence['llama_significant_seeds']} individually significant).",
        ),
        (
            "Five-seed arm has no stable increase in harmful compliance",
            f"Qwen3-4B positive in {evidence['qwen4_positive_seeds']}/5 seeds "
            f"({evidence['qwen4_significant_seeds']} p<.05); Llama negative in "
            f"{evidence['llama_negative_seeds']}/5 ({evidence['llama_significant_seeds']} p<.05)",
        ),
        (
            "The empirical evidence demonstrates the absence of a statistically significant increase in "
            "harmful compliance after correction within the detection limits of the experimental design,",
            "Within the detection limits of the experimental design, no increase in harmful compliance "
            "reaches statistical significance after correction,",
        ),
        (
            "A more accurate characterisation of low-bit quantization is prompt-level decision boundary "
            "instability accompanied by capability degradation. Small weight perturbations shift borderline "
            "outputs in both directions without generating systematic aggregate alignment failure.",
            "The post hoc refusal-margin analysis is consistent with model-specific prompt-level boundary "
            "instability, but it does not establish a universal mechanism. Across the pooled flip set, "
            f"{evidence['harmful_ward_flips']} prompts move harmful-ward and "
            f"{evidence['safe_ward_flips']} safe-ward; the within-pair predictive strength varies substantially. "
            "These diagnostic results support bidirectional, model-dependent instability rather than targeted "
            "alignment erosion.",
        ),
        (
            "Qwen3-4B rises from 0.115 to 0.125 to 0.155, but neither increase is significant after correction.",
            "Qwen3-4B rises from 0.115 to 0.125 to 0.155. The INT8 step is non-significant on its separate "
            "uncorrected McNemar test, while the NF4 contrast does not survive correction in the primary family.",
        ),
        (
            "+0." "055. the matched 512-token rerun",
            "+0." "055. The matched 512-token rerun",
        ),
        (
            "Two repository actions remain useful before the final freeze: record the exact TC1 storage "
            "quota with MyTCinfo, and, if time permits, add a second annotator",
            "The main repository-strengthening actions before the final freeze are to record the exact TC1 "
            "storage quota with MyTCinfo, export the fyp-tc1 conda environment, and, if time permits, add a "
            "second annotator",
        ),
        (
            "Address feedback; record TC1 quota; complete optional second-annotator overlap if feasible",
            "Address feedback; record TC1 quota and environment; complete optional second-annotator overlap "
            "if feasible",
        ),
        (
            "Broader quantization methods, additional model scales, jailbreak evaluations, and wider "
            "stochastic-decoding work remain future-study extensions rather than evidence claimed here.",
            "Broader quantization methods, additional model scales, jailbreak evaluations, and wider "
            "stochastic-decoding work remain future-study extensions rather than evidence claimed here. This "
            "section is the supervisor-facing scientific subset of the backlog, not an exhaustive task list; "
            "PROJECT_LOG.md separately tracks correspondence, Overleaf delivery, and internal gate hardening.",
        ),
    )
    for old, new in repairs:
        hits = _replace_all_text(doc, old, new)
        if hits != 1:
            raise AssertionError(f"Correctness-audit repair expected one hit, found {hits}: {old!r}")


def _rewrite_colon_heavy_prose(doc: Document) -> None:
    """Replace habitual claim-colon-explanation phrasing with academic prose."""
    rewrites = (
        (
            "These transformations can affect more than storage: a model may preserve some task accuracy "
            "while changing how it answers harmful or benign requests.",
            "Their effects can extend beyond storage. A model may preserve some task accuracy while "
            "changing how it answers harmful or benign requests.",
        ),
        (
            "Safety evaluation has a second complication: a model may comply less with harmful prompts "
            "while also refusing benign prompts more often, or it may appear safer because its responses "
            "have become less capable.",
            "Safety evaluation is further complicated when a model complies less with harmful prompts "
            "while also refusing benign prompts more often, or appears safer because its responses have "
            "become less capable.",
        ),
        (
            "The problem addressed by this study is interpretive rather than merely operational: a change "
            "in harmful compliance cannot be labelled an alignment shift until it is read alongside benign "
            "over-refusal, capability benchmarks, and evidence that the scorer measures what it is taken "
            "to measure.",
            "The problem addressed by this study is interpretive rather than merely operational. A change "
            "in harmful compliance cannot be labelled an alignment shift until it is read alongside benign "
            "over-refusal, capability benchmarks, and evidence that the scorer measures what it is taken "
            "to measure.",
        ),
        (
            "This study therefore measures all four outcomes on the same prompts under a single controlled "
            "precision contrast, and treats the scorer as an object of study rather than as a fixed instrument.",
            "This study therefore measures four outcomes under a single controlled precision contrast: "
            "harmful compliance, benign over-refusal, MMLU accuracy, and ARC-Challenge accuracy. Prompts are "
            "paired between precision conditions within each benchmark, and the scorer is treated as an "
            "object of study rather than as a fixed instrument.",
        ),
        (
            "The safety threat model is HarmBench DirectRequest: models receive unsafe instructions without "
            "an optimised jailbreak suffix or attack loop.",
            "The safety threat model is HarmBench DirectRequest, under which models receive unsafe "
            "instructions without an optimised jailbreak suffix or attack loop.",
        ),
        (
            "The same concern applies to safety: an overall average is insufficient if scorer error is "
            "concentrated in one model family or response style.",
            "The same concern applies to safety because an overall average would be insufficient if scorer "
            "error were concentrated in one model family or response style.",
        ),
        (
            "The appropriate comparison is therefore qualitative: both studies find that four-bit "
            "compression need not create a large general safety regression, while both leave room for "
            "model- and dimension-specific effects.",
            "The appropriate comparison is therefore qualitative. Both studies find that four-bit "
            "compression need not create a large general safety regression, while leaving room for "
            "model- and dimension-specific effects.",
        ),
        (
            "The main gap concerns interpretation: across the studies reviewed there, a reported metric "
            "change is not consistently separated from capability, scorer error, truncation, or multiplicity.",
            "The main gap concerns interpretation because the studies reviewed do not consistently separate "
            "a reported metric change from capability, scorer error, truncation, or multiplicity.",
        ),
        (
            "Accordingly, the conclusions of this report apply specifically to standard direct requests in "
            "unmodified checkpoints: detectable capability losses occurred, but no statistically significant "
            "increases in harmful compliance were observed after correction.",
            "Accordingly, the conclusions of this report apply specifically to standard direct requests in "
            "unmodified checkpoints. Detectable capability losses occurred, but no statistically significant "
            "increases in harmful compliance were observed after correction.",
        ),
        (
            "The primary NF4 set of corrected comparisons contains 20 contrasts: five model pairs multiplied "
            "by four outcomes.",
            "The primary NF4 set of corrected comparisons contains 20 contrasts, corresponding to five model "
            "pairs evaluated on four outcomes.",
        ),
        (
            "Each section follows the same order: finding, magnitude, statistical support, and interpretation.",
            "Each section presents the finding, magnitude, statistical support, and interpretation in the "
            "same order.",
        ),
        (
            "The aggregate null does not mean every prompt is stable: sixteen prompts flip toward harmful "
            "compliance and sixteen flip toward safety, producing McNemar p = 1.000.",
            "The aggregate null does not mean that every prompt is stable. Sixteen prompts flip toward "
            "harmful compliance and sixteen flip toward safety, producing McNemar p = 1.000.",
        ),
        (
            "Section 4.4 explains the caution: ARC accuracy falls by 0.0316 in the same precision contrast.",
            "Section 4.4 provides the necessary context because ARC accuracy falls by 0.0316 in the same "
            "precision contrast.",
        ),
        (
            "Llama-3.2-3B has an ARC contrast of −0.032 (unrounded: −0.0316),",
            "Llama-3.2-3B has an ARC contrast of −0.032 (unrounded value −0.0316),",
        ),
        (
            "However, that contrast is much smaller on ARC: the 1.7B decrease is −0.0094 and the 4B "
            "decrease is −0.0162.",
            "However, that contrast is much smaller on ARC, where the 1.7B decrease is −0.0094 and the "
            "4B decrease is −0.0162.",
        ),
        (
            "The remaining classifications are based only on the direction of the point estimates: "
            "Qwen3-4B and Phi show point estimates in the direction of higher harmful compliance, while "
            "Mistral has a point estimate in the direction of lower harmful compliance.",
            "The remaining classifications are based only on the direction of the point estimates. "
            "Qwen3-4B and Phi show point estimates in the direction of higher harmful compliance, while "
            "Mistral has a point estimate in the direction of lower harmful compliance.",
        ),
        (
            "Mistral is the clearest case: its refusal proxy is far more pessimistic than the semantic classifier.",
            "Mistral provides the clearest example because its refusal proxy is far more pessimistic than "
            "the semantic classifier.",
        ),
        (
            "Figure 4.2 shows that harmful compliance does not follow one shared path across fp16, INT8, and "
            "NF4: two pairs move non-monotonically, and the rest disagree with each other in direction.",
            "Figure 4.2 shows that harmful compliance does not follow one shared path across fp16, INT8, and "
            "NF4 because two pairs move non-monotonically and the others disagree in direction.",
        ),
        (
            "However, this result is method-specific: LLM.int8 and NF4 use different numerical treatments, "
            "so the pattern should not be simplified to a universal bit-width curve.",
            "However, this result is method-specific because LLM.int8 and NF4 use different numerical "
            "treatments, so the pattern should not be simplified to a universal bit-width curve.",
        ),
        (
            "Qwen3-1.7B illustrates the large capability decrease: MMLU is slightly higher at INT8 than "
            "fp16, then drops by more than ten points relative to INT8 at NF4.",
            "Qwen3-1.7B illustrates the large capability decrease. MMLU is slightly higher at INT8 than "
            "fp16, then drops by more than ten points relative to INT8 at NF4.",
        ),
        (
            "This budget-sensitive non-replication is attributed to cross-run greedy divergence rather than "
            "truncation: Llama has the study's lowest 128-token truncation rate (3–4%) and the largest "
            "prefix-mismatch count.",
            "This budget-sensitive non-replication is attributed to cross-run greedy divergence rather than "
            "truncation because Llama has the study's lowest 128-token truncation rate (3–4%) and the largest "
            "prefix-mismatch count.",
        ),
        (
            "This chapter tests the central assumptions directly: whether the classifier captures harmful "
            "compliance better than a refusal proxy, whether the XSTest regex captures refusal, and whether "
            "the conclusions persist across judges, generation budgets, and decoding seeds.",
            "This chapter directly tests whether the classifier captures harmful compliance better than a "
            "refusal proxy, whether the XSTest regex captures refusal, and whether the conclusions persist "
            "across judges, generation budgets, and decoding seeds.",
        ),
        (
            "This preserves a record of both scoring versions: the classifier result supports the current "
            "claim, while the proxy remains visible as evidence for why the interpretation changed.",
            "This preserves a record of both scoring versions. The classifier result supports the current "
            "claim, while the proxy remains visible as evidence for why the interpretation changed.",
        ),
        (
            "Qwen3-1.7B is the exception: the classifier is flat and Llama Guard decreases by 0.050, and "
            "neither result is significant.",
            "Qwen3-1.7B is the exception. The classifier is flat and Llama Guard decreases by 0.050, and "
            "neither result is significant.",
        ),
        (
            "The key takeaway is the relative fidelity of the evaluators: the official classifier exhibits "
            "substantially higher agreement with the human annotator than the keyword refusal proxy.",
            "The key takeaway is that the official classifier exhibits substantially higher agreement with "
            "the human annotator than the keyword refusal proxy.",
        ),
        (
            "This error profile matches the conceptual concern: absence of a refusal phrase is treated as "
            "harmful completion.",
            "This error profile matches the conceptual concern that absence of a refusal phrase is treated "
            "as harmful completion.",
        ),
        (
            "The XSTest human validation protocol incorporated one disclosed procedural deviation: the "
            "blindness exclusion was applied across all responses for an affected prompt rather than "
            "restricting exclusions to specific model-response pairs.",
            "The XSTest human validation protocol incorporated one disclosed procedural deviation. The "
            "blindness exclusion was applied across all responses for an affected prompt rather than "
            "restricting exclusions to specific model-response pairs.",
        ),
        (
            "In a validation-informed sensitivity analysis where the keyword regex is replaced by the "
            "validated strict semantic judge, exactly two contrasts survive Benjamini–Hochberg correction "
            "across the entire study: the Qwen MMLU and Llama ARC capability declines.",
            "In a validation-informed sensitivity analysis where the keyword regex is replaced by the "
            "validated strict semantic judge, exactly two contrasts survive Benjamini–Hochberg correction "
            "across the entire study. These are the Qwen MMLU and Llama ARC capability declines.",
        ),
        (
            "The reported null finding is supported along four partly independent dimensions: no HarmBench "
            "contrast survives multiplicity correction, the primary classifier and two further judges find "
            "no significant ASR increases, the 128-token Qwen3-1.7B artifact does not survive the move to the "
            "512-token reference budget, and the five-seed arm shows pair-specific rather than common "
            "movement: Qwen3-4B is positive in 5/5 seeds (1 individually significant), whereas Llama is "
            "negative in 5/5 (2 individually significant).",
            "Four partly independent findings support the reported null. No HarmBench contrast survives "
            "multiplicity correction; the primary classifier and two further judges find no significant ASR "
            "increases; and the 128-token Qwen3-1.7B artifact does not survive the move to the 512-token "
            "reference budget. The five-seed arm also shows pair-specific rather than common movement. "
            "Qwen3-4B is positive in 5/5 seeds (1 individually significant), whereas Llama is negative in "
            "5/5 (2 individually significant).",
        ),
        (
            "Each validation experiment addresses one competing explanation: human annotations expose the "
            "over-flagging bias of refusal proxies, cross-model judges test whether the result depends on a "
            "single evaluator, the 512-token protocol removes the truncation that affected the 128-token "
            "runs, and multi-seed testing bounds sampling noise.",
            "Each validation experiment addresses one competing explanation. Human annotations expose the "
            "over-flagging bias of refusal proxies, cross-model judges test whether the result depends on a "
            "single evaluator, the 512-token protocol removes the truncation that affected the 128-token "
            "runs, and multi-seed testing bounds sampling noise.",
        ),
        (
            "The results support neither common oversimplification: neither that quantization inherently "
            "improves safety alignment, nor that it systematically degrades it.",
            "The results support neither the claim that quantization inherently improves safety alignment "
            "nor the claim that it systematically degrades it.",
        ),
        (
            "The study's strongest generalisable contribution is methodological: match the weights, evaluate "
            "safety and capability together, pair the statistics, declare multiplicity, and validate the scorer.",
            "The study's strongest generalisable contribution lies in its method of matching the weights, "
            "evaluating safety and capability together, pairing the statistics, declaring multiplicity, and "
            "validating the scorer.",
        ),
        (
            "The practical takeaway for system designers is procedural: quantized language models should be "
            "evaluated in their exact deployment configuration, combining task-specific capability benchmarks "
            "with semantic safety evaluators tuned to the target operational domain.",
            "For system designers, the practical takeaway is that quantized language models should be "
            "evaluated in their exact deployment configuration, combining task-specific capability benchmarks "
            "with semantic safety evaluators tuned to the target operational domain.",
        ),
        (
            "The committed results_512 tree contains 60 benchmark summaries: five matched fp16–NF4 pairs and "
            "five additional INT8 aliases, each evaluated on HarmBench, XSTest, the six-subject MMLU subset, "
            "and ARC-Challenge.",
            "The committed results_512 tree contains 60 benchmark summaries, comprising five matched "
            "fp16–NF4 pairs and five additional INT8 aliases evaluated on HarmBench, XSTest, the six-subject "
            "MMLU subset, and ARC-Challenge.",
        ),
        (
            "The principal result is stable across the committed analysis surfaces: no harmful-compliance "
            "contrast survives Benjamini–Hochberg false-discovery-rate correction at the 512-token reference "
            "budget.",
            "The principal result is stable across the committed analysis surfaces. No harmful-compliance "
            "contrast survives Benjamini–Hochberg false-discovery-rate correction at the 512-token reference "
            "budget.",
        ),
        (
            "The corrected family contains 3 survivors, all outside harmful compliance: the Qwen3-1.7B MMLU "
            "decrease, the Llama-3.2-3B ARC decrease, and the Phi-4-mini regex-scored over-refusal decrease.",
            "The corrected family contains 3 survivors, all outside harmful compliance. They are the "
            "Qwen3-1.7B MMLU decrease, the Llama-3.2-3B ARC decrease, and the Phi-4-mini regex-scored "
            "over-refusal decrease.",
        ),
    )
    for old, new in rewrites:
        hits = _replace_all_text(doc, old, new)
        if hits != 1:
            raise AssertionError(f"Colon-style rewrite expected one hit, found {hits}: {old!r}")


def _refresh_package_metadata(path: Path, doc: Document) -> None:
    """Replace inherited template metadata with deterministic artifact metadata."""
    paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    text = "\n".join(paragraph.text for paragraph in paragraphs if paragraph.text)
    stats = {
        "Pages": PHYSICAL_PAGE_COUNT,
        "Words": len(re.findall(r"\b[\w’'-]+\b", text)),
        "Characters": sum(not char.isspace() for char in text),
        "CharactersWithSpaces": len(text),
        "Paragraphs": sum(bool(paragraph.text.strip()) for paragraph in paragraphs),
    }

    with zipfile.ZipFile(path, "r") as source:
        infos = source.infolist()
        members = {info.filename: source.read(info.filename) for info in infos}

    app_name = "docProps/app.xml"
    app = ET.fromstring(members[app_name])
    app_ns = "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"
    for name, value in stats.items():
        element = app.find(f"{{{app_ns}}}{name}")
        if element is None:
            element = ET.SubElement(app, f"{{{app_ns}}}{name}")
        element.text = str(value)
    members[app_name] = ET.tostring(app, encoding="utf-8", xml_declaration=True)

    bibliography_name = "customXml/item1.xml"
    bibliography = ET.fromstring(members[bibliography_name])
    bibliography.set("SelectedStyle", "/IEEE.XSL")
    bibliography.set("StyleName", "IEEE")
    members[bibliography_name] = ET.tostring(bibliography, encoding="utf-8", xml_declaration=True)

    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temporary = Path(handle.name)
    try:
        with zipfile.ZipFile(temporary, "w") as target:
            for info in infos:
                target.writestr(info, members[info.filename])
        temporary.replace(path)
    finally:
        temporary.unlink(missing_ok=True)


def _move_recent_blocks_before(doc: Document, reference, previous_children: list) -> None:
    body = doc.element.body
    old = set(previous_children)
    added = [child for child in body.iterchildren() if child not in old and child.tag != qn("w:sectPr")]
    if not added:
        raise AssertionError("Expected newly-added document blocks")
    for element in added:
        reference._element.addprevious(element)


def _add_table_before(doc: Document, reference, v1, *args, **kwargs) -> None:
    before = list(doc.element.body.iterchildren())
    v1.add_table(doc, *args, **kwargs)
    _move_recent_blocks_before(doc, reference, before)


def _insert_progress_chapter(doc: Document, v1, evidence: dict) -> None:
    chapter = _paragraph(doc, "Chapter 7: Conclusion and Future Work")
    _set_text(chapter, "Chapter 7: Progress and Remaining Work")
    references = _paragraph(doc, "References")
    current = chapter._element.getnext()
    while current is not None and current is not references._element:
        next_el = current.getnext()
        current.getparent().remove(current)
        current = next_el
    if current is None:
        raise AssertionError("References heading did not follow Chapter 7")

    references.insert_paragraph_before("7.1 Progress Achieved", style="Heading 2")
    references.insert_paragraph_before(
        "At the interim-submission point, the planned experimental matrix is complete. The committed "
        f"results_512 tree contains {evidence['summary_count']} benchmark summaries: five matched fp16–NF4 "
        "pairs and five additional INT8 aliases, each evaluated on HarmBench, XSTest, the six-subject MMLU "
        "subset, and ARC-Challenge. The primary paired analysis and the scorer-validation layers have also "
        "been completed. This means the remaining FYP work is principally verification, supervisor-led "
        "revision, and submission rather than further execution of the registered matrix.",
        style="Normal",
    )
    references.insert_paragraph_before(
        "The principal result is stable across the committed analysis surfaces: no harmful-compliance "
        "contrast survives Benjamini–Hochberg false-discovery-rate correction at the 512-token reference "
        f"budget. The corrected family contains {evidence['survivor_count']} survivors, all outside harmful "
        "compliance: the Qwen3-1.7B MMLU decrease, the Llama-3.2-3B ARC decrease, and the Phi-4-mini "
        "regex-scored over-refusal decrease. The last result is treated as scorer-dependent because it is "
        "not reproduced by the semantic refusal judge and the human audit favours that judge over the regex.",
        style="Normal",
    )
    _add_table_before(
        doc,
        references,
        v1,
        "7.1",
        "Completed work and repository evidence",
        ["Workstream", "Status", "Primary repository evidence"],
        [
            ("Experiment matrix", "Complete", "results_512/<alias>/<benchmark>/summary.json (60 summaries)"),
            ("Primary paired analysis", "Complete", "results_512/analysis/multiple_comparisons.json; pairwise_deltas.json"),
            ("HarmBench scorer validation", "Complete", "human_validation.json; judge_pairwise_agreement*.json"),
            ("XSTest scorer validation", "Complete", "xstest_judge_agreement.json; xstest_human_validation.json"),
            ("Robustness checks", "Complete", "genlen_robustness.json; sensitivity_multiseed.json; precision_sweep.json"),
        ],
        [3.4, 2.4, 9.8],
        font_size=7.7,
    )

    references.insert_paragraph_before("7.2 Problems Encountered and Resolutions", style="Heading 2")
    references.insert_paragraph_before(
        "Three issues materially changed the analysis. First, the original 128-token runs frequently ended "
        "at the generation limit; in that 128-token tree, the apparent Qwen3-1.7B harmful-compliance increase was +0.055. "
        "the matched 512-token rerun gives a classifier delta of 0.000. The 512-token tree therefore became "
        "the primary study, while the earlier tree is retained as a generation-length comparison. Second, a "
        "refusal regex over-counted harmful compliance, so the official HarmBench classifier became the "
        "primary scorer and was checked against an API judge, Llama Guard, and a human-labelled sample. Third, "
        "the primary XSTest regex and a semantic judge disagreed on the Phi-4-mini contrast; the separate human "
        "audit aligned much more closely with the semantic judge. The report consequently preserves the "
        "pre-specified regex result but identifies it as scorer-dependent rather than treating it as a robust "
        "alignment effect.",
        style="Normal",
    )
    _add_table_before(
        doc,
        references,
        v1,
        "7.2",
        "Problems encountered and evidence-led responses",
        ["Problem", "Response", "Evidence"],
        [
            ("Generation truncation at 128 tokens", "Re-ran the matched study at 512 tokens and made it primary", "genlen_robustness.json"),
            ("Refusal proxy over-counted harmful compliance", "Adopted the official classifier and triangulated it", "judge_agreement.json; human_validation.json"),
            ("Phi over-refusal result depended on scorer", "Retained the registered result but added judge and human sensitivity", "xstest_judge_agreement.json; xstest_human_validation.json"),
        ],
        [4.0, 6.5, 5.1],
        font_size=7.6,
    )

    references.insert_paragraph_before("7.3 Remaining Work", style="Heading 2")
    references.insert_paragraph_before(
        "The scientific submission is not blocked by an unfinished benchmark run. The immediate work is to "
        "submit this interim document, collect the supervisor's feedback, and incorporate justified revisions "
        "into the final-report line. Two repository actions remain useful before the final freeze: record the "
        "exact TC1 storage quota with MyTCinfo, and, if time permits, add a second annotator on the planned "
        "overlap slice so inter-rater agreement can supplement the disclosed single-annotator limitation. "
        "Residual resume/configuration hardening is tracked separately and is not a blocker for the current "
        "scientific claims. Broader quantization methods, additional model scales, jailbreak evaluations, and "
        "wider stochastic-decoding work remain future-study extensions rather than evidence claimed here.",
        style="Normal",
    )

    references.insert_paragraph_before("7.4 Submission Timeline", style="Heading 2")
    _add_table_before(
        doc,
        references,
        v1,
        "7.3",
        "Remaining milestones recorded in PROJECT_LOG.md",
        ["Date or period", "Milestone", "Planned action"],
        [
            ("24 Aug 2026", "Interim report", "Submit to Dr. Zhang and obtain feedback"),
            ("Aug–Oct 2026", "Revision period", "Address feedback; record TC1 quota; complete optional second-annotator overlap if feasible"),
            ("19 Oct 2026", "Final report", "Submit the graded softcopy to the supervisor and examiners"),
            ("13 Nov 2026", "Amended final", "Submit amendments and deposit the electronic copy in DR-NTU"),
            ("4, 7 or 8 Dec 2026", "Oral presentation", "Deliver the 20-minute presentation and 10-minute question-and-answer session"),
        ],
        [3.3, 3.9, 8.4],
        font_size=7.7,
    )


def _rebuild_toc(doc: Document, v1) -> None:
    """Replace the inherited TOC with a stable, verified compact page map."""
    toc_heading = _paragraph(doc, "Table of Contents")
    current = toc_heading._element.getnext()
    boundary = None
    while current is not None:
        if current.tag == qn("w:p") and current.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
            boundary = current
            break
        next_el = current.getnext()
        current.getparent().remove(current)
        current = next_el
    if boundary is None:
        raise AssertionError("Could not locate the front/main section break after the TOC")

    def append_before_boundary(create) -> None:
        before = list(doc.element.body.iterchildren())
        create()
        old = set(before)
        added = [
            child
            for child in doc.element.body.iterchildren()
            if child not in old and child.tag != qn("w:sectPr")
        ]
        if len(added) != 1:
            raise AssertionError(f"Expected one TOC block, found {len(added)}")
        boundary.addprevious(added[0])

    # The source report requests automatic field updates on open. That setting
    # is unsafe for this generated multiline TOC in Word, which resolves every
    # heading target to page 1. Remove it and emit the verified map as ordinary
    # paragraphs so opening or saving the DOCX cannot corrupt the page numbers.
    update_fields = doc.settings.element.find(qn("w:updateFields"))
    if update_fields is not None:
        doc.settings.element.remove(update_fields)

    headings = [
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.style and paragraph.style.name in {"Heading 1", "Heading 2"}
    ]
    for heading in headings:
        title = heading.text.strip()
        if title not in TOC_PAGE_NUMBERS:
            raise AssertionError(f"No cached TOC page for retained heading {title!r}")
        level = 0 if heading.style.name == "Heading 1" else 1
        append_before_boundary(
            lambda title=title, level=level: v1.add_contents_entry(
                doc,
                title,
                str(TOC_PAGE_NUMBERS[title]),
                level=level,
                bold=level == 0,
            )
        )

def _assert_compact_structure(doc: Document) -> None:
    headings = [
        paragraph.text
        for paragraph in doc.paragraphs
        if paragraph.style and paragraph.style.name.startswith("Heading")
    ]
    required = {
        "Chapter 1: Introduction",
        "Chapter 3: Methodology and Experimental Setup",
        "Chapter 4: Results",
        "Chapter 5: Scorer Validation and Robustness",
        "Chapter 6: Discussion",
        "Chapter 7: Progress and Remaining Work",
        "7.3 Remaining Work",
        "7.4 Submission Timeline",
        "References",
    }
    missing = sorted(required - set(headings))
    if missing:
        raise AssertionError(f"Compact interim is missing required headings: {missing}")
    forbidden = {
        "1.5 Contributions",
        "1.6 Report Structure",
        "3.6 Implementation and Execution",
        "3.7 Reproducibility Controls",
        "3.8 Analysis Workflow and Quality Control",
        "5.6 Multi-Seed Decoding Sensitivity",
        "6.2 Relationship to Prior Work",
        "6.3 Capability Effects and Deployment",
        "6.4 Validity of the Scoring Methods",
        "Appendix A: Reproducibility Summary",
    }
    residue = sorted(forbidden & set(headings))
    if residue:
        raise AssertionError(f"Compact interim retained excluded detail: {residue}")
    text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
    for stale in ("List of Figures", "List of Tables", "Abbreviations", "Chapter 7: Conclusion and Future Work"):
        if stale in text:
            raise AssertionError(f"Compact interim retained stale front/body label: {stale!r}")
    for dangling in ("Table 2.1", "Figure 3.1", "Figure 5.1", "Table A.1", "Appendix A", "Appendix B", "Appendix C"):
        if dangling in text:
            raise AssertionError(f"Compact interim cites removed material: {dangling!r}")


def build_interim() -> Path:
    evidence = _load_progress_evidence()
    v4 = _load_module(V4_BUILDER, "build_rewritten_fyp_report_v4_for_compact")
    v1 = _load_module(V1_BUILDER, "build_rewritten_fyp_report_v1_for_compact")

    source_hash = _sha256(ROOT / "docs" / "rewritten_fyp_interim_v4.docx")
    with tempfile.TemporaryDirectory(prefix="fyp_interim_compact_") as tmp:
        tmp_report = Path(tmp) / "v4_report.docx"
        original_out = v4.OUT
        try:
            v4.OUT = tmp_report
            v4.build_report()
        finally:
            v4.OUT = original_out

        doc = Document(tmp_report)
        _retitle(doc)

        _remove_range(doc, "Acknowledgements", "Abstract")
        _remove_front_lists(doc)
        _remove_range(doc, "1.5 Contributions", "Chapter 2: Literature Review")

        _remove_range(doc, "2.1 Trustworthiness and Safety Evaluation", "2.2 Quantization Methods and Capability Effects")
        _remove_range(doc, "2.4 Refusal, Harmfulness, and Threat Models", "2.6 Synthesis and Research Gap")
        _remove_range(doc, "2.7 Implications for Experimental Design", "Chapter 3: Methodology and Experimental Setup")

        _remove_range(
            doc,
            "Table 2.1 compares the studies relevant to this experiment rather than a claim that this design is unique. It records whether the cited papers' abstracts and stated scope explicitly establish each listed dimension. “Not reported” means that the consulted source did not establish the property; it is not evidence that the authors did not consider it.",
            "The literature review identifies five gaps. First, the quantized and full-precision models should come from the same checkpoint revision. Second, the quantization method and precision must be explicit. Third, harmful compliance, benign refusal, and capability should be measured together. Fourth, the conclusion should be tested against scorer choice and human labels. Fifth, the report must distinguish direct requests from jailbreak robustness.",
        )

        _remove_range(doc, "3.6 Implementation and Execution", "Chapter 4: Results")
        _remove_figure(
            doc,
            "Figure 3.1. Matched pairs, precision arms, benchmark battery, and analysis flow.",
        )
        _remove_range(doc, "5.6 Multi-Seed Decoding Sensitivity", "5.7 Power and the Meaning of the Null")
        figure_intro = _paragraph(
            doc,
            "Figure 5.1 compares the HarmBench classifier with the refusal-regex proxy. The proxy systematically reports higher attack-success rates for several aliases because it treats non-refusal as compliance. However, that is not the construct HarmBench defines: a response may avoid an explicit refusal and still fail to provide the requested harmful information.",
        )
        _set_text(
            figure_intro,
            "The HarmBench classifier and the refusal-regex proxy diverge because the proxy systematically "
            "reports higher attack-success rates for several aliases by treating non-refusal as compliance. "
            "That is not the construct HarmBench defines. The classifier asks whether a response contains "
            "substantive harmful content that fulfils the behaviour, so disclaimers, vague discussion, and "
            "safe redirection are not positive cases.",
        )
        _remove_figure(
            doc,
            "Figure 5.1. HarmBench classifier ASR versus the refusal-regex proxy and their agreement.",
        )
        _remove_range(doc, "6.2 Relationship to Prior Work", "6.5 Threats to Validity")
        _remove_appendices(doc)

        _rename_heading(doc, "2.2 Quantization Methods and Capability Effects", "2.1 Quantization Methods and Capability Effects")
        _rename_heading(doc, "2.3 Quantization and Safety Behaviour", "2.2 Quantization and Safety Behaviour")
        _rename_heading(doc, "2.6 Synthesis and Research Gap", "2.3 Synthesis and Research Gap")
        _rename_heading(doc, "5.7 Power and the Meaning of the Null", "5.6 Power and the Meaning of the Null")
        _rename_heading(doc, "5.8 Robustness Synthesis", "5.7 Robustness Synthesis")
        _rename_heading(doc, "6.5 Threats to Validity", "6.2 Threats to Validity")
        _rename_heading(doc, "6.5.1 Internal Validity", "6.2.1 Internal Validity")
        _rename_heading(doc, "6.5.2 Construct Validity", "6.2.2 Construct Validity")
        _rename_heading(doc, "6.5.3 External Validity", "6.2.3 External Validity")
        _rename_heading(doc, "6.6 Practical Implications", "6.3 Practical Implications")

        for old, new in (
            ("Section 5.7", "Section 5.6"),
            ("Section 6.5.1", "Section 6.2.1"),
            ("Section 6.5.2", "Section 6.2.2"),
            ("Section 6.5.3", "Section 6.2.3"),
            ("Section 6.5", "Section 6.2"),
            ("Section 6.6", "Section 6.3"),
            ("the absence, in Table 2.1, of the same benchmark combination", "the absence of the same benchmark combination in prior work"),
            (
                "Table A.1 records the quantization settings and the declared dependency constraints.",
                "The quantization settings and dependency constraints are recorded in the revision-pinned configuration and run metadata.",
            ),
            (" Figure 3.1 shows the study design.", ""),
            ("Appendix B therefore reports the thresholds", "The analysis therefore reports the thresholds"),
            ("As detailed in Appendix B, observed capability changes", "Observed capability changes"),
        ):
            _replace_all_text(doc, old, new)

        _insert_progress_chapter(doc, v1, evidence)
        _apply_correctness_audit_repairs(doc, evidence)
        _rewrite_colon_heavy_prose(doc)
        _rebuild_toc(doc, v1)

        citation_summary = v4._renumber_citations(doc)
        v4._assert_ieee_order(doc)
        _assert_compact_structure(doc)

        doc.core_properties.title = "Benchmarking Ethical Performance of Open-Source LLMs"
        doc.core_properties.subject = "Compact interim progress report"
        doc.core_properties.author = "TAN UEI HORNG"
        doc.core_properties.last_modified_by = "TAN UEI HORNG"
        doc.core_properties.created = datetime(2026, 8, 24)
        doc.core_properties.modified = datetime(2026, 8, 24)
        doc.core_properties.comments = (
            "Compact interim derived non-destructively from the verified v4 report. Progress and remaining "
            "work are generated only after repository evidence and PROJECT_LOG checks pass."
        )
        OUT.parent.mkdir(parents=True, exist_ok=True)
        doc.save(OUT)
        _refresh_package_metadata(OUT, doc)

    if _sha256(ROOT / "docs" / "rewritten_fyp_interim_v4.docx") != source_hash:
        raise AssertionError("The verified v4 interim changed during compact build")
    print(
        f"    compact citations: {citation_summary['kept']} references; "
        f"dropped uncited references: {citation_summary['dropped'] or 'none'}"
    )
    return OUT


if __name__ == "__main__":
    print(build_interim())
